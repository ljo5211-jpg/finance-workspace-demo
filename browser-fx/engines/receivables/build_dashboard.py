from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import statistics
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = ROOT / "dashboard-template.html"
WORKING_HTML_PATH = ROOT / "receivables-dashboard-layout.html"
PRIVATE_DIR = ROOT / "private"
PUBLISH_DIR = ROOT / "publish"
DB_PATH = PRIVATE_DIR / "receivables.db"
CONFIG_PATH = ROOT / "dashboard-settings.json"

SOURCE_TYPE_KEYWORDS = {
    "aging": ("에이징", "aging"),
    "credit": ("여신", "credit"),
    "allowance": ("대손충당금", "allowance"),
}
SOURCE_TYPE_LABELS = {
    "aging": "에이징",
    "credit": "여신",
    "allowance": "대손충당금",
}
FILENAME_DATE_PATTERNS = (
    re.compile(r"(?<!\d)(?P<year>\d{4})(?P<month>\d{2})(?P<day>\d{2})(?!\d)"),
    re.compile(r"(?<!\d)(?P<year>\d{2})(?P<month>\d{2})(?P<day>\d{2})(?!\d)"),
    re.compile(r"(?<!\d)(?P<year>\d{4})[.\-_/\s]+(?P<month>\d{1,2})[.\-_/\s]+(?P<day>\d{1,2})(?!\d)"),
    re.compile(r"(?<!\d)(?P<year>\d{2})[.\-_/\s]+(?P<month>\d{1,2})[.\-_/\s]+(?P<day>\d{1,2})(?!\d)"),
    re.compile(r"(?<!\d)(?P<year>\d{2}|\d{4})\s*년\s*(?P<month>\d{1,2})\s*월\s*(?P<day>\d{1,2})\s*일?"),
)
TRADE_ACCOUNT_RE = re.compile(r"매출채권|미수금\s*\(거래처\)")
EXCLUDED_ACCOUNT_RE = re.compile(r"단기대여금|임직원")


@dataclass
class ImportResult:
    snapshot_date: date
    source_type: str
    file_name: str
    file_hash: str
    raw_rows: int
    included_rows: int
    excluded_rows: int
    amount_sum: float


def parse_args() -> argparse.Namespace:
    default_source = ROOT / "적재"
    parser = argparse.ArgumentParser(description="SAP 채권 Aging·여신 스냅샷을 SQLite와 공유용 HTML로 생성합니다.")
    parser.add_argument("--source-dir", type=Path, default=default_source, help="SAP XLSX가 저장된 폴더")
    parser.add_argument("--usd-krw-rate", type=float, default=None, help="원화환산용 USD/KRW 관리환율")
    parser.add_argument("--allowance-file", type=Path, default=None, help="전이율계산·주석 시트가 있는 대손충당금 XLSX")
    parser.add_argument("--latest-date-only", action="store_true", help="최종 Aging 기준일만 출력")
    return parser.parse_args()


def load_config() -> dict[str, Any]:
    defaults = {
        "usdKrwRate": 1400.0,
        "rateBasis": "재무팀 관리환율(수동)",
        "rateDate": "2026-08-18",
        "rateQuoteDate": "2026-08-18",
    }
    if not CONFIG_PATH.exists():
        return defaults
    loaded = json.loads(CONFIG_PATH.read_text(encoding="utf-8-sig"))
    return {**defaults, **loaded}


def source_type_from_name(name: str) -> str | None:
    """Use the left-most recognized type keyword, regardless of its position."""
    lowered = name.casefold()
    matches: list[tuple[int, int, int, str]] = []
    for type_order, (source_type, keywords) in enumerate(SOURCE_TYPE_KEYWORDS.items()):
        for keyword in keywords:
            position = lowered.find(keyword.casefold())
            if position >= 0:
                matches.append((position, -len(keyword), type_order, source_type))
    return min(matches)[3] if matches else None


def filename_date(name: str) -> date | None:
    """Return the left-most valid YMD date in common Korean filename formats."""
    candidates: list[tuple[int, int, date]] = []
    stem = Path(name).stem
    for pattern_order, pattern in enumerate(FILENAME_DATE_PATTERNS):
        for match in pattern.finditer(stem):
            year = int(match.group("year"))
            if year < 100:
                year += 2000
            try:
                parsed = date(year, int(match.group("month")), int(match.group("day")))
            except ValueError:
                continue
            candidates.append((match.start(), pattern_order, parsed))
    return min(candidates)[2] if candidates else None


def parse_source_filename(name: str) -> tuple[str, date] | None:
    """Classify an XLSX using the first type keyword and first valid date."""
    if name.startswith("~$") or Path(name).suffix.lower() != ".xlsx":
        return None
    source_type = source_type_from_name(name)
    parsed_date = filename_date(name)
    if source_type is None or parsed_date is None:
        return None
    return source_type, parsed_date


def source_workbook_date(path: Path, expected_type: str) -> date:
    parsed = parse_source_filename(path.name)
    if parsed is None or parsed[0] != expected_type:
        label = SOURCE_TYPE_LABELS[expected_type]
        raise ValueError(f"{label} 파일명에서 자료종류와 기준일을 인식하지 못했습니다: {path.name}")
    return parsed[1]


def discover_source_workbooks(source_dir: Path) -> tuple[dict[str, list[Path]], list[str]]:
    """Discover usable workbooks without stopping on unrelated or malformed files."""
    discovered = {source_type: [] for source_type in SOURCE_TYPE_KEYWORDS}
    warnings: list[str] = []
    selected_keys: dict[tuple[str, date], Path] = {}
    for path in sorted(source_dir.iterdir(), key=lambda item: item.name.casefold()):
        if not path.is_file() or path.name.startswith("~$") or path.suffix.lower() != ".xlsx":
            continue
        source_type = source_type_from_name(path.name)
        if source_type is None:
            continue
        parsed_date = filename_date(path.name)
        if parsed_date is None:
            warnings.append(f"기준일을 찾지 못해 건너뜀: {path.name}")
            continue
        key = (source_type, parsed_date)
        if key in selected_keys:
            warnings.append(
                f"동일 자료·기준일 중 첫 파일을 사용하고 건너뜀: {path.name} "
                f"(사용: {selected_keys[key].name})"
            )
            continue
        selected_keys[key] = path
        discovered[source_type].append(path)
    for source_type, paths in discovered.items():
        paths.sort(key=lambda path: (source_workbook_date(path, source_type), path.name.casefold()))
    return discovered, warnings


def latest_source_workbook(paths: list[Path], source_type: str) -> Path:
    """Return the newest dated workbook independently of other source types."""
    if not paths:
        raise ValueError(f"{SOURCE_TYPE_LABELS[source_type]} 파일이 없습니다.")
    return max(
        paths,
        key=lambda path: (source_workbook_date(path, source_type), path.name.casefold()),
    )


def find_allowance_workbook(
    source_dir: Path,
    explicit_path: Path | None,
    discovered_candidates: list[Path] | None = None,
) -> Path:
    if explicit_path:
        candidate = explicit_path.expanduser().resolve()
        if not candidate.exists():
            raise FileNotFoundError(f"대손충당금 파일을 찾을 수 없습니다: {candidate}")
        source_workbook_date(candidate, "allowance")
        return candidate
    if discovered_candidates is None:
        discovered_candidates = discover_source_workbooks(source_dir)[0]["allowance"]
    candidates = sorted(discovered_candidates, key=lambda path: source_workbook_date(path, "allowance"), reverse=True)
    if not candidates:
        raise FileNotFoundError(
            "적재 폴더에서 날짜와 '대손충당금'이 포함된 XLSX 파일을 찾지 못했습니다."
        )
    return candidates[0]


def load_allowance_model(path: Path) -> dict[str, Any]:
    """Read the KPMG roll-rate workbook without modifying it."""
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
    values = openpyxl.load_workbook(path, data_only=True, read_only=False)
    required_sheets = {"전이율계산", "주석", "장기채권검토"}
    missing = required_sheets.difference(values.sheetnames)
    if missing:
        raise ValueError(f"{path.name}: 대손모형 필수 시트 누락 - {', '.join(sorted(missing))}")

    roll = values["전이율계산"]
    note = values["주석"]
    long_term = values["장기채권검토"]
    aging_analysis = values["연령분석표"]
    raw_rates = [roll.cell(row, 47).value for row in (46, 47, 48, 49)] + [roll["Z28"].value]
    applied_rates = list(raw_rates)
    prior_note_rates = [note.cell(row, 5).value for row in (6, 7, 8, 9, 10)]
    transition_steps = [roll.cell(row, 47).value for row in (21, 22, 23, 24)]
    for label, rates in (("2분기 Roll-rate", raw_rates), ("1분기 주석 적용률", prior_note_rates), ("전이확률", transition_steps)):
        if any(not isinstance(rate, (int, float)) or not 0 <= float(rate) <= 1 for rate in rates):
            raise ValueError(f"{path.name}: {label}을 숫자 0~1로 읽지 못했습니다.")
    if list(map(float, applied_rates)) != sorted(map(float, applied_rates)):
        raise ValueError(f"{path.name}: KPMG 권고 적용률이 연령구간 순서대로 증가하지 않습니다.")

    relevant_sheets = ("전이율계산", "주석", "장기채권검토", "연령분석표")
    formula_errors: list[str] = []
    for sheet_name in relevant_sheets:
        sheet = formulas[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "#REF!" in cell.value:
                    formula_errors.append(f"{sheet_name}!{cell.coordinate}")

    model_date_value = roll["AZ3"].value
    if isinstance(model_date_value, datetime):
        model_date = model_date_value.date().isoformat()
    elif isinstance(model_date_value, date):
        model_date = model_date_value.isoformat()
    else:
        model_date = str(model_date_value or "")
    file_date = source_workbook_date(path, "allowance").isoformat()
    if model_date != file_date:
        raise ValueError(
            f"{path.name}: 파일명 기준일 {file_date}과 "
            f"전이율계산 시트 모형기준일 {model_date or '없음'}이 다릅니다."
        )
    buckets = [
        {"key": "age_0_90", "label": "90일 이하", "fromDays": 0, "toDays": 90},
        {"key": "age_91_180", "label": "91–180일", "fromDays": 91, "toDays": 180},
        {"key": "age_181_270", "label": "181–270일", "fromDays": 181, "toDays": 270},
        {"key": "age_271_360", "label": "271–360일", "fromDays": 271, "toDays": 360},
        {"key": "age_361_plus", "label": "361일 이상", "fromDays": 361, "toDays": None},
    ]
    for index, bucket in enumerate(buckets):
        bucket["appliedRate"] = round(float(applied_rates[index]) * 100, 6)
        bucket["rawRate"] = round(float(raw_rates[index]) * 100, 6)
        bucket["priorRate"] = round(float(prior_note_rates[index]) * 100, 6)
        bucket["method"] = "개별평가" if bucket["key"] == "age_361_plus" else "집합평가"

    referenced_subtotal_rows: list[int] = []
    for coordinate in ("AU4", "AU5", "AU6", "AU7", "AU8", "AU9"):
        formula = formulas["전이율계산"][coordinate].value
        if isinstance(formula, str):
            referenced_subtotal_rows.extend(
                int(row_number)
                for row_number in re.findall(r"\$?[A-Z]{1,3}\$?(\d+)", formula)
            )
    last_included_row = max(referenced_subtotal_rows, default=0)
    excluded_customer_codes: list[str] = []
    excluded_exposure_krw = 0.0
    for row_number in range(last_included_row + 1, aging_analysis.max_row + 1):
        label = aging_analysis.cell(row_number, 1).value
        if label == "총계":
            break
        code_value = aging_analysis.cell(row_number, 3).value
        if code_value is None:
            continue
        code = str(code_value).strip()
        if code.endswith(".0") and code[:-2].isdigit():
            code = code[:-2]
        if code.isdigit():
            excluded_customer_codes.append(code.lstrip("0") or "0")
            excluded_exposure_krw += float(aging_analysis.cell(row_number, 5).value or 0)

    baseline_exposures = [
        float(roll.cell(row, 47).value or 0) for row in (4, 5, 6, 7)
    ] + [float(roll["AU8"].value or 0) + float(roll["AU9"].value or 0)]
    baseline_losses = [
        float(roll.cell(row, 26).value or 0) for row in (53, 54, 55, 56)
    ] + [float(roll["Z62"].value or 0)]
    baseline_buckets = [
        {
            "key": bucket["key"],
            "exposureKRW": baseline_exposures[index],
            "expectedLossKRW": baseline_losses[index],
        }
        for index, bucket in enumerate(buckets)
    ]
    return {
        "file": path.name,
        "modelDate": model_date,
        "basis": "전이율계산 시트 2026.2Q 누적평균 Roll-rate",
        "rawBasis": "전이율계산 시트 2026.2Q 누적평균 Roll-rate",
        "previousBasis": "주석 시트 2026.1Q 적용률",
        "buckets": buckets,
        "baselineBuckets": baseline_buckets,
        "transitionSteps": [round(float(value) * 100, 6) for value in transition_steps],
        "modelFxRate": float(roll["AY5"].value or 0),
        "excludedCustomerCodes": sorted(set(excluded_customer_codes)),
        "baselineExcludedExposureKRW": excluded_exposure_krw,
        "collectiveAllowance": float(roll["Z59"].value or 0),
        "individualLongTermAllowance": float(long_term["P71"].value or 0),
        "adjustedAllowance": float(roll["Z67"].value or 0),
        "controlChecks": {
            "rollforwardTie": roll["AU43"].value is True,
            "longTermTie": long_term["Q71"].value is True and long_term["R71"].value is True,
            "formulaReferenceWarnings": formula_errors,
        },
        "methodNote": "매출채권 계정의 순액을 대상으로 360일 이하는 집합평가, 361일 이상은 개별평가 100%를 적용",
    }


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_identifier(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def clean_customer_code(value: Any) -> str:
    text = clean_identifier(value)
    if text.isdigit():
        return text.lstrip("0") or "0"
    return text


def iso_date(value: Any) -> str | None:
    parsed = pd.to_datetime(value, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.date().isoformat()


def number(value: Any) -> float:
    parsed = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    return 0.0 if pd.isna(parsed) else float(parsed)


def detect_aging_date(frame: pd.DataFrame, path: Path) -> date:
    file_date = source_workbook_date(path, "aging")
    if "현재 날짜" in frame.columns:
        dates = pd.to_datetime(frame["현재 날짜"], errors="coerce").dropna()
        if not dates.empty:
            internal_dates = sorted({value.date() for value in dates})
            if len(internal_dates) != 1:
                rendered = ", ".join(value.isoformat() for value in internal_dates)
                raise ValueError(f"{path.name}: 엑셀 '현재 날짜'가 하나로 통일되지 않았습니다: {rendered}")
            internal_date = internal_dates[0]
            if internal_date != file_date:
                raise ValueError(
                    f"{path.name}: 파일명 기준일 {file_date.isoformat()}과 "
                    f"엑셀 '현재 날짜' {internal_date.isoformat()}가 다릅니다."
                )
    return file_date


def due_bucket(days: int | None) -> str:
    if days is None:
        return "unknown"
    if days < 0:
        return "not_due"
    if days == 0:
        return "due_today"
    if days <= 30:
        return "overdue_1_30"
    if days <= 60:
        return "overdue_31_60"
    if days <= 90:
        return "overdue_61_90"
    return "overdue_90_plus"


def invoice_bucket(days: float | None) -> str:
    if days is None or math.isnan(days):
        return "unknown"
    if days <= 30:
        return "invoice_1_30"
    if days <= 60:
        return "invoice_31_60"
    if days <= 90:
        return "invoice_61_90"
    return "invoice_91_plus"


def create_schema(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        PRAGMA foreign_keys = ON;
        PRAGMA journal_mode = DELETE;
        PRAGMA synchronous = NORMAL;

        CREATE TABLE import_batch (
            id INTEGER PRIMARY KEY,
            snapshot_date TEXT NOT NULL,
            source_type TEXT NOT NULL CHECK (source_type IN ('aging','credit')),
            source_filename TEXT NOT NULL,
            file_hash TEXT NOT NULL,
            raw_rows INTEGER NOT NULL,
            included_rows INTEGER NOT NULL,
            excluded_rows INTEGER NOT NULL,
            amount_sum REAL NOT NULL,
            imported_at TEXT NOT NULL,
            UNIQUE(snapshot_date, source_type)
        );

        CREATE TABLE ar_snapshot (
            snapshot_date TEXT NOT NULL,
            source_row INTEGER NOT NULL,
            customer_code TEXT NOT NULL,
            customer_name TEXT NOT NULL,
            document_no TEXT,
            invoice_date TEXT,
            posting_date TEXT,
            due_date TEXT,
            days_overdue INTEGER,
            aging_days_source REAL,
            payment_terms TEXT,
            payment_terms_desc TEXT,
            currency TEXT,
            recon_account TEXT,
            recon_account_name TEXT,
            scope TEXT NOT NULL CHECK (scope IN ('receivable','other')),
            amount REAL NOT NULL,
            due_bucket TEXT NOT NULL,
            invoice_bucket TEXT NOT NULL,
            PRIMARY KEY(snapshot_date, source_row)
        );

        CREATE TABLE credit_snapshot (
            snapshot_date TEXT NOT NULL,
            customer_code TEXT NOT NULL,
            customer_name TEXT,
            credit_control_area TEXT NOT NULL,
            hold_reason TEXT,
            credit_limit REAL NOT NULL,
            total_receivables REAL NOT NULL,
            order_amount REAL NOT NULL,
            delivery_amount REAL NOT NULL,
            billing_value REAL NOT NULL,
            special_exposure REAL NOT NULL,
            credit_used REAL NOT NULL,
            credit_available REAL NOT NULL,
            currency TEXT,
            PRIMARY KEY(snapshot_date, customer_code, credit_control_area)
        );

        CREATE TABLE monthly_summary (
            snapshot_date TEXT PRIMARY KEY,
            total_amount REAL NOT NULL,
            overdue_amount REAL NOT NULL,
            over90_amount REAL NOT NULL,
            not_due_amount REAL NOT NULL,
            due_today_amount REAL NOT NULL,
            overdue_1_30 REAL NOT NULL,
            overdue_31_60 REAL NOT NULL,
            overdue_61_90 REAL NOT NULL,
            invoice_1_30 REAL NOT NULL,
            invoice_31_60 REAL NOT NULL,
            invoice_61_90 REAL NOT NULL,
            invoice_91_plus REAL NOT NULL,
            customer_count INTEGER NOT NULL,
            row_count INTEGER NOT NULL
        );

        CREATE TABLE currency_summary (
            snapshot_date TEXT NOT NULL,
            currency TEXT NOT NULL,
            total_amount REAL NOT NULL,
            overdue_amount REAL NOT NULL,
            over90_amount REAL NOT NULL,
            not_due_amount REAL NOT NULL,
            due_today_amount REAL NOT NULL,
            overdue_1_30 REAL NOT NULL,
            overdue_31_60 REAL NOT NULL,
            overdue_61_90 REAL NOT NULL,
            invoice_1_30 REAL NOT NULL,
            invoice_31_60 REAL NOT NULL,
            invoice_61_90 REAL NOT NULL,
            invoice_91_plus REAL NOT NULL,
            customer_count INTEGER NOT NULL,
            row_count INTEGER NOT NULL,
            PRIMARY KEY(snapshot_date, currency)
        );

        CREATE TABLE customer_summary (
            snapshot_date TEXT NOT NULL,
            customer_code TEXT NOT NULL,
            customer_name TEXT NOT NULL,
            scope TEXT NOT NULL,
            total_amount REAL NOT NULL,
            overdue_amount REAL NOT NULL,
            over90_amount REAL NOT NULL,
            document_count INTEGER NOT NULL,
            PRIMARY KEY(snapshot_date, customer_code)
        );

        CREATE TABLE customer_currency_summary (
            snapshot_date TEXT NOT NULL,
            currency TEXT NOT NULL,
            customer_code TEXT NOT NULL,
            customer_name TEXT NOT NULL,
            scope TEXT NOT NULL,
            total_amount REAL NOT NULL,
            overdue_amount REAL NOT NULL,
            over90_amount REAL NOT NULL,
            gross_debit REAL NOT NULL,
            credit_amount REAL NOT NULL,
            document_count INTEGER NOT NULL,
            PRIMARY KEY(snapshot_date, currency, customer_code)
        );

        CREATE INDEX idx_ar_snapshot_customer_date
            ON ar_snapshot(customer_code, snapshot_date);
        CREATE INDEX idx_ar_snapshot_date_bucket
            ON ar_snapshot(snapshot_date, due_bucket);
        CREATE INDEX idx_credit_snapshot_customer_date
            ON credit_snapshot(customer_code, snapshot_date);
        CREATE INDEX idx_customer_summary_date_risk
            ON customer_summary(snapshot_date, over90_amount, overdue_amount);
        CREATE INDEX idx_ar_snapshot_date_currency_customer
            ON ar_snapshot(snapshot_date, currency, customer_code);
        CREATE INDEX idx_ar_snapshot_date_currency_due
            ON ar_snapshot(snapshot_date, currency, due_date);
        CREATE INDEX idx_customer_currency_date_risk
            ON customer_currency_summary(snapshot_date, currency, over90_amount, overdue_amount);
        """
    )


def import_aging(connection: sqlite3.Connection, path: Path) -> ImportResult:
    frame = pd.read_excel(path)
    snapshot = detect_aging_date(frame, path)
    raw_rows = len(frame)

    required = {"고객", "고객 이름", "전표 번호", "순만기일", "조정계정명", "총 금액", "Aging Days"}
    missing = required.difference(frame.columns)
    if missing:
        raise ValueError(f"{path.name}: 필수 열 누락 - {', '.join(sorted(missing))}")

    account_names = frame["조정계정명"].fillna("").astype(str).str.strip()
    customer_codes = frame["고객"].map(clean_customer_code)
    amounts = pd.to_numeric(frame["총 금액"], errors="coerce")
    include = (
        account_names.str.contains(TRADE_ACCOUNT_RE, na=False)
        & ~account_names.str.contains(EXCLUDED_ACCOUNT_RE, na=False)
        & customer_codes.ne("")
        & amounts.notna()
    )
    selected = frame.loc[include].copy()
    selected["_source_row"] = selected.index + 2
    selected["_customer_code"] = customer_codes.loc[include]
    selected["_amount"] = amounts.loc[include].astype(float)

    rows: list[tuple[Any, ...]] = []
    for _, row in selected.iterrows():
        due = pd.to_datetime(row.get("순만기일"), errors="coerce")
        days_overdue = None if pd.isna(due) else (snapshot - due.date()).days
        source_days_value = pd.to_numeric(pd.Series([row.get("Aging Days")]), errors="coerce").iloc[0]
        source_days = None if pd.isna(source_days_value) else float(source_days_value)
        account_name = str(row.get("조정계정명") or "").strip()
        scope = "receivable" if "매출채권" in account_name else "other"
        rows.append(
            (
                snapshot.isoformat(),
                int(row["_source_row"]),
                row["_customer_code"],
                str(row.get("고객 이름") or "").strip(),
                clean_identifier(row.get("전표 번호")),
                iso_date(row.get("송장일")),
                iso_date(row.get("계산일")),
                None if pd.isna(due) else due.date().isoformat(),
                days_overdue,
                source_days,
                clean_identifier(row.get("지급 기간")),
                str(row.get("기간 설명") or "").strip(),
                str(row.get("전표 통화") or "").strip(),
                clean_identifier(row.get("조정계정")),
                account_name,
                scope,
                float(row["_amount"]),
                due_bucket(days_overdue),
                invoice_bucket(source_days),
            )
        )

    connection.executemany(
        """
        INSERT INTO ar_snapshot (
            snapshot_date, source_row, customer_code, customer_name, document_no,
            invoice_date, posting_date, due_date, days_overdue, aging_days_source,
            payment_terms, payment_terms_desc, currency, recon_account,
            recon_account_name, scope, amount, due_bucket, invoice_bucket
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        rows,
    )

    connection.execute(
        """
        INSERT INTO monthly_summary
        SELECT
            snapshot_date,
            SUM(amount),
            SUM(CASE WHEN days_overdue > 0 THEN amount ELSE 0 END),
            SUM(CASE WHEN days_overdue > 90 THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'not_due' THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'due_today' THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'overdue_1_30' THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'overdue_31_60' THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'overdue_61_90' THEN amount ELSE 0 END),
            SUM(CASE WHEN invoice_bucket = 'invoice_1_30' THEN amount ELSE 0 END),
            SUM(CASE WHEN invoice_bucket = 'invoice_31_60' THEN amount ELSE 0 END),
            SUM(CASE WHEN invoice_bucket = 'invoice_61_90' THEN amount ELSE 0 END),
            SUM(CASE WHEN invoice_bucket = 'invoice_91_plus' THEN amount ELSE 0 END),
            COUNT(DISTINCT customer_code),
            COUNT(*)
        FROM ar_snapshot
        WHERE snapshot_date = ?
        GROUP BY snapshot_date
        """,
        (snapshot.isoformat(),),
    )

    connection.execute(
        """
        INSERT INTO customer_summary
        SELECT
            snapshot_date,
            customer_code,
            MAX(customer_name),
            CASE WHEN SUM(CASE WHEN scope='other' THEN ABS(amount) ELSE 0 END) >
                           SUM(CASE WHEN scope='receivable' THEN ABS(amount) ELSE 0 END)
                 THEN 'other' ELSE 'receivable' END,
            SUM(amount),
            SUM(CASE WHEN days_overdue > 0 THEN amount ELSE 0 END),
            SUM(CASE WHEN days_overdue > 90 THEN amount ELSE 0 END),
            COUNT(DISTINCT document_no)
        FROM ar_snapshot
        WHERE snapshot_date = ?
        GROUP BY snapshot_date, customer_code
        """,
        (snapshot.isoformat(),),
    )

    connection.execute(
        """
        INSERT INTO currency_summary
        SELECT
            snapshot_date,
            CASE WHEN TRIM(COALESCE(currency,''))='' THEN '미지정' ELSE TRIM(currency) END,
            SUM(amount),
            SUM(CASE WHEN days_overdue > 0 THEN amount ELSE 0 END),
            SUM(CASE WHEN days_overdue > 90 THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'not_due' THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'due_today' THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'overdue_1_30' THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'overdue_31_60' THEN amount ELSE 0 END),
            SUM(CASE WHEN due_bucket = 'overdue_61_90' THEN amount ELSE 0 END),
            SUM(CASE WHEN invoice_bucket = 'invoice_1_30' THEN amount ELSE 0 END),
            SUM(CASE WHEN invoice_bucket = 'invoice_31_60' THEN amount ELSE 0 END),
            SUM(CASE WHEN invoice_bucket = 'invoice_61_90' THEN amount ELSE 0 END),
            SUM(CASE WHEN invoice_bucket = 'invoice_91_plus' THEN amount ELSE 0 END),
            COUNT(DISTINCT customer_code),
            COUNT(*)
        FROM ar_snapshot
        WHERE snapshot_date = ?
        GROUP BY snapshot_date, CASE WHEN TRIM(COALESCE(currency,''))='' THEN '미지정' ELSE TRIM(currency) END
        """,
        (snapshot.isoformat(),),
    )

    connection.execute(
        """
        INSERT INTO customer_currency_summary
        SELECT
            snapshot_date,
            CASE WHEN TRIM(COALESCE(currency,''))='' THEN '미지정' ELSE TRIM(currency) END,
            customer_code,
            MAX(customer_name),
            CASE WHEN SUM(CASE WHEN scope='other' THEN ABS(amount) ELSE 0 END) >
                           SUM(CASE WHEN scope='receivable' THEN ABS(amount) ELSE 0 END)
                 THEN 'other' ELSE 'receivable' END,
            SUM(amount),
            SUM(CASE WHEN days_overdue > 0 THEN amount ELSE 0 END),
            SUM(CASE WHEN days_overdue > 90 THEN amount ELSE 0 END),
            SUM(CASE WHEN amount > 0 THEN amount ELSE 0 END),
            SUM(CASE WHEN amount < 0 THEN amount ELSE 0 END),
            COUNT(DISTINCT document_no)
        FROM ar_snapshot
        WHERE snapshot_date = ?
        GROUP BY snapshot_date,
                 CASE WHEN TRIM(COALESCE(currency,''))='' THEN '미지정' ELSE TRIM(currency) END,
                 customer_code
        """,
        (snapshot.isoformat(),),
    )

    result = ImportResult(
        snapshot_date=snapshot,
        source_type="aging",
        file_name=path.name,
        file_hash=sha256_file(path),
        raw_rows=raw_rows,
        included_rows=len(selected),
        excluded_rows=raw_rows - len(selected),
        amount_sum=float(selected["_amount"].sum()),
    )
    save_import_batch(connection, result)
    return result


def import_credit(connection: sqlite3.Connection, path: Path, snapshot: date) -> ImportResult:
    frame = pd.read_excel(path)
    required = {"고객", "고객명", "Credit control area", "여신한도", "여신사용액", "여신잔액"}
    missing = required.difference(frame.columns)
    if missing:
        raise ValueError(f"{path.name}: 필수 열 누락 - {', '.join(sorted(missing))}")

    records: dict[tuple[str, str], dict[str, Any]] = {}
    for _, row in frame.iterrows():
        customer = clean_customer_code(row.get("고객"))
        control_area = clean_identifier(row.get("Credit control area")) or "미지정"
        if not customer:
            continue
        key = (customer, control_area)
        item = records.setdefault(
            key,
            {
                "name": str(row.get("고객명") or "").strip(),
                "hold": str(row.get("보류 내역") or "").strip(),
                "limit": 0.0,
                "receivables": 0.0,
                "orders": 0.0,
                "deliveries": 0.0,
                "billing": 0.0,
                "special": 0.0,
                "used": 0.0,
                "available": 0.0,
                "currency": str(row.get("Currency") or "").strip(),
            },
        )
        item["limit"] += number(row.get("여신한도"))
        item["receivables"] += number(row.get("총채권"))
        item["orders"] += number(row.get("주문금액"))
        item["deliveries"] += number(row.get("납품금액"))
        item["billing"] += number(row.get("Billing Value"))
        item["special"] += number(row.get("Special Credit Exposure"))
        item["used"] += number(row.get("여신사용액"))
        item["available"] += number(row.get("여신잔액"))

    rows = [
        (
            snapshot.isoformat(), customer, item["name"], area, item["hold"],
            item["limit"], item["receivables"], item["orders"], item["deliveries"],
            item["billing"], item["special"], item["used"], item["available"], item["currency"],
        )
        for (customer, area), item in records.items()
    ]
    connection.executemany(
        """
        INSERT INTO credit_snapshot (
            snapshot_date, customer_code, customer_name, credit_control_area,
            hold_reason, credit_limit, total_receivables, order_amount,
            delivery_amount, billing_value, special_exposure, credit_used,
            credit_available, currency
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        rows,
    )

    result = ImportResult(
        snapshot_date=snapshot,
        source_type="credit",
        file_name=path.name,
        file_hash=sha256_file(path),
        raw_rows=len(frame),
        included_rows=len(records),
        excluded_rows=len(frame) - len(records),
        amount_sum=sum(item["limit"] for item in records.values()),
    )
    save_import_batch(connection, result)
    return result


def save_import_batch(connection: sqlite3.Connection, result: ImportResult) -> None:
    connection.execute(
        """
        INSERT INTO import_batch (
            snapshot_date, source_type, source_filename, file_hash, raw_rows,
            included_rows, excluded_rows, amount_sum, imported_at
        ) VALUES (?,?,?,?,?,?,?,?,?)
        """,
        (
            result.snapshot_date.isoformat(), result.source_type, result.file_name,
            result.file_hash, result.raw_rows, result.included_rows,
            result.excluded_rows, result.amount_sum,
            datetime.now().astimezone().isoformat(timespec="seconds"),
        ),
    )


def rows_as_dicts(cursor: sqlite3.Cursor) -> list[dict[str, Any]]:
    columns = [column[0] for column in cursor.description]
    return [dict(zip(columns, row)) for row in cursor.fetchall()]


def won_to_eok(value: float | int | None) -> float:
    return round(float(value or 0) / 100_000_000, 4)


def build_dashboard_data(connection: sqlite3.Connection) -> dict[str, Any]:
    summary_rows = rows_as_dicts(
        connection.execute("SELECT * FROM monthly_summary ORDER BY snapshot_date")
    )
    if not summary_rows:
        raise ValueError("적재된 Aging 스냅샷이 없습니다.")
    latest_date = summary_rows[-1]["snapshot_date"]
    previous_date = summary_rows[-2]["snapshot_date"] if len(summary_rows) > 1 else None

    current_customers = rows_as_dicts(
        connection.execute(
            "SELECT * FROM customer_summary WHERE snapshot_date=? ORDER BY total_amount DESC",
            (latest_date,),
        )
    )
    previous_amounts = {
        row["customer_code"]: row["total_amount"]
        for row in rows_as_dicts(
            connection.execute(
                "SELECT customer_code,total_amount FROM customer_summary WHERE snapshot_date=?",
                (previous_date,),
            )
        )
    } if previous_date else {}

    credit_date_row = connection.execute(
        "SELECT MAX(snapshot_date) FROM credit_snapshot"
    ).fetchone()
    credit_date = credit_date_row[0] if credit_date_row else None
    credit_rows = rows_as_dicts(
        connection.execute(
            """
            SELECT customer_code, MAX(customer_name) AS customer_name,
                   SUM(credit_limit) AS credit_limit,
                   SUM(credit_used) AS credit_used,
                   SUM(credit_available) AS credit_available
            FROM credit_snapshot
            WHERE snapshot_date=?
            GROUP BY customer_code
            """,
            (credit_date,),
        )
    ) if credit_date else []
    credit = {row["customer_code"]: row for row in credit_rows}

    customer_payload: list[dict[str, Any]] = []
    matched_count = 0
    unmatched_amount = 0.0
    high_util_count = 0
    high_util_amount = 0.0
    over_limit_count = 0
    over_limit_amount = 0.0
    credit_distribution = {"50% 미만": 0, "50–80%": 0, "80–100%": 0, "100% 초과": 0}

    for row in current_customers:
        total = float(row["total_amount"] or 0)
        overdue = float(row["overdue_amount"] or 0)
        over90 = float(row["over90_amount"] or 0)
        credit_row = credit.get(row["customer_code"])
        limit_value = used_value = util = None
        has_credit = credit_row is not None
        if has_credit:
            matched_count += 1
            limit_value = float(credit_row["credit_limit"] or 0)
            used_value = float(credit_row["credit_used"] or 0)
            if limit_value > 0:
                util = used_value / limit_value * 100
            elif used_value > 0:
                util = 999.0
            else:
                util = 0.0
            if util < 50:
                credit_distribution["50% 미만"] += 1
            elif util < 80:
                credit_distribution["50–80%"] += 1
            elif util <= 100:
                credit_distribution["80–100%"] += 1
            else:
                credit_distribution["100% 초과"] += 1
            if util >= 80:
                high_util_count += 1
                high_util_amount += total
            if used_value > limit_value:
                over_limit_count += 1
                over_limit_amount += total
        else:
            unmatched_amount += total

        if over90 >= 500_000_000 or overdue >= 1_000_000_000 or (util is not None and util > 100 and total >= 500_000_000):
            grade = "R1"
        elif over90 > 0 or overdue >= 200_000_000 or (util is not None and util > 100):
            grade = "R2"
        elif overdue > 0 or (util is not None and util >= 80) or (not has_credit and total >= 1_000_000_000):
            grade = "R3"
        else:
            grade = "R4"

        signals: list[str] = []
        if over90 > 0:
            signals.append("90일 초과")
        elif overdue > 0:
            signals.append("연체 발생")
        if util is not None and util > 100:
            signals.append("여신한도 초과")
        elif util is not None and util >= 80:
            signals.append("여신 고사용률")
        if not has_credit:
            signals.append("여신정보 없음")
        change = total - float(previous_amounts.get(row["customer_code"], 0))
        if abs(change) >= 500_000_000:
            signals.append("전월 대비 급변")
        if not signals:
            signals.append("정상 모니터링")

        customer_payload.append(
            {
                "id": row["customer_code"],
                "name": row["customer_name"],
                "grade": grade,
                "total": won_to_eok(total),
                "overdue": won_to_eok(overdue),
                "over90": won_to_eok(over90),
                "limit": None if limit_value is None else won_to_eok(limit_value),
                "used": None if used_value is None else won_to_eok(used_value),
                "util": None if util is None else round(util, 1),
                "credit": has_credit,
                "scope": row["scope"],
                "signal": " · ".join(signals[:3]),
                "action": "영업 확인 필요" if grade in {"R1", "R2"} else "정기 모니터링",
                "change": won_to_eok(change),
            }
        )

    grade_order = {"R1": 0, "R2": 1, "R3": 2, "R4": 3}
    customer_payload.sort(key=lambda item: (grade_order[item["grade"]], -item["over90"], -item["overdue"], -item["total"]))

    latest_summary = summary_rows[-1]
    import_rows = rows_as_dicts(
        connection.execute(
            "SELECT snapshot_date,source_type,source_filename,raw_rows,included_rows FROM import_batch ORDER BY snapshot_date,source_type"
        )
    )
    latest_aging_import = next(
        row for row in reversed(import_rows)
        if row["source_type"] == "aging" and row["snapshot_date"] == latest_date
    )
    negative = connection.execute(
        "SELECT COUNT(*), COALESCE(SUM(amount),0) FROM ar_snapshot WHERE snapshot_date=? AND amount<0",
        (latest_date,),
    ).fetchone()

    trend = [
        {
            "month": datetime.strptime(row["snapshot_date"], "%Y-%m-%d").strftime("%m.%d"),
            "date": row["snapshot_date"],
            "total": won_to_eok(row["total_amount"]),
            "overdue": won_to_eok(row["overdue_amount"]),
            "over90": won_to_eok(row["over90_amount"]),
        }
        for row in summary_rows
    ]
    aging = {
        "due": [
            {"label": "만기 미도래", "value": won_to_eok(latest_summary["not_due_amount"]), "risk": False},
            {"label": "당일 만기", "value": won_to_eok(latest_summary["due_today_amount"]), "risk": False},
            {"label": "1–30일 만기경과", "value": won_to_eok(latest_summary["overdue_1_30"]), "risk": True},
            {"label": "31–60일 만기경과", "value": won_to_eok(latest_summary["overdue_31_60"]), "risk": True},
            {"label": "61–90일 만기경과", "value": won_to_eok(latest_summary["overdue_61_90"]), "risk": True},
            {"label": "90일 초과 만기경과", "value": won_to_eok(latest_summary["over90_amount"]), "risk": True},
        ],
        "invoice": [
            {"label": "송장 1–30일", "value": won_to_eok(latest_summary["invoice_1_30"]), "risk": False},
            {"label": "송장 31–60일", "value": won_to_eok(latest_summary["invoice_31_60"]), "risk": False},
            {"label": "송장 61–90일", "value": won_to_eok(latest_summary["invoice_61_90"]), "risk": False},
            {"label": "송장 91일 이상", "value": won_to_eok(latest_summary["invoice_91_plus"]), "risk": True},
        ],
    }

    all_customer_rows = rows_as_dicts(
        connection.execute(
            "SELECT * FROM customer_summary ORDER BY snapshot_date, customer_code"
        )
    )
    customers_by_date: dict[str, list[dict[str, Any]]] = {}
    for row in all_customer_rows:
        customers_by_date.setdefault(row["snapshot_date"], []).append(row)
    imports_by_date = {
        row["snapshot_date"]: row
        for row in import_rows
        if row["source_type"] == "aging"
    }

    periods: list[dict[str, Any]] = []
    prior_customer_map: dict[str, dict[str, Any]] = {}
    for period_summary in summary_rows:
        period_date = period_summary["snapshot_date"]
        period_rows = customers_by_date.get(period_date, [])
        current_map = {row["customer_code"]: row for row in period_rows}
        period_customers: list[dict[str, Any]] = []

        if period_date == latest_date:
            period_customers = customer_payload
        else:
            for row in period_rows:
                total = float(row["total_amount"] or 0)
                overdue = float(row["overdue_amount"] or 0)
                over90 = float(row["over90_amount"] or 0)
                prior = prior_customer_map.get(row["customer_code"])
                prior_total = float(prior["total_amount"] or 0) if prior else 0.0
                change = total - prior_total
                if over90 >= 500_000_000 or overdue >= 1_000_000_000:
                    grade = "R1"
                elif over90 > 0 or overdue >= 200_000_000:
                    grade = "R2"
                elif overdue > 0 or total >= 1_000_000_000:
                    grade = "R3"
                else:
                    grade = "R4"
                signals: list[str] = []
                if over90 > 0:
                    signals.append("90일 초과")
                elif overdue > 0:
                    signals.append("연체 발생")
                if abs(change) >= 500_000_000:
                    signals.append("직전 대비 급변")
                if not signals:
                    signals.append("정상 모니터링")
                period_customers.append(
                    {
                        "id": row["customer_code"],
                        "name": row["customer_name"],
                        "grade": grade,
                        "total": won_to_eok(total),
                        "overdue": won_to_eok(overdue),
                        "over90": won_to_eok(over90),
                        "limit": None,
                        "used": None,
                        "util": None,
                        "credit": False,
                        "scope": row["scope"],
                        "signal": " · ".join(signals),
                        "action": "재무 모니터링",
                        "change": won_to_eok(change),
                    }
                )
            period_customers.sort(
                key=lambda item: (
                    grade_order[item["grade"]], -item["over90"],
                    -item["overdue"], -item["total"]
                )
            )

        movements: list[dict[str, Any]] = []
        for code in set(prior_customer_map) | set(current_map):
            before = prior_customer_map.get(code)
            after = current_map.get(code)
            before_total = float(before["total_amount"] or 0) if before else 0.0
            after_total = float(after["total_amount"] or 0) if after else 0.0
            before_overdue = float(before["overdue_amount"] or 0) if before else 0.0
            after_overdue = float(after["overdue_amount"] or 0) if after else 0.0
            before_over90 = float(before["over90_amount"] or 0) if before else 0.0
            after_over90 = float(after["over90_amount"] or 0) if after else 0.0
            overdue_change = after_overdue - before_overdue
            total_change = after_total - before_total
            movement_type = None
            if before_overdue <= 0 < after_overdue:
                movement_type = "SAP 만기경과 진입"
            elif before_overdue > 0 >= after_overdue:
                movement_type = "SAP 만기경과 해소"
            elif overdue_change >= 50_000_000:
                movement_type = "SAP 만기경과 증가"
            elif overdue_change <= -50_000_000:
                movement_type = "SAP 만기경과 감소"
            elif total_change >= 500_000_000:
                movement_type = "채권 증가"
            elif total_change <= -500_000_000:
                movement_type = "채권 감소"
            if movement_type:
                source = after or before
                movements.append(
                    {
                        "type": movement_type,
                        "id": code,
                        "name": source["customer_name"],
                        "beforeTotal": won_to_eok(before_total),
                        "total": won_to_eok(after_total),
                        "totalChange": won_to_eok(total_change),
                        "beforeOverdue": won_to_eok(before_overdue),
                        "overdue": won_to_eok(after_overdue),
                        "overdueChange": won_to_eok(overdue_change),
                        "beforeOver90": won_to_eok(before_over90),
                        "over90": won_to_eok(after_over90),
                    }
                )
        movements.sort(
            key=lambda item: (
                0 if item["type"] in {"SAP 만기경과 진입", "SAP 만기경과 증가"} else 1,
                -abs(item["overdueChange"]), -abs(item["totalChange"]),
            )
        )

        period_import = imports_by_date[period_date]
        quality_row = connection.execute(
            """
            SELECT
                SUM(CASE WHEN due_date IS NULL THEN 1 ELSE 0 END),
                SUM(CASE WHEN TRIM(customer_name)='' THEN 1 ELSE 0 END),
                COUNT(*),
                SUM(amount)
            FROM ar_snapshot WHERE snapshot_date=?
            """,
            (period_date,),
        ).fetchone()
        duplicate_count = connection.execute(
            """
            SELECT COALESCE(SUM(cnt-1),0) FROM (
                SELECT COUNT(*) AS cnt
                FROM ar_snapshot
                WHERE snapshot_date=?
                GROUP BY customer_code, document_no, due_date, amount
                HAVING COUNT(*) > 1
            )
            """,
            (period_date,),
        ).fetchone()[0]
        is_latest = period_date == latest_date
        period_summary_payload = {
            "total": won_to_eok(period_summary["total_amount"]),
            "overdue": won_to_eok(period_summary["overdue_amount"]),
            "over90": won_to_eok(period_summary["over90_amount"]),
            "customerCount": int(period_summary["customer_count"]),
            "highUtilCount": high_util_count if is_latest else 0,
            "highUtilAmount": won_to_eok(high_util_amount) if is_latest else 0,
            "overLimitCount": over_limit_count if is_latest else 0,
            "overLimitAmount": won_to_eok(over_limit_amount) if is_latest else 0,
            "matchedCount": matched_count if is_latest else 0,
            "unmatchedCount": len(period_rows) - matched_count if is_latest else 0,
            "unmatchedAmount": won_to_eok(unmatched_amount) if is_latest else 0,
            "rawRows": int(period_import["raw_rows"]),
            "includedRows": int(period_import["included_rows"]),
            "excludedRows": int(period_import["raw_rows"] - period_import["included_rows"]),
            "negativeRows": int(connection.execute(
                "SELECT COUNT(*) FROM ar_snapshot WHERE snapshot_date=? AND amount<0",
                (period_date,),
            ).fetchone()[0]),
            "negativeAmount": won_to_eok(connection.execute(
                "SELECT COALESCE(SUM(amount),0) FROM ar_snapshot WHERE snapshot_date=? AND amount<0",
                (period_date,),
            ).fetchone()[0]),
            "creditAvailable": is_latest and bool(credit_rows),
        }
        period_aging = {
            "due": [
                {"label": "만기 미도래", "value": won_to_eok(period_summary["not_due_amount"]), "risk": False},
                {"label": "당일 만기", "value": won_to_eok(period_summary["due_today_amount"]), "risk": False},
                {"label": "1–30일 만기경과", "value": won_to_eok(period_summary["overdue_1_30"]), "risk": True},
                {"label": "31–60일 만기경과", "value": won_to_eok(period_summary["overdue_31_60"]), "risk": True},
                {"label": "61–90일 만기경과", "value": won_to_eok(period_summary["overdue_61_90"]), "risk": True},
                {"label": "90일 초과 만기경과", "value": won_to_eok(period_summary["over90_amount"]), "risk": True},
            ],
            "invoice": [
                {"label": "송장 1–30일", "value": won_to_eok(period_summary["invoice_1_30"]), "risk": False},
                {"label": "송장 31–60일", "value": won_to_eok(period_summary["invoice_31_60"]), "risk": False},
                {"label": "송장 61–90일", "value": won_to_eok(period_summary["invoice_61_90"]), "risk": False},
                {"label": "송장 91일 이상", "value": won_to_eok(period_summary["invoice_91_plus"]), "risk": True},
            ],
        }
        periods.append(
            {
                "date": period_date,
                "label": datetime.strptime(period_date, "%Y-%m-%d").strftime("%Y.%m.%d"),
                "summary": period_summary_payload,
                "aging": period_aging,
                "customers": period_customers,
                "movements": movements[:200],
                "quality": {
                    "missingDueDate": int(quality_row[0] or 0),
                    "missingCustomerName": int(quality_row[1] or 0),
                    "duplicateSuspects": int(duplicate_count or 0),
                    "rowCount": int(quality_row[2] or 0),
                    "amountDifference": won_to_eok(float(quality_row[3] or 0) - float(period_summary["total_amount"] or 0)),
                },
            }
        )
        prior_customer_map = current_map

    return {
        "meta": {
            "asOf": latest_date,
            "generatedAt": datetime.now().astimezone().strftime("%Y-%m-%d %H:%M"),
            "readonly": True,
        },
        "summary": {
            "total": won_to_eok(latest_summary["total_amount"]),
            "overdue": won_to_eok(latest_summary["overdue_amount"]),
            "over90": won_to_eok(latest_summary["over90_amount"]),
            "customerCount": int(latest_summary["customer_count"]),
            "highUtilCount": high_util_count,
            "highUtilAmount": won_to_eok(high_util_amount),
            "overLimitCount": over_limit_count,
            "overLimitAmount": won_to_eok(over_limit_amount),
            "matchedCount": matched_count,
            "unmatchedCount": len(current_customers) - matched_count,
            "unmatchedAmount": won_to_eok(unmatched_amount),
            "rawRows": int(latest_aging_import["raw_rows"]),
            "includedRows": int(latest_aging_import["included_rows"]),
            "excludedRows": int(latest_aging_import["raw_rows"] - latest_aging_import["included_rows"]),
            "negativeRows": int(negative[0]),
            "negativeAmount": won_to_eok(negative[1]),
        },
        "customers": customer_payload,
        "aging": aging,
        "trend": trend,
        "creditDistribution": [
            {"label": label, "value": value, "risk": label in {"80–100%", "100% 초과"}}
            for label, value in credit_distribution.items()
        ],
        "snapshots": [
            {"date": row["snapshot_date"], "label": "현재 Aging" if row["snapshot_date"] == latest_date else "과거 Aging"}
            for row in summary_rows
        ],
        "imports": [
            {
                "date": row["snapshot_date"],
                "type": row["source_type"],
                "file": row["source_filename"],
                "rawRows": int(row["raw_rows"]),
                "includedRows": int(row["included_rows"]),
            }
            for row in import_rows
        ],
        "periods": periods,
        "movements": periods[-1]["movements"],
        "quality": periods[-1]["quality"],
    }


def display_amount(value: float | int | None, currency: str) -> float:
    divisor = 100_000_000 if currency == "KRW" else 1_000_000
    return round(float(value or 0) / divisor, 4)


def clamp(value: float, lower: float = 0.0, upper: float = 1.0) -> float:
    return min(upper, max(lower, value))


def normalized_repayment_rate(repaid_amount: float, base_amount: float, elapsed_days: int) -> float:
    """Convert an observed Aging-based repayment rate to a comparable 30-day rate."""
    if base_amount <= 0 or elapsed_days <= 0:
        return 0.0
    observed = clamp(repaid_amount / base_amount)
    return clamp(1 - (1 - observed) ** (30 / elapsed_days))


def net_open_documents(document_map: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Return customer-level debit/credit-netted documents for risk analysis.

    The current customer balance is authoritative for collection risk.  Credit
    balances therefore extinguish debit documents even when SAP has not linked
    the clearing items at line-item level.  Credits are applied to the oldest
    debit documents first so duration, Aging amounts, recovery and transition
    metrics all use the same post-offset basis.  Signed source documents remain
    unchanged for accounting/allowance calculations.
    """
    by_customer: dict[str, list[tuple[str, dict[str, Any]]]] = {}
    for key, item in document_map.items():
        by_customer.setdefault(str(item.get("customer") or ""), []).append((key, item))

    netted: dict[str, dict[str, Any]] = {}
    for items in by_customer.values():
        credit_pool = -sum(min(0.0, float(item.get("amount") or 0)) for _, item in items)
        debit_items = [
            (key, item)
            for key, item in items
            if float(item.get("amount") or 0) > 0
        ]
        debit_items.sort(
            key=lambda pair: (
                -(int(pair[1]["invoiceAge"]) if pair[1].get("invoiceAge") is not None else -1),
                pair[1].get("invoiceDate") or "",
                pair[0],
            )
        )
        for key, item in debit_items:
            source_amount = float(item.get("amount") or 0)
            offset = min(source_amount, credit_pool)
            credit_pool -= offset
            residual = source_amount - offset
            if residual <= 1e-9:
                continue
            net_item = dict(item)
            net_item["amount"] = residual
            net_item["preOffsetAmount"] = source_amount
            netted[key] = net_item
    return netted


def piecewise_score(value: float, anchors: list[tuple[float, float]]) -> float:
    """Interpolate an explainable 0-100 score between fixed business anchors."""
    if not anchors:
        return 0.0
    numeric = max(0.0, float(value or 0))
    if numeric <= anchors[0][0]:
        return float(anchors[0][1])
    for (left_value, left_score), (right_value, right_score) in zip(anchors, anchors[1:]):
        if numeric <= right_value:
            width = right_value - left_value
            ratio = (numeric - left_value) / width if width else 0.0
            return left_score + (right_score - left_score) * ratio
    return float(anchors[-1][1])


def weighted_mean(values: list[tuple[float, float]], default: float = 0.0) -> float:
    total_weight = sum(max(0.0, weight) for _, weight in values)
    if total_weight <= 0:
        return default
    return sum(value * max(0.0, weight) for value, weight in values) / total_weight


def weighted_median(values: list[tuple[float, float]], default: float = 0.0) -> float:
    weighted = sorted((value, max(0.0, weight)) for value, weight in values if weight > 0)
    total_weight = sum(weight for _, weight in weighted)
    if total_weight <= 0:
        return default
    halfway = total_weight / 2
    cumulative = 0.0
    for value, weight in weighted:
        cumulative += weight
        if cumulative >= halfway:
            return value
    return weighted[-1][0]


def duration_risk_score(max_open_days: int | float) -> float:
    """Score the balance-weighted average invoice age of current open items."""
    days = max(0, float(max_open_days or 0))
    if days <= 90:
        return 0.0
    if days <= 180:
        return 25.0
    if days <= 270:
        return 50.0
    if days <= 360:
        return 75.0
    return 100.0


def absolute_recovery_weakness_score(current_speed_pct: float) -> float:
    """Score the absolute weakness of the recent 90-day collection rate.

    A persistently stalled customer must not receive zero risk merely because
    its historical collection rate was equally weak.  Seventy-five percent or
    more is treated as healthy for this factor; zero collection receives 100.
    """
    return piecewise_score(
        current_speed_pct,
        [(0.0, 100.0), (10.0, 85.0), (25.0, 60.0), (50.0, 25.0), (75.0, 0.0)],
    )


def credit_pressure_profile(
    credit_limit_eok: float | None,
    credit_used_eok: float | None,
    credit_available: bool,
) -> dict[str, Any]:
    """Return a materiality-adjusted credit exposure pressure score.

    Credit is an additive corroborating signal, not a replacement for observed
    collection behaviour.  A high utilisation percentage on an immaterial or
    zero-limit residual must therefore be moderated by the absolute amount used.
    """
    if not credit_available or credit_limit_eok is None or credit_used_eok is None:
        return {
            "available": False,
            "utilization": None,
            "rawScore": 0.0,
            "materiality": 0.0,
            "score": 0.0,
            "adjustment": 0.0,
        }
    limit_value = max(0.0, float(credit_limit_eok or 0))
    used_value = max(0.0, float(credit_used_eok or 0))
    utilization = (
        used_value / limit_value * 100
        if limit_value > 0
        else (999.0 if used_value > 0 else 0.0)
    )
    raw_score = piecewise_score(
        utilization,
        [(0.0, 0.0), (70.0, 0.0), (80.0, 20.0), (90.0, 50.0),
         (100.0, 80.0), (110.0, 100.0)],
    )
    # One hundred million KRW of used credit gives this corroborating factor
    # full credibility.  Smaller exposures scale linearly to prevent a tiny
    # zero-limit balance from dominating the customer grade.
    materiality = clamp(used_value / 1.0)
    pressure_score = raw_score * materiality
    return {
        "available": True,
        "utilization": round(utilization, 1),
        "rawScore": round(raw_score, 1),
        "materiality": round(materiality * 100, 1),
        "score": round(pressure_score, 1),
        "adjustment": round(pressure_score * 0.10, 1),
    }


def receivable_risk_profile(
    aged_amount_eok: float,
    positive_open_eok: float,
    average_open_days: int | float,
    current_speed_pct: float,
    baseline_speed_pct: float,
    transition_rate_pct: float,
    transition_eligible_eok: float,
    speed_available: bool,
    transition_available: bool,
    transition_status: str | None = None,
    credit_limit_eok: float | None = None,
    credit_used_eok: float | None = None,
    credit_available: bool = False,
) -> dict[str, Any]:
    """Return the five-factor collection score plus an auditable credit adjustment.

    Amount inputs are KRW-equivalent 100M units.  The score deliberately uses
    fixed anchors instead of peer percentiles, so the same exposure produces the
    same score regardless of which customers happen to be in a snapshot.
    """
    aged_amount = max(0.0, float(aged_amount_eok or 0))
    positive_open = max(0.0, float(positive_open_eok or 0))
    aged_share = aged_amount / positive_open * 100 if positive_open else 0.0
    speed_decline = max(0.0, float(baseline_speed_pct or 0) - float(current_speed_pct or 0))

    amount_score = piecewise_score(
        aged_amount,
        [(0.0, 0.0), (0.1, 5.0), (0.5, 15.0), (1.0, 30.0), (2.0, 45.0),
         (5.0, 70.0), (10.0, 90.0), (20.0, 100.0)],
    )
    share_raw_score = piecewise_score(
        aged_share,
        [(0.0, 0.0), (15.0, 20.0), (30.0, 40.0), (45.0, 60.0), (60.0, 80.0), (75.0, 100.0)],
    )
    # A 100% share on a trivial residual balance must not receive the same
    # structural-risk score as a material aged exposure.  The existing 0.5 EOK
    # amount anchor is used as the full-credibility threshold.
    share_credibility = clamp(aged_amount / 0.5)
    share_score = share_raw_score * share_credibility
    duration_score_basis_days = max(0.0, float(average_open_days or 0))
    duration_score = duration_risk_score(duration_score_basis_days)
    speed_decline_score = (
        piecewise_score(
            speed_decline,
            [(0.0, 0.0), (5.0, 0.0), (10.0, 20.0), (15.0, 40.0), (25.0, 70.0), (35.0, 90.0), (40.0, 100.0)],
        )
        if speed_available else 0.0
    )
    absolute_weakness_score = (
        absolute_recovery_weakness_score(current_speed_pct)
        if speed_available else 0.0
    )
    # Validation showed that decline-only scoring misses customers whose
    # collection rate has remained at zero.  Absolute weakness therefore carries
    # 60% of this factor and deterioration versus baseline carries 40%.
    speed_score = (
        absolute_weakness_score * 0.60 + speed_decline_score * 0.40
        if speed_available else 0.0
    )
    resolved_transition_status = transition_status or (
        "available" if transition_available else "insufficient_history"
    )
    if transition_available:
        resolved_transition_status = "available"
    if resolved_transition_status not in {"available", "not_applicable", "insufficient_history"}:
        resolved_transition_status = "insufficient_history"
    transition_raw_score = (
        piecewise_score(
            transition_rate_pct,
            [(0.0, 0.0), (10.0, 0.0), (25.0, 25.0), (50.0, 50.0), (75.0, 75.0), (100.0, 100.0)],
        )
        if resolved_transition_status == "available" else 0.0
    )
    transition_eligible = max(0.0, float(transition_eligible_eok or 0))
    transition_credibility = (
        clamp(transition_eligible / 0.5)
        if resolved_transition_status == "available" else 0.0
    )
    transition_score = transition_raw_score * transition_credibility
    weights = {"amount": 40, "share": 10, "duration": 25, "speed": 15, "transition": 10}
    scores = {
        "amount": amount_score,
        "share": share_score,
        "duration": duration_score,
        "speed": speed_score,
        "transition": transition_score,
    }
    contributions = {
        key: round(scores[key] * weights[key] / 100, 1)
        for key in weights
    }
    collection_risk_score = round(
        sum(scores[key] * weights[key] / 100 for key in weights),
        1,
    )
    credit_pressure = credit_pressure_profile(
        credit_limit_eok,
        credit_used_eok,
        credit_available,
    )
    risk_score = round(
        min(100.0, collection_risk_score + float(credit_pressure["adjustment"])),
        1,
    )
    if risk_score >= 75:
        grade, risk_type = "R1", "최고위험"
    elif risk_score >= 60:
        grade, risk_type = "R2", "고위험"
    elif risk_score >= 30:
        grade, risk_type = "R3", "주의"
    else:
        grade, risk_type = "R4", "낮음"
    coverage = 75 + (15 if speed_available else 0) + (
        10 if resolved_transition_status in {"available", "not_applicable"} else 0
    )
    return {
        "riskScore": risk_score,
        "collectionRiskScore": collection_risk_score,
        "grade": grade,
        "riskType": risk_type,
        "amountScore": round(amount_score, 1),
        "shareScore": round(share_score, 1),
        "durationScore": round(duration_score, 1),
        "durationScoreBasisDays": round(duration_score_basis_days, 6),
        "speedScore": round(speed_score, 1),
        "recoveryScore": round(speed_score, 1),
        "absoluteRecoveryWeaknessScore": round(absolute_weakness_score, 1),
        "speedDeclineScore": round(speed_decline_score, 1),
        "transitionScore": round(transition_score, 1),
        "creditPressureScore": credit_pressure["score"],
        "creditUsageRawScore": credit_pressure["rawScore"],
        "creditMateriality": credit_pressure["materiality"],
        "creditAdjustment": credit_pressure["adjustment"],
        "creditAvailable": credit_pressure["available"],
        "creditUtilization": credit_pressure["utilization"],
        # Backward-compatible aliases for older static consumers.
        "frequencyScore": round(share_score, 1),
        "trendScore": round(transition_score, 1),
        "weights": weights,
        "contributions": contributions,
        "agedAmount": round(aged_amount, 4),
        "agedShare": round(aged_share, 1),
        "averageOpenDays": round(duration_score_basis_days, 1),
        "speedDecline": round(speed_decline, 1),
        "transitionRate": round(max(0.0, float(transition_rate_pct or 0)), 1),
        "shareCredibility": round(share_credibility * 100, 1),
        "transitionCredibility": round(transition_credibility * 100, 1),
        "speedAvailable": bool(speed_available),
        "transitionAvailable": resolved_transition_status == "available",
        "transitionStatus": resolved_transition_status,
        "coverage": coverage,
        "provisional": coverage < 80,
    }


def estimated_repayment_components(
    prior_documents: dict[str, dict[str, Any]],
    current_documents: dict[str, dict[str, Any]],
) -> tuple[float, float, float, float]:
    """Estimate full and partial repayment from consecutive Aging reports."""
    prior_positive = {
        key: item for key, item in prior_documents.items() if float(item.get("amount") or 0) > 0
    }
    current_positive = {
        key: item for key, item in current_documents.items() if float(item.get("amount") or 0) > 0
    }
    common_keys = set(prior_positive) & set(current_positive)
    base_amount = sum(float(item["amount"]) for item in prior_positive.values())
    disappeared_amount = sum(
        float(item["amount"])
        for key, item in prior_positive.items()
        if key not in current_positive
    )
    partial_repayment = sum(
        max(0.0, float(prior_positive[key]["amount"]) - float(current_positive[key]["amount"]))
        for key in common_keys
    )
    return base_amount, disappeared_amount + partial_repayment, disappeared_amount, partial_repayment


def invoice_allowance_bucket(invoice_age: int | None) -> str | None:
    if invoice_age is None:
        return None
    if invoice_age <= 90:
        return "age_0_90"
    if invoice_age <= 180:
        return "age_91_180"
    if invoice_age <= 270:
        return "age_181_270"
    if invoice_age <= 360:
        return "age_271_360"
    return "age_361_plus"


def repayment_age_bucket(invoice_age: int | None) -> str:
    if invoice_age is None:
        return "unclassified"
    if invoice_age <= 30:
        return "age_0_30"
    if invoice_age <= 60:
        return "age_31_60"
    if invoice_age <= 90:
        return "age_61_90"
    if invoice_age <= 180:
        return "age_91_180"
    if invoice_age <= 360:
        return "age_181_360"
    return "age_361_plus"


OPEN_AGE_BUCKETS = (
    ("age_0_90", "90일 이하", 0, 90),
    ("age_91_180", "91~180일", 91, 180),
    ("age_181_270", "181~270일", 181, 270),
    ("age_271_360", "271~360일", 271, 360),
    ("age_361_plus", "361일 이상", 361, None),
)


def open_age_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Return an auditable amount-weighted invoice-age summary for open items."""
    positive = [
        row for row in rows
        if float(row.get("amount") or 0) > 0
    ]
    total = sum(float(row.get("amount") or 0) for row in positive)
    numerator = sum(
        float(row.get("amount") or 0) * max(0, int(row.get("invoiceAge") or 0))
        for row in positive
    )
    average_days = numerator / total if total > 0 else 0.0
    buckets: list[dict[str, Any]] = []
    for key, label, minimum, maximum in OPEN_AGE_BUCKETS:
        bucket_rows = [
            row for row in positive
            if max(0, int(row.get("invoiceAge") or 0)) >= minimum
            and (maximum is None or max(0, int(row.get("invoiceAge") or 0)) <= maximum)
        ]
        amount = sum(float(row.get("amount") or 0) for row in bucket_rows)
        buckets.append(
            {
                "key": key,
                "label": label,
                "amount": amount,
                "share": amount / total * 100 if total > 0 else 0.0,
                "documentCount": len(bucket_rows),
            }
        )
    contributors = sorted(
        (
            {
                "id": str(row.get("id") or ""),
                "invoiceDate": row.get("invoiceDate"),
                "invoiceAge": max(0, int(row.get("invoiceAge") or 0)),
                "amount": float(row.get("amount") or 0),
                "sourceCurrency": row.get("sourceCurrency"),
                "contributionDays": (
                    float(row.get("amount") or 0)
                    / total
                    * max(0, int(row.get("invoiceAge") or 0))
                    if total > 0 else 0.0
                ),
            }
            for row in positive
        ),
        key=lambda row: (row["contributionDays"], row["amount"]),
        reverse=True,
    )[:10]
    return {
        "averageDays": average_days,
        "totalAmount": total,
        "documentCount": len(positive),
        "buckets": buckets,
        "contributors": contributors,
    }


def net_open_age_basis(
    rows: list[dict[str, Any]],
    cross_currency_offset: float,
    net_open_amount: float,
) -> list[dict[str, Any]]:
    """Apply customer-level currency offsets without losing item precision.

    Combined KRW-equivalent reporting uses 억원 as its display unit.  Rounding
    each invoice to four decimals before aggregation can therefore discard
    material amounts for customers with many open items.  Keep full precision
    here and round only the final customer/bucket outputs.
    """
    if net_open_amount <= 0:
        return []
    remaining_offset = max(0.0, float(cross_currency_offset or 0))
    netted: list[dict[str, Any]] = []
    for basis in sorted(
        rows,
        key=lambda item: -(int(item.get("invoiceAge") or 0)),
    ):
        basis_amount = max(0.0, float(basis.get("amount") or 0))
        basis_offset = min(basis_amount, remaining_offset)
        remaining_offset -= basis_offset
        residual_amount = basis_amount - basis_offset
        if residual_amount <= 1e-12:
            continue
        net_basis = dict(basis)
        net_basis["amount"] = residual_amount
        netted.append(net_basis)
    return netted


def enrich_native_periods(
    periods: list[dict[str, Any]],
    doc_maps: dict[str, dict[str, dict[str, Any]]],
    dates: list[str],
    currency: str,
    allowance_model: dict[str, Any],
    usd_krw_rate: float,
) -> None:
    """Enrich periods with cadence-independent document lifecycle risk metrics.

    A snapshot is an observation only.  Repeated appearances of the same SAP
    document never become multiple risk events; each distinct document counts
    once, and duration is measured in calendar days.
    """
    histories: dict[str, list[dict[str, Any]]] = {}
    prior_profiles: dict[str, dict[str, Any]] = {}
    portfolio_speed_history: list[tuple[float, float, int]] = []
    parsed_dates = {value: datetime.strptime(value, "%Y-%m-%d").date() for value in dates}
    risk_doc_maps = {
        observed_date: net_open_documents(doc_maps.get(observed_date, {}))
        for observed_date in dates
    }
    rates = {bucket["key"]: float(bucket["appliedRate"]) / 100 for bucket in allowance_model["buckets"]}
    raw_rates = {bucket["key"]: float(bucket["rawRate"]) for bucket in allowance_model["buckets"]}
    labels = {bucket["key"]: bucket["label"] for bucket in allowance_model["buckets"]}
    methods = {
        bucket["key"]: bucket.get("method", "개별평가" if bucket["key"] == "age_361_plus" else "집합평가")
        for bucket in allowance_model["buckets"]
    }
    allowance_excluded_codes = set(allowance_model.get("excludedCustomerCodes") or [])
    grade_order = {"R1": 0, "R2": 1, "R3": 2, "R4": 3}

    def lifecycle_for(document_key: str, date_index: int) -> dict[str, Any]:
        observed_dates = [
            value
            for value in dates[: date_index + 1]
            if document_key in risk_doc_maps.get(value, {})
            and float(risk_doc_maps[value][document_key].get("amount") or 0) > 0
        ]
        flags = [
            document_key in risk_doc_maps.get(value, {})
            and float(risk_doc_maps[value][document_key].get("amount") or 0) > 0
            for value in dates[: date_index + 1]
        ]
        run_count = sum(flag and (index == 0 or not flags[index - 1]) for index, flag in enumerate(flags))
        continuous_start_index = date_index
        while continuous_start_index > 0 and flags[continuous_start_index - 1]:
            continuous_start_index -= 1
        first_seen = observed_dates[0]
        continuous_start = dates[continuous_start_index]
        return {
            "firstSeenDate": first_seen,
            "lastSeenDate": observed_dates[-1],
            "continuousStartDate": continuous_start,
            "observationCount": len(observed_dates),
            "observedOpenDays": (parsed_dates[dates[date_index]] - parsed_dates[continuous_start]).days,
            "reappearanceCount": max(0, run_count - 1),
            "firstSeenLowerBound": first_seen == dates[0],
        }

    def closed_lifecycle_for(document_key: str, date_index: int) -> dict[str, Any]:
        flags = [
            document_key in risk_doc_maps.get(value, {})
            and float(risk_doc_maps[value][document_key].get("amount") or 0) > 0
            for value in dates[: date_index + 1]
        ]
        observed_indices = [index for index, flag in enumerate(flags) if flag]
        first_index = observed_indices[0]
        last_index = observed_indices[-1]
        continuous_start_index = last_index
        while continuous_start_index > 0 and flags[continuous_start_index - 1]:
            continuous_start_index -= 1
        run_count = sum(flag and (index == 0 or not flags[index - 1]) for index, flag in enumerate(flags))
        closure_index = min(last_index + 1, date_index)
        return {
            "firstSeenDate": dates[first_index],
            "lastSeenDate": dates[last_index],
            "continuousStartDate": dates[continuous_start_index],
            "closureConfirmedDate": dates[closure_index],
            "observationCount": len(observed_indices),
            "observedOpenDays": (parsed_dates[dates[last_index]] - parsed_dates[dates[continuous_start_index]]).days,
            "reappearanceCount": max(0, run_count - 1),
            "firstSeenLowerBound": first_index == 0,
        }

    def aging_transition_for(customer_code: str, date_index: int) -> tuple[float, float, int, str]:
        """Current residual roll-rate from invoice-age 61-90 into 91+.

        Each SAP document enters the denominator once.  Snapshot frequency never
        becomes event frequency.  The numerator is the balance of the same
        document that remains open above 90 days at the current as-of date, so a
        document cleared after crossing the boundary no longer raises current risk.
        """
        as_of = parsed_dates[dates[date_index]]
        lookback = as_of - timedelta(days=183)
        candidate_keys: set[str] = set()
        customer_observed_dates: list[date] = []
        for observed_date in dates[: date_index + 1]:
            if parsed_dates[observed_date] < lookback - timedelta(days=45):
                continue
            customer_seen = False
            for key, item in risk_doc_maps.get(observed_date, {}).items():
                if item.get("customer") == customer_code and float(item.get("amount") or 0) > 0:
                    candidate_keys.add(key)
                    customer_seen = True
            if customer_seen:
                customer_observed_dates.append(parsed_dates[observed_date])

        eligible_amount = transitioned_amount = 0.0
        eligible_documents = 0
        for key in candidate_keys:
            observations: list[tuple[int, dict[str, Any]]] = []
            for index, observed_date in enumerate(dates[: date_index + 1]):
                item = risk_doc_maps.get(observed_date, {}).get(key)
                if item and float(item.get("amount") or 0) > 0:
                    observations.append((index, item))
            pre_candidates = [
                (index, item)
                for index, item in observations
                if parsed_dates[dates[index]] >= lookback
                and item.get("invoiceAge") is not None
                and 61 <= int(item["invoiceAge"]) <= 90
            ]
            if not pre_candidates:
                continue
            pre_index, pre_item = pre_candidates[-1]
            pre_age = int(pre_item["invoiceAge"])
            boundary = parsed_dates[dates[pre_index]] + timedelta(days=max(1, 91 - pre_age))
            post_index = next(
                (
                    index
                    for index in range(pre_index + 1, date_index + 1)
                    if parsed_dates[dates[index]] >= boundary
                ),
                None,
            )
            if post_index is None:
                continue
            denominator = max(0.0, float(pre_item.get("amount") or 0))
            if denominator <= 0:
                continue
            current_item = risk_doc_maps.get(dates[date_index], {}).get(key)
            numerator = (
                min(denominator, max(0.0, float(current_item.get("amount") or 0)))
                if current_item
                and current_item.get("invoiceAge") is not None
                and int(current_item["invoiceAge"]) > 90
                else 0.0
            )
            eligible_amount += denominator
            transitioned_amount += numerator
            eligible_documents += 1
        if eligible_amount > 0:
            status = "available"
        elif customer_observed_dates and (as_of - min(customer_observed_dates)).days >= 90:
            status = "not_applicable"
        else:
            status = "insufficient_history"
        return eligible_amount, transitioned_amount, eligible_documents, status

    for date_index, (period_date, period) in enumerate(zip(dates, periods)):
        signed_current_docs = doc_maps.get(period_date, {})
        current_docs = risk_doc_maps.get(period_date, {})
        prior_docs = risk_doc_maps.get(dates[date_index - 1], {}) if date_index else {}
        elapsed_days = (
            max(1, (parsed_dates[period_date] - parsed_dates[dates[date_index - 1]]).days)
            if date_index
            else 0
        )
        current_by_customer: dict[str, dict[str, dict[str, Any]]] = {}
        signed_current_by_customer: dict[str, dict[str, dict[str, Any]]] = {}
        prior_by_customer: dict[str, dict[str, dict[str, Any]]] = {}
        for key, item in signed_current_docs.items():
            signed_current_by_customer.setdefault(item["customer"], {})[key] = item
        for key, item in current_docs.items():
            current_by_customer.setdefault(item["customer"], {})[key] = item
        for key, item in prior_docs.items():
            prior_by_customer.setdefault(item["customer"], {})[key] = item
        known_by_customer: dict[str, set[str]] = {}
        if date_index == len(dates) - 1:
            for observed_date in dates:
                for key, item in risk_doc_maps.get(observed_date, {}).items():
                    if float(item.get("amount") or 0) > 0:
                        known_by_customer.setdefault(item["customer"], set()).add(key)

        raw_profiles: dict[str, dict[str, Any]] = {}
        for customer in period["customers"]:
            code = customer["id"]
            current_positive = {
                key: item
                for key, item in current_by_customer.get(code, {}).items()
                if float(item.get("amount") or 0) > 0
            }
            prior_positive = {
                key: item
                for key, item in prior_by_customer.get(code, {}).items()
                if float(item.get("amount") or 0) > 0
            }
            common_keys = set(current_positive) & set(prior_positive)
            base_amount, estimated_repayment, disappeared_amount, partial_repayment = (
                estimated_repayment_components(prior_positive, current_positive)
            )
            carried_amount = sum(float(current_positive[key]["amount"]) for key in common_keys)
            interval_speed = normalized_repayment_rate(estimated_repayment, base_amount, elapsed_days)
            history = histories.setdefault(code, [])
            if base_amount > 0 and elapsed_days > 0:
                history.append(
                    {
                        "endDate": parsed_dates[period_date],
                        "speed": interval_speed,
                        "base": base_amount,
                        "days": elapsed_days,
                    }
                )
            recent_cutoff = parsed_dates[period_date] - timedelta(days=90)
            baseline_cutoff = recent_cutoff - timedelta(days=365)
            recent_records = [record for record in history if record["endDate"] > recent_cutoff]
            baseline_records = [
                record
                for record in history
                if baseline_cutoff < record["endDate"] <= recent_cutoff
            ]
            resolution_speed = weighted_mean(
                [
                    (record["speed"], record["base"] * record["days"])
                    for record in recent_records
                ],
                interval_speed,
            )
            historical_speed = weighted_median(
                [
                    (record["speed"], record["base"] * record["days"])
                    for record in baseline_records
                ],
                resolution_speed,
            )
            speed_available = bool(recent_records and baseline_records)

            lifecycle_rows: list[dict[str, Any]] = []
            aged90_amount = 0.0
            aged361_amount = 0.0
            age_numerator = 0.0
            observed_numerator = 0.0
            reappeared_document_count = 0
            for key, item in current_positive.items():
                amount = float(item["amount"])
                invoice_age = max(0, int(item.get("invoiceAge") or 0))
                lifecycle = lifecycle_for(key, date_index)
                is_aged90 = item.get("invoiceAge") is not None and invoice_age > 90
                if is_aged90:
                    aged90_amount += amount
                if item.get("invoiceAge") is not None and invoice_age > 360:
                    aged361_amount += amount
                age_numerator += amount * invoice_age
                observed_numerator += amount * lifecycle["observedOpenDays"]
                if lifecycle["reappearanceCount"]:
                    reappeared_document_count += 1
                lifecycle_rows.append({**item, **lifecycle, "aged90": is_aged90})

            closed_lifecycle_rows: list[dict[str, Any]] = []
            if date_index == len(dates) - 1:
                for key in known_by_customer.get(code, set()) - set(current_positive):
                    lifecycle = closed_lifecycle_for(key, date_index)
                    last_item = risk_doc_maps[lifecycle["lastSeenDate"]][key]
                    closed_lifecycle_rows.append({**last_item, **lifecycle})

            open_document_count = len(current_positive)
            aged90_document_count = sum(1 for item in lifecycle_rows if item["aged90"])
            aged90_document_share = (
                aged90_document_count / open_document_count if open_document_count else 0.0
            )
            current_positive_amount = sum(float(item["amount"]) for item in current_positive.values())
            aged90_amount_share = (
                aged90_amount / current_positive_amount if current_positive_amount > 0 else 0.0
            )
            weighted_invoice_age = (
                age_numerator / current_positive_amount if current_positive_amount > 0 else 0.0
            )
            native_age_basis = [
                {
                    "id": item.get("id"),
                    "invoiceDate": item.get("invoiceDate"),
                    "invoiceAge": max(0, int(item.get("invoiceAge") or 0)),
                    "amount": float(item.get("amount") or 0),
                    "sourceCurrency": currency,
                }
                for item in lifecycle_rows
                if float(item.get("amount") or 0) > 0
            ]
            age_summary = open_age_summary(native_age_basis)
            weighted_observed_days = (
                observed_numerator / current_positive_amount if current_positive_amount > 0 else 0.0
            )
            longest_observed_row = max(
                lifecycle_rows,
                key=lambda item: int(item.get("observedOpenDays") or 0),
                default=None,
            )
            min_observed_open_days = int(longest_observed_row.get("observedOpenDays") or 0) if longest_observed_row else 0
            min_observed_lower_bound = bool(longest_observed_row and longest_observed_row.get("firstSeenLowerBound"))
            max_open_days = max(
                (
                    max(0, int(item["invoiceAge"]))
                    for item in current_positive.values()
                    if item.get("invoiceAge") is not None
                ),
                default=0,
            )
            aged361_document_count = sum(
                1
                for item in lifecycle_rows
                if item.get("invoiceAge") is not None and int(item["invoiceAge"]) > 360
            )
            transition_eligible, transitioned_amount, transition_documents, transition_status = aging_transition_for(
                code, date_index
            )
            aging_transition_rate = (
                transitioned_amount / transition_eligible * 100
                if transition_eligible > 0 else 0.0
            )
            transition_available = transition_eligible > 0

            previous_profile = prior_profiles.get(code, {})
            prior_delay_amount = float(previous_profile.get("aged90DocumentAmountRaw") or 0)
            prior_weighted_age = float(previous_profile.get("weightedInvoiceAge") or 0)
            if prior_delay_amount > 0:
                delay_amount_change_pct = (aged90_amount - prior_delay_amount) / prior_delay_amount * 100
            else:
                delay_amount_change_pct = 0.0
            duration_delta = weighted_invoice_age - prior_weighted_age if previous_profile else 0.0
            duration_mix_delta = duration_delta - elapsed_days if previous_profile and elapsed_days else 0.0

            allowance_exposure = allowance_loss = invoice_over180 = unclassified = 0.0
            scope_excluded = other_receivables = 0.0
            bucket_exposure: dict[str, float] = {key: 0.0 for key in rates}
            for item in signed_current_by_customer.get(code, {}).values():
                amount = float(item["amount"])
                if item.get("scope") == "other":
                    other_receivables += amount
                    continue
                if item.get("scope") != "receivable":
                    continue
                if code in allowance_excluded_codes:
                    scope_excluded += amount
                    continue
                bucket_key = invoice_allowance_bucket(item.get("invoiceAge"))
                if bucket_key is None:
                    unclassified += amount
                    continue
                bucket_exposure[bucket_key] += amount
                allowance_exposure += amount
                allowance_loss += amount * rates[bucket_key]
                if bucket_key in {"age_181_270", "age_271_360", "age_361_plus"}:
                    invoice_over180 += amount

            raw_profiles[code] = {
                "delayAmountRaw": aged90_amount,
                "carriedAmountRaw": carried_amount,
                "priorDelayAmountRaw": prior_delay_amount,
                "delayAmountChangePct": delay_amount_change_pct,
                "weightedInvoiceAge": weighted_invoice_age,
                "openAgeSummary": age_summary,
                "openAgeBasis": native_age_basis,
                "weightedObservedDays": weighted_observed_days,
                "minObservedOpenDays": min_observed_open_days,
                "minObservedLowerBound": min_observed_lower_bound,
                "maxOpenDays": max_open_days,
                "durationDelta": duration_delta,
                "durationMixDelta": duration_mix_delta,
                "speed": resolution_speed,
                "historicalSpeed": historical_speed,
                "speedDelta": (resolution_speed - historical_speed) * 100,
                "speedAvailable": speed_available,
                "speedBaseRaw": base_amount,
                "resolvedRaw": estimated_repayment,
                "disappearedRaw": disappeared_amount,
                "partialRepaymentRaw": partial_repayment,
                "openDocumentCount": open_document_count,
                "aged90DocumentCount": aged90_document_count,
                "aged90DocumentShare": aged90_document_share,
                "aged90DocumentAmountRaw": aged90_amount,
                "aged90AmountShare": aged90_amount_share,
                "aged361DocumentCount": aged361_document_count,
                "aged361DocumentAmountRaw": aged361_amount,
                "positiveOpenAmountRaw": current_positive_amount,
                "agingTransitionEligibleRaw": transition_eligible,
                "agingTransitionedRaw": transitioned_amount,
                "agingTransitionDocuments": transition_documents,
                "agingTransitionRate": aging_transition_rate,
                "transitionAvailable": transition_available,
                "transitionStatus": transition_status,
                "reappearedDocumentCount": reappeared_document_count,
                "lifecycleRows": lifecycle_rows,
                "closedLifecycleRows": closed_lifecycle_rows,
                "allowanceExposureRaw": allowance_exposure,
                "allowanceLossRaw": allowance_loss,
                "invoiceOver180Raw": invoice_over180,
                "unclassifiedRaw": unclassified,
                "scopeExcludedRaw": scope_excluded,
                "otherReceivablesRaw": other_receivables,
                "allowanceBucketsRaw": bucket_exposure,
            }

        for customer in period["customers"]:
            code = customer["id"]
            profile = raw_profiles[code]
            delay_amount = profile["delayAmountRaw"]
            krw_factor = 1.0 if currency == "KRW" else usd_krw_rate
            has_current_exposure = profile["positiveOpenAmountRaw"] > 1e-9
            behavior = receivable_risk_profile(
                aged_amount_eok=profile["aged90DocumentAmountRaw"] * krw_factor / 100_000_000,
                positive_open_eok=profile["positiveOpenAmountRaw"] * krw_factor / 100_000_000,
                average_open_days=profile["weightedInvoiceAge"],
                current_speed_pct=profile["speed"] * 100,
                baseline_speed_pct=profile["historicalSpeed"] * 100,
                transition_rate_pct=profile["agingTransitionRate"],
                transition_eligible_eok=profile["agingTransitionEligibleRaw"] * krw_factor / 100_000_000,
                speed_available=profile["speedAvailable"] and has_current_exposure,
                transition_available=profile["transitionAvailable"] and has_current_exposure,
                transition_status=profile["transitionStatus"] if has_current_exposure else "not_applicable",
                credit_limit_eok=customer.get("limit") if currency == "KRW" else None,
                credit_used_eok=customer.get("used") if currency == "KRW" else None,
                credit_available=bool(customer.get("credit")) and currency == "KRW",
            )
            risk_score = behavior["riskScore"]
            grade, risk_type = behavior["grade"], behavior["riskType"]

            signals: list[str] = []
            if profile["aged90DocumentAmountRaw"] > 0:
                aged_amount_text = (
                    f"{display_amount(profile['aged90DocumentAmountRaw'], currency):,.1f}억"
                    if currency == "KRW"
                    else f"USD {display_amount(profile['aged90DocumentAmountRaw'], currency):,.1f}M"
                )
                signals.append(
                    f"90일 초과 미결액 {aged_amount_text}"
                    f"({profile['aged90AmountShare'] * 100:.0f}%)"
                )
            if profile["weightedInvoiceAge"] > 90:
                signals.append(
                    f"잔액 기준 평균 미결일수 {profile['weightedInvoiceAge']:.0f}일"
                )
            if behavior["speedAvailable"] and (
                behavior["speedScore"] >= 60 or behavior["speedDecline"] >= 15
            ):
                recovery_signal = f"최근 채권회수율 {profile['speed'] * 100:.1f}%"
                if behavior["speedDecline"] >= 10:
                    recovery_signal += f"·기준 대비 {behavior['speedDecline']:.1f}%p 감소"
                signals.append(recovery_signal)
            if behavior["transitionAvailable"] and behavior["transitionRate"] >= 25:
                signals.append(f"장기화 잔존율 {behavior['transitionRate']:.1f}%")
            if customer.get("util") is not None and customer["util"] > 100:
                signals.append("여신한도 초과")
            elif behavior["creditAdjustment"] >= 3:
                signals.append(
                    f"여신노출 압력 {behavior['creditPressureScore']:.1f}점"
                )
            if not signals:
                signals.append("현재 뚜렷한 악화신호 없음" if behavior["coverage"] >= 80 else "행동자료 부족")

            documents = []
            closed_documents = []
            if date_index == len(dates) - 1:
                documents = sorted(
                    profile["lifecycleRows"],
                    key=lambda item: (
                        -(item.get("invoiceAge") if item.get("invoiceAge") is not None else -1),
                        -abs(float(item.get("amount") or 0)),
                    ),
                )[:15]
                closed_documents = sorted(
                    profile["closedLifecycleRows"],
                    key=lambda item: (item["closureConfirmedDate"], abs(float(item.get("amount") or 0))),
                    reverse=True,
                )[:10]
            customer.update(
                {
                    "grade": grade,
                    "riskType": risk_type,
                    "riskScore": risk_score,
                    "persistence": profile["aged90DocumentCount"],
                    "delayAmount": display_amount(delay_amount, currency),
                    "delayAmountChange": display_amount(delay_amount - profile["priorDelayAmountRaw"], currency),
                    "delayAmountChangePct": round(profile["delayAmountChangePct"], 1),
                    "delayDays": round(profile["weightedInvoiceAge"], 1),
                    "balanceWeightedAverageOpenDays": round(profile["weightedInvoiceAge"], 1),
                    "weightedObservedDays": round(profile["weightedObservedDays"], 1),
                    "minObservedOpenDays": profile["minObservedOpenDays"],
                    "minObservedLowerBound": profile["minObservedLowerBound"],
                    "maxOpenDays": int(profile["maxOpenDays"]),
                    "durationDelta": round(profile["durationDelta"], 1),
                    "durationMixDelta": round(profile["durationMixDelta"], 1),
                    "speed": round(profile["speed"] * 100, 1),
                    "speedPrior": round(profile["historicalSpeed"] * 100, 1),
                    "speedDelta": round(profile["speedDelta"], 1),
                    "speedBase": display_amount(profile["speedBaseRaw"], currency),
                    "repaymentBase": display_amount(profile["speedBaseRaw"], currency),
                    "resolvedAmount": display_amount(profile["resolvedRaw"], currency),
                    "estimatedRepaymentAmount": display_amount(profile["resolvedRaw"], currency),
                    "disappearedRepaymentAmount": display_amount(profile["disappearedRaw"], currency),
                    "partialRepaymentAmount": display_amount(profile["partialRepaymentRaw"], currency),
                    "openDocumentCount": profile["openDocumentCount"],
                    "aged90DocumentCount": profile["aged90DocumentCount"],
                    "aged90DocumentShare": round(profile["aged90DocumentShare"] * 100, 1),
                    "aged90DocumentAmount": display_amount(profile["aged90DocumentAmountRaw"], currency),
                    "aged90AmountShare": round(profile["aged90AmountShare"] * 100, 1),
                    "aged361DocumentCount": profile["aged361DocumentCount"],
                    "aged361DocumentAmount": display_amount(profile["aged361DocumentAmountRaw"], currency),
                    "positiveOpenAmount": display_amount(profile["positiveOpenAmountRaw"], currency),
                    "openAgeBuckets": [
                        {
                            **bucket,
                            "amount": display_amount(bucket["amount"], currency),
                            "share": round(bucket["share"], 1),
                        }
                        for bucket in profile["openAgeSummary"]["buckets"]
                    ],
                    "averageOpenDayContributors": [
                        {
                            **contributor,
                            "amount": display_amount(contributor["amount"], currency),
                            "contributionDays": round(contributor["contributionDays"], 1),
                        }
                        for contributor in profile["openAgeSummary"]["contributors"]
                    ],
                    "_openAgeBasis": profile["openAgeBasis"],
                    "agingTransitionEligibleAmount": display_amount(profile["agingTransitionEligibleRaw"], currency),
                    "agingTransitionedAmount": display_amount(profile["agingTransitionedRaw"], currency),
                    "agingTransitionDocuments": profile["agingTransitionDocuments"],
                    "agingTransitionEligibleEok": round(
                        profile["agingTransitionEligibleRaw"] * krw_factor / 100_000_000, 4
                    ),
                    "agingTransitionRate": round(profile["agingTransitionRate"], 1),
                    "reappearedDocumentCount": profile["reappearedDocumentCount"],
                    # Backward-compatible aliases now mean distinct documents, not snapshots.
                    "delayFrequency": profile["aged90DocumentCount"],
                    "delayEligible": profile["openDocumentCount"],
                    "allowanceExposure": display_amount(profile["allowanceExposureRaw"], currency),
                    "allowanceExpectedLoss": display_amount(profile["allowanceLossRaw"], currency),
                    "invoiceOver180": display_amount(profile["invoiceOver180Raw"], currency),
                    "allowanceOver180": display_amount(
                        sum(
                            profile["allowanceBucketsRaw"].get(key, 0.0)
                            for key in ("age_181_270", "age_271_360", "age_361_plus")
                        ),
                        currency,
                    ),
                    "allowanceOver360": display_amount(
                        profile["allowanceBucketsRaw"].get("age_361_plus", 0.0),
                        currency,
                    ),
                    "allowanceScopeExcluded": display_amount(profile["scopeExcludedRaw"], currency),
                    "otherReceivablesExposure": display_amount(profile["otherReceivablesRaw"], currency),
                    "subsidiaryCandidate": code.isdigit() and len(code) == 4,
                    "signal": " · ".join(signals[:4]),
                    "action": "재무 중점 점검" if grade in {"R1", "R2"} else "정기 모니터링",
                    "documents": [
                        {
                            "id": item["id"],
                            "documentKey": item["documentKey"],
                            "invoiceDate": item["invoiceDate"],
                            "invoiceAge": item["invoiceAge"],
                            "dueDate": item["dueDate"],
                            "days": item["days"],
                            "scope": item["scope"],
                            "amount": display_amount(item["amount"], currency),
                            "terms": item["terms"],
                            "firstSeenDate": item["firstSeenDate"],
                            "lastSeenDate": item["lastSeenDate"],
                            "continuousStartDate": item["continuousStartDate"],
                            "observationCount": item["observationCount"],
                            "observedOpenDays": item["observedOpenDays"],
                            "reappearanceCount": item["reappearanceCount"],
                            "firstSeenLowerBound": item["firstSeenLowerBound"],
                            "aged90": item["aged90"],
                        }
                        for item in documents
                    ],
                    "closedDocuments": [
                        {
                            "id": item["id"],
                            "documentKey": item["documentKey"],
                            "invoiceDate": item["invoiceDate"],
                            "invoiceAge": item["invoiceAge"],
                            "scope": item["scope"],
                            "amount": display_amount(item["amount"], currency),
                            "firstSeenDate": item["firstSeenDate"],
                            "lastSeenDate": item["lastSeenDate"],
                            "continuousStartDate": item["continuousStartDate"],
                            "closureConfirmedDate": item["closureConfirmedDate"],
                            "observationCount": item["observationCount"],
                            "observedOpenDays": item["observedOpenDays"],
                            "reappearanceCount": item["reappearanceCount"],
                            "firstSeenLowerBound": item["firstSeenLowerBound"],
                        }
                        for item in closed_documents
                    ],
                    "behavior": behavior,
                }
            )
            prior_profiles[code] = profile

        period["customers"].sort(
            key=lambda item: (
                grade_order[item["grade"]],
                -item["riskScore"],
                -item["delayAmount"],
                -item["total"],
            )
        )
        high_risk = [customer for customer in period["customers"] if customer["grade"] in {"R1", "R2"}]
        slowdown = [
            customer
            for customer in period["customers"]
            if float((customer.get("behavior") or {}).get("speedDecline") or 0) >= 10
        ]
        deteriorating = [
            customer
            for customer in period["customers"]
            if customer["grade"] in {"R1", "R2"}
            or (
                float((customer.get("behavior") or {}).get("amountScore") or 0) >= 70
                and sum(
                (
                    float((customer.get("behavior") or {}).get("speedScore") or 0) >= 60,
                    float((customer.get("behavior") or {}).get("agedShare") or 0) >= 45,
                    int(customer.get("maxOpenDays") or 0) > 270,
                    float((customer.get("behavior") or {}).get("transitionRate") or 0) >= 50,
                )
                ) >= 2
            )
        ]
        speed_base_total = sum(profile["speedBaseRaw"] for profile in raw_profiles.values())
        resolved_total = sum(profile["resolvedRaw"] for profile in raw_profiles.values())
        eligible_speed_profiles = [
            profile for profile in raw_profiles.values() if profile["speedAvailable"]
        ]
        portfolio_speed = weighted_mean(
            [(profile["speed"], max(profile["speedBaseRaw"], 0.0)) for profile in eligible_speed_profiles],
            0.0,
        )
        portfolio_prior_speed = weighted_mean(
            [(profile["historicalSpeed"], max(profile["speedBaseRaw"], 0.0)) for profile in eligible_speed_profiles],
            portfolio_speed,
        )
        transition_eligible_total = sum(profile["agingTransitionEligibleRaw"] for profile in raw_profiles.values())
        transitioned_total = sum(profile["agingTransitionedRaw"] for profile in raw_profiles.values())
        portfolio_transition_rate = (
            transitioned_total / transition_eligible_total * 100
            if transition_eligible_total > 0 else 0.0
        )
        positive_open_total = sum(profile["positiveOpenAmountRaw"] for profile in raw_profiles.values())
        aged361_total = sum(profile["aged361DocumentAmountRaw"] for profile in raw_profiles.values())
        aged361_customers = sum(1 for profile in raw_profiles.values() if profile["aged361DocumentAmountRaw"] > 0)
        period["summary"].update(
            {
                "highRiskCount": len(high_risk),
                "highRiskAmount": round(sum(max(0.0, customer["positiveOpenAmount"]) for customer in high_risk), 4),
                "delayAmount": round(sum(max(0.0, customer["delayAmount"]) for customer in period["customers"]), 4),
                "positiveOpenAmount": display_amount(positive_open_total, currency),
                "slowdownCount": len(slowdown),
                "slowdownAmount": round(sum(max(0.0, customer["positiveOpenAmount"]) for customer in slowdown), 4),
                "deterioratingCount": len(deteriorating),
                "deterioratingAmount": round(sum(max(0.0, customer["total"]) for customer in deteriorating), 4),
                "openDocumentCount": sum(customer["openDocumentCount"] for customer in period["customers"]),
                "maxOpenDays": max((int(customer.get("maxOpenDays") or 0) for customer in period["customers"]), default=0),
                "aged90DocumentCount": sum(customer["aged90DocumentCount"] for customer in period["customers"]),
                "aged90DocumentAmount": round(sum(max(0.0, customer["aged90DocumentAmount"]) for customer in period["customers"]), 4),
                "aged90AmountShare": round(
                    sum(max(0.0, customer["aged90DocumentAmount"]) for customer in period["customers"])
                    / display_amount(positive_open_total, currency) * 100,
                    1,
                ) if positive_open_total > 0 else 0.0,
                "aged361DocumentCount": sum(customer["aged361DocumentCount"] for customer in period["customers"]),
                "aged361DocumentAmount": display_amount(aged361_total, currency),
                "aged361CustomerCount": aged361_customers,
                "reappearedDocumentCount": sum(customer["reappearedDocumentCount"] for customer in period["customers"]),
                # Retained for older HTML consumers; value is now 90-day document exposure.
                "repeatDelayAmount": round(sum(max(0.0, customer["aged90DocumentAmount"]) for customer in period["customers"]), 4),
                "repaymentBase": display_amount(speed_base_total, currency),
                "estimatedRepaymentAmount": display_amount(resolved_total, currency),
                "resolutionSpeed": round(portfolio_speed * 100, 1),
                "priorResolutionSpeed": round(portfolio_prior_speed * 100, 1),
                "recoverySpeedDecline": round(max(0.0, (portfolio_prior_speed - portfolio_speed) * 100), 1),
                "agingTransitionRate": round(portfolio_transition_rate, 1),
                "agingTransitionEligibleAmount": display_amount(transition_eligible_total, currency),
            }
        )
        if speed_base_total > 0 and elapsed_days > 0:
            portfolio_speed_history.append((portfolio_speed, speed_base_total, elapsed_days))

        lifecycle_groups = {
            "관찰 30일 미만": 0.0,
            "관찰 30–89일": 0.0,
            "관찰 90–179일": 0.0,
            "관찰 180일 이상": 0.0,
        }
        for profile in raw_profiles.values():
            for item in profile["lifecycleRows"]:
                observed_days = int(item["observedOpenDays"])
                amount = display_amount(float(item["amount"]), currency)
                if observed_days < 30:
                    lifecycle_groups["관찰 30일 미만"] += amount
                elif observed_days < 90:
                    lifecycle_groups["관찰 30–89일"] += amount
                elif observed_days < 180:
                    lifecycle_groups["관찰 90–179일"] += amount
                else:
                    lifecycle_groups["관찰 180일 이상"] += amount
        period["persistence"] = [
            {"label": label, "value": round(value, 4), "risk": label in {"관찰 90–179일", "관찰 180일 이상"}}
            for label, value in lifecycle_groups.items()
        ]
        period["aging"]["behavior"] = [
            {
                "label": label,
                "value": round(sum(max(0.0, customer["total"]) for customer in period["customers"] if customer["grade"] == grade), 4),
                "risk": grade in {"R1", "R2"},
            }
            for grade, label in (("R1", "R1 최고위험"), ("R2", "R2 고위험"), ("R3", "R3 주의"), ("R4", "R4 낮음"))
        ]

        allowance_buckets: list[dict[str, Any]] = []
        for model_bucket in allowance_model["buckets"]:
            key = model_bucket["key"]
            exposure_raw = sum(profile["allowanceBucketsRaw"][key] for profile in raw_profiles.values())
            expected_raw = exposure_raw * rates[key]
            customer_count = sum(1 for profile in raw_profiles.values() if profile["allowanceBucketsRaw"][key] > 0)
            allowance_buckets.append(
                {
                    "key": key,
                    "label": labels[key],
                    "exposure": display_amount(exposure_raw, currency),
                    "appliedRate": round(rates[key] * 100, 4),
                    "rawRate": round(raw_rates[key], 4),
                    "method": methods[key],
                    "expectedLoss": display_amount(expected_raw, currency),
                    "customerCount": customer_count,
                    "risk": key in {"age_181_270", "age_271_360", "age_361_plus"},
                }
            )
        allowance_total = sum(row["exposure"] for row in allowance_buckets)
        allowance_loss = sum(row["expectedLoss"] for row in allowance_buckets)
        period["allowance"] = {
            "buckets": allowance_buckets,
            "totalExposure": round(allowance_total, 4),
            "expectedLoss": round(allowance_loss, 4),
            "weightedRate": round(allowance_loss / allowance_total * 100, 2) if allowance_total else 0.0,
            "over360Exposure": next(row["exposure"] for row in allowance_buckets if row["key"] == "age_361_plus"),
            "unclassifiedExposure": display_amount(sum(profile["unclassifiedRaw"] for profile in raw_profiles.values()), currency),
            "scopeExcludedExposure": display_amount(sum(profile["scopeExcludedRaw"] for profile in raw_profiles.values()), currency),
            "otherReceivablesExposure": display_amount(sum(profile["otherReceivablesRaw"] for profile in raw_profiles.values()), currency),
            "scopeExcludedCustomerCodes": sorted(allowance_excluded_codes),
            "signedBasis": True,
        }
        top_notices = sorted(high_risk, key=lambda item: (-item["riskScore"], -item["aged90DocumentAmount"]))[:3]
        period["notices"] = [
            {
                "id": customer["id"],
                "grade": customer["grade"],
                "name": customer["name"],
                "score": customer["riskScore"],
                "amount": customer["total"],
                "detail": customer["signal"],
            }
            for customer in top_notices
        ]


def combine_currency_views(
    krw_data: dict[str, Any],
    usd_data: dict[str, Any],
    usd_krw_rate: float,
    credit_map: dict[str, dict[str, Any]],
    allowance_model: dict[str, Any],
) -> dict[str, Any]:
    """Combine native KRW and USD views into KRW-equivalent display units (억원)."""
    usd_to_eok = lambda value: round(float(value or 0) * usd_krw_rate / 100, 4)

    def merge_amount_rows(krw_rows: list[dict[str, Any]], usd_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        usd_by_label = {row["label"]: row for row in usd_rows}
        labels = [row["label"] for row in krw_rows] + [row["label"] for row in usd_rows if row["label"] not in {item["label"] for item in krw_rows}]
        krw_by_label = {row["label"]: row for row in krw_rows}
        return [
            {
                "label": label,
                "value": round(float(krw_by_label.get(label, {}).get("value", 0)) + usd_to_eok(usd_by_label.get(label, {}).get("value", 0)), 4),
                "risk": bool(krw_by_label.get(label, {}).get("risk") or usd_by_label.get(label, {}).get("risk")),
            }
            for label in labels
        ]

    def concentration(customers: list[dict[str, Any]]) -> dict[str, float]:
        total = sorted((max(0.0, float(row["total"] or 0)) for row in customers), reverse=True)
        overdue = sorted((max(0.0, float(row["overdue"] or 0)) for row in customers), reverse=True)
        over90 = sorted((max(0.0, float(row["over90"] or 0)) for row in customers), reverse=True)
        return {
            "top1Total": round(total[0] / sum(total) * 100 if total and sum(total) else 0, 1),
            "top5Total": round(sum(total[:5]) / sum(total) * 100 if sum(total) else 0, 1),
            "top5Overdue": round(sum(overdue[:5]) / sum(overdue) * 100 if sum(overdue) else 0, 1),
            "top5Over90": round(sum(over90[:5]) / sum(over90) * 100 if sum(over90) else 0, 1),
        }

    krw_periods = {period["date"]: period for period in krw_data["periods"]}
    usd_periods = {period["date"]: period for period in usd_data["periods"]}
    dates = sorted(set(krw_periods) | set(usd_periods))
    periods: list[dict[str, Any]] = []
    prior_customers: dict[str, dict[str, Any]] = {}
    portfolio_speed_history: list[float] = []

    for date_index, period_date in enumerate(dates):
        kp = krw_periods.get(period_date)
        up = usd_periods.get(period_date)
        if not kp or not up:
            continue
        is_latest = period_date == dates[-1]
        is_model_date = period_date == allowance_model.get("modelDate")
        allowance_fx_rate = (
            float(allowance_model.get("modelFxRate") or usd_krw_rate)
            if is_model_date
            else usd_krw_rate
        )
        allowance_usd_to_eok = lambda value: round(float(value or 0) * allowance_fx_rate / 100, 4)
        customer_map: dict[str, dict[str, Any]] = {}
        for source_currency, period, converter in (("KRW", kp, float), ("USD", up, usd_to_eok)):
            allowance_converter = float if source_currency == "KRW" else allowance_usd_to_eok
            for source in period["customers"]:
                code = source["id"]
                target = customer_map.setdefault(
                    code,
                    {
                        "id": code,
                        "name": source["name"],
                        "total": 0.0,
                        "gross": 0.0,
                        "credits": 0.0,
                        "overdue": 0.0,
                        "over90": 0.0,
                        "change": 0.0,
                        "persistence": 0,
                        "delayAmount": 0.0,
                        "priorDelayAmount": 0.0,
                        "repaymentBase": 0.0,
                        "estimatedRepaymentAmount": 0.0,
                        "disappearedRepaymentAmount": 0.0,
                        "partialRepaymentAmount": 0.0,
                        "openDocumentCount": 0,
                        "maxOpenDays": 0,
                        "minObservedOpenDays": 0,
                        "minObservedLowerBound": False,
                        "aged90DocumentCount": 0,
                        "aged90DocumentAmount": 0.0,
                        "aged361DocumentCount": 0,
                        "aged361DocumentAmount": 0.0,
                        "positiveOpenAmount": 0.0,
                        "agingTransitionEligibleAmount": 0.0,
                        "agingTransitionedAmount": 0.0,
                        "agingTransitionDocuments": 0,
                        "reappearedDocumentCount": 0,
                        "allowanceExposure": 0.0,
                        "allowanceExpectedLoss": 0.0,
                        "invoiceOver180": 0.0,
                        "allowanceOver180": 0.0,
                        "allowanceOver360": 0.0,
                        "allowanceScopeExcluded": 0.0,
                        "otherReceivablesExposure": 0.0,
                        "subsidiaryCandidate": False,
                        "speedWeighted": 0.0,
                        "speedPriorWeighted": 0.0,
                        "speedWeight": 0.0,
                        "speedAvailable": False,
                        "transitionAvailable": False,
                        "transitionDataSufficient": False,
                        "delayDaysWeighted": 0.0,
                        "observedDaysWeighted": 0.0,
                        "delayDaysWeight": 0.0,
                        "durationDeltaWeighted": 0.0,
                        "durationMixDeltaWeighted": 0.0,
                        "scopeAmounts": {"receivable": 0.0, "other": 0.0},
                        "currencies": [],
                        "documents": [],
                        "closedDocuments": [],
                        "_openAgeBasis": [],
                    },
                )
                target["name"] = target["name"] or source["name"]
                for field in ("total", "gross", "credits", "overdue", "over90", "change"):
                    target[field] = round(target[field] + converter(source.get(field, 0)), 4)
                for field in (
                    "delayAmount", "aged90DocumentAmount", "repaymentBase",
                    "estimatedRepaymentAmount", "disappearedRepaymentAmount",
                    "partialRepaymentAmount", "aged361DocumentAmount",
                    "positiveOpenAmount", "agingTransitionEligibleAmount",
                    "agingTransitionedAmount",
                ):
                    target[field] = round(target[field] + converter(source.get(field, 0)), 4)
                for field in (
                    "allowanceExposure", "allowanceExpectedLoss", "invoiceOver180",
                    "allowanceOver180", "allowanceOver360",
                    "allowanceScopeExcluded", "otherReceivablesExposure",
                ):
                    target[field] = round(target[field] + allowance_converter(source.get(field, 0)), 4)
                target["subsidiaryCandidate"] = bool(
                    target["subsidiaryCandidate"] or source.get("subsidiaryCandidate")
                )
                converted_delay_change = converter(source.get("delayAmountChange", 0))
                target["priorDelayAmount"] += max(0.0, converter(source.get("delayAmount", 0)) - converted_delay_change)
                target["openDocumentCount"] += int(source.get("openDocumentCount") or 0)
                target["maxOpenDays"] = max(target["maxOpenDays"], int(source.get("maxOpenDays") or 0))
                source_observed_days = int(source.get("minObservedOpenDays") or 0)
                if source_observed_days > target["minObservedOpenDays"]:
                    target["minObservedOpenDays"] = source_observed_days
                    target["minObservedLowerBound"] = bool(source.get("minObservedLowerBound"))
                target["aged90DocumentCount"] += int(source.get("aged90DocumentCount") or 0)
                target["aged361DocumentCount"] += int(source.get("aged361DocumentCount") or 0)
                target["agingTransitionDocuments"] += int(source.get("agingTransitionDocuments") or 0)
                target["reappearedDocumentCount"] += int(source.get("reappearedDocumentCount") or 0)
                speed_weight = max(abs(converter(source.get("speedBase", 0))), 0.001)
                delay_days_weight = max(abs(allowance_converter(source.get("allowanceExposure", 0))), 0.001)
                target["speedWeighted"] += float(source.get("speed") or 0) * speed_weight
                target["speedPriorWeighted"] += float(source.get("speedPrior") or 0) * speed_weight
                target["speedWeight"] += speed_weight
                target["speedAvailable"] = bool(
                    target["speedAvailable"] or (source.get("behavior") or {}).get("speedAvailable")
                )
                target["transitionAvailable"] = bool(
                    target["transitionAvailable"] or (source.get("behavior") or {}).get("transitionAvailable")
                )
                target["transitionDataSufficient"] = bool(
                    target["transitionDataSufficient"]
                    or (source.get("behavior") or {}).get("transitionStatus")
                    in {"available", "not_applicable"}
                )
                target["delayDaysWeighted"] += float(source.get("delayDays") or 0) * delay_days_weight
                target["observedDaysWeighted"] += float(source.get("weightedObservedDays") or 0) * delay_days_weight
                target["durationDeltaWeighted"] += float(source.get("durationDelta") or 0) * delay_days_weight
                target["durationMixDeltaWeighted"] += float(source.get("durationMixDelta") or 0) * delay_days_weight
                target["delayDaysWeight"] += delay_days_weight
                target["scopeAmounts"][source["scope"]] += abs(converter(source.get("total", 0)))
                target["currencies"].append(source_currency)
                target["documents"].extend(
                    [
                        {
                            **document,
                            "sourceAmount": document.get("amount", 0),
                            "amount": converter(document.get("amount", 0)),
                            "sourceCurrency": source_currency,
                        }
                        for document in source.get("documents", [])
                    ]
                )
                target["_openAgeBasis"].extend(
                    [
                        {
                            **basis,
                            "amount": (
                                float(basis.get("amount", 0)) / 100_000_000
                                if source_currency == "KRW"
                                else float(basis.get("amount", 0)) / 1_000_000 * usd_krw_rate / 100
                            ),
                            "sourceCurrency": source_currency,
                        }
                        for basis in source.get("_openAgeBasis", [])
                    ]
                )
                target["closedDocuments"].extend(
                    [
                        {
                            **document,
                            "sourceAmount": document.get("amount", 0),
                            "amount": converter(document.get("amount", 0)),
                            "sourceCurrency": source_currency,
                        }
                        for document in source.get("closedDocuments", [])
                    ]
                )

        elapsed_days = (
            max(1, (datetime.strptime(period_date, "%Y-%m-%d") - datetime.strptime(dates[date_index - 1], "%Y-%m-%d")).days)
            if date_index else 0
        )
        customers: list[dict[str, Any]] = []
        for code, row in customer_map.items():
            credit = credit_map.get(code) if is_latest else None
            limit_value = float(credit["credit_limit"] or 0) / 100_000_000 if credit else None
            used_value = float(credit["credit_used"] or 0) / 100_000_000 if credit else None
            util = None
            if credit:
                util = used_value / limit_value * 100 if limit_value and limit_value > 0 else (999.0 if used_value and used_value > 0 else 0.0)
            open_document_count = int(row["openDocumentCount"])
            aged90_document_count = int(row["aged90DocumentCount"])
            aged90_document_share = (
                aged90_document_count / open_document_count * 100 if open_document_count else 0.0
            )
            delay_days = row["delayDaysWeighted"] / row["delayDaysWeight"] if row["delayDaysWeight"] else 0.0
            weighted_observed_days = row["observedDaysWeighted"] / row["delayDaysWeight"] if row["delayDaysWeight"] else 0.0
            speed = row["speedWeighted"] / row["speedWeight"] if row["speedWeight"] else 0.0
            speed_prior = row["speedPriorWeighted"] / row["speedWeight"] if row["speedWeight"] else speed
            speed_delta = speed - speed_prior
            aging_transition_rate = (
                row["agingTransitionedAmount"] / row["agingTransitionEligibleAmount"] * 100
                if row["agingTransitionEligibleAmount"] > 0 else 0.0
            )
            transition_status = (
                "available"
                if row["transitionAvailable"] and row["agingTransitionEligibleAmount"] > 0
                else "not_applicable"
                if row["transitionDataSufficient"]
                else "insufficient_history"
            )
            # Native-currency views are already customer-netted.  This second
            # cap removes any remaining cross-currency credit balance after KRW
            # conversion so the combined view also uses the current net claim.
            pre_cross_positive = max(0.0, float(row["positiveOpenAmount"] or 0))
            net_open_amount = max(0.0, float(row["total"] or 0))
            cross_currency_offset = max(0.0, pre_cross_positive - net_open_amount)
            row["aged90DocumentAmount"] = round(
                max(0.0, float(row["aged90DocumentAmount"] or 0) - cross_currency_offset), 4
            )
            row["aged361DocumentAmount"] = round(
                max(0.0, float(row["aged361DocumentAmount"] or 0) - cross_currency_offset), 4
            )
            row["agingTransitionedAmount"] = round(
                max(0.0, float(row["agingTransitionedAmount"] or 0) - cross_currency_offset), 4
            )
            row["positiveOpenAmount"] = round(net_open_amount, 4)
            row["delayAmount"] = row["aged90DocumentAmount"]
            aging_transition_rate = (
                row["agingTransitionedAmount"] / row["agingTransitionEligibleAmount"] * 100
                if row["agingTransitionEligibleAmount"] > 0 else 0.0
            )
            remaining_cross_offset = cross_currency_offset
            netted_detail_documents: list[dict[str, Any]] = []
            for document in sorted(
                row["documents"],
                key=lambda item: -(int(item.get("invoiceAge") or 0)),
            ):
                document_amount = max(0.0, float(document.get("amount") or 0))
                document_offset = min(document_amount, remaining_cross_offset)
                remaining_cross_offset -= document_offset
                residual_amount = document_amount - document_offset
                if residual_amount <= 0.00005 or net_open_amount <= 0:
                    continue
                net_document = dict(document)
                net_document["amount"] = round(residual_amount, 4)
                netted_detail_documents.append(net_document)
            row["documents"] = netted_detail_documents
            netted_age_basis = net_open_age_basis(
                row["_openAgeBasis"],
                cross_currency_offset,
                net_open_amount,
            )
            row["_openAgeBasis"] = netted_age_basis
            combined_age_summary = open_age_summary(netted_age_basis)
            delay_days = combined_age_summary["averageDays"]
            if net_open_amount <= 0:
                row["openDocumentCount"] = 0
                row["aged90DocumentCount"] = 0
                row["aged361DocumentCount"] = 0
                row["maxOpenDays"] = 0
            elif row["aged90DocumentAmount"] <= 0:
                row["aged90DocumentCount"] = 0
                row["aged361DocumentCount"] = 0
                row["maxOpenDays"] = min(int(row["maxOpenDays"] or 0), 90)
            elif row["aged361DocumentAmount"] <= 0:
                row["aged361DocumentCount"] = 0
                row["maxOpenDays"] = min(int(row["maxOpenDays"] or 0), 360)
            behavior = receivable_risk_profile(
                aged_amount_eok=row["aged90DocumentAmount"],
                positive_open_eok=row["positiveOpenAmount"],
                average_open_days=delay_days,
                current_speed_pct=speed,
                baseline_speed_pct=speed_prior,
                transition_rate_pct=aging_transition_rate,
                transition_eligible_eok=row["agingTransitionEligibleAmount"],
                speed_available=row["speedAvailable"] and net_open_amount > 0,
                transition_available=(
                    row["transitionAvailable"]
                    and row["agingTransitionEligibleAmount"] > 0
                    and net_open_amount > 0
                ),
                transition_status=transition_status if net_open_amount > 0 else "not_applicable",
                credit_limit_eok=limit_value,
                credit_used_eok=used_value,
                credit_available=credit is not None,
            )
            duration_delta = row["durationDeltaWeighted"] / row["delayDaysWeight"] if row["delayDaysWeight"] else 0.0
            duration_mix_delta = row["durationMixDeltaWeighted"] / row["delayDaysWeight"] if row["delayDaysWeight"] else 0.0
            delay_change = row["delayAmount"] - row["priorDelayAmount"]
            delay_change_pct = delay_change / row["priorDelayAmount"] * 100 if row["priorDelayAmount"] > 0 else (100.0 if row["delayAmount"] > 0 else 0.0)
            risk_score = behavior["riskScore"]
            grade, risk_type = behavior["grade"], behavior["riskType"]
            signals: list[str] = []
            if len(set(row["currencies"])) > 1:
                signals.append("KRW·USD 혼합")
            if row["aged90DocumentAmount"] > 0:
                signals.append(f"90일 초과 미결액 {row['aged90DocumentAmount']:.1f}억({behavior['agedShare']:.0f}%)")
            if delay_days > 90:
                signals.append(f"잔액 기준 평균 미결일수 {delay_days:.0f}일")
            if behavior["speedAvailable"] and (
                behavior["speedScore"] >= 60 or behavior["speedDecline"] >= 15
            ):
                recovery_signal = f"최근 채권회수율 {speed:.1f}%"
                if behavior["speedDecline"] >= 10:
                    recovery_signal += f"·기준 대비 {behavior['speedDecline']:.1f}%p 감소"
                signals.append(recovery_signal)
            if behavior["transitionAvailable"] and behavior["transitionRate"] >= 25:
                signals.append(f"장기화 잔존율 {behavior['transitionRate']:.1f}%")
            if util is not None and util > 100:
                signals.append("여신한도 초과")
            elif behavior["creditAdjustment"] >= 3:
                signals.append(
                    f"여신노출 압력 {behavior['creditPressureScore']:.1f}점"
                )
            if not signals:
                signals.append("현재 뚜렷한 악화신호 없음" if behavior["coverage"] >= 80 else "행동자료 부족")
            scope = "other" if row["scopeAmounts"]["other"] > row["scopeAmounts"]["receivable"] else "receivable"
            customers.append(
                {
                    **{key: row[key] for key in (
                        "id", "name", "total", "gross", "credits", "overdue", "over90",
                        "change", "delayAmount", "aged90DocumentAmount", "allowanceExposure",
                        "allowanceExpectedLoss", "invoiceOver180", "allowanceOver180",
                        "allowanceOver360",
                        "allowanceScopeExcluded",
                        "otherReceivablesExposure", "subsidiaryCandidate", "repaymentBase",
                        "estimatedRepaymentAmount", "disappearedRepaymentAmount",
                        "partialRepaymentAmount",
                        "maxOpenDays", "minObservedOpenDays", "minObservedLowerBound",
                        "aged361DocumentCount", "aged361DocumentAmount",
                        "positiveOpenAmount", "agingTransitionEligibleAmount",
                        "agingTransitionedAmount", "agingTransitionDocuments",
                    )},
                    "grade": grade,
                    "riskType": risk_type,
                    "riskScore": risk_score,
                    "delayAmountChange": round(delay_change, 4),
                    "delayAmountChangePct": round(delay_change_pct, 1),
                    "delayDays": round(delay_days, 1),
                    "balanceWeightedAverageOpenDays": round(delay_days, 1),
                    "weightedObservedDays": round(weighted_observed_days, 1),
                    "durationDelta": round(duration_delta, 1),
                    "durationMixDelta": round(duration_mix_delta, 1),
                    "speed": round(speed, 1),
                    "speedPrior": round(speed_prior, 1),
                    "speedDelta": round(speed_delta, 1),
                    "openDocumentCount": open_document_count,
                    "aged90DocumentCount": aged90_document_count,
                    "aged90DocumentShare": round(aged90_document_share, 1),
                    "aged90AmountShare": round(behavior["agedShare"], 1),
                    "agingTransitionRate": round(aging_transition_rate, 1),
                    "agingTransitionEligibleEok": round(row["agingTransitionEligibleAmount"], 4),
                    "openAgeBuckets": [
                        {
                            **bucket,
                            "amount": round(bucket["amount"], 4),
                            "share": round(bucket["share"], 1),
                        }
                        for bucket in combined_age_summary["buckets"]
                    ],
                    "averageOpenDayContributors": [
                        {
                            **contributor,
                            "amount": round(contributor["amount"], 4),
                            "contributionDays": round(contributor["contributionDays"], 1),
                        }
                        for contributor in combined_age_summary["contributors"]
                    ],
                    "_openAgeBasis": row["_openAgeBasis"],
                    "reappearedDocumentCount": int(row["reappearedDocumentCount"]),
                    "persistence": aged90_document_count,
                    "delayFrequency": aged90_document_count,
                    "delayEligible": open_document_count,
                    "behavior": behavior,
                    "limit": None if limit_value is None else round(limit_value, 4),
                    "used": None if used_value is None else round(used_value, 4),
                    "util": None if util is None else round(util, 1),
                    "credit": credit is not None,
                    "scope": scope,
                    "signal": " · ".join(signals[:4]),
                    "action": "재무 중점 점검" if grade in {"R1", "R2"} else "정기 모니터링",
                    "documents": sorted(row["documents"], key=lambda item: -abs(item["amount"]))[:15],
                    "closedDocuments": sorted(
                        row["closedDocuments"],
                        key=lambda item: (item.get("closureConfirmedDate") or "", abs(item["amount"])),
                        reverse=True,
                    )[:10],
                }
            )
        grade_order = {"R1": 0, "R2": 1, "R3": 2, "R4": 3}
        customers.sort(key=lambda item: (grade_order[item["grade"]], -item["riskScore"], -item["delayAmount"], -item["total"]))

        movements: list[dict[str, Any]] = []
        current_by_code = {row["id"]: row for row in customers}
        for code in set(prior_customers) | set(current_by_code):
            before = prior_customers.get(code)
            after = current_by_code.get(code)
            bt, at = (before or {}).get("total", 0), (after or {}).get("total", 0)
            bo, ao = (before or {}).get("overdue", 0), (after or {}).get("overdue", 0)
            b90, a90 = (before or {}).get("over90", 0), (after or {}).get("over90", 0)
            total_change, overdue_change = at - bt, ao - bo
            movement_type = None
            if bo <= 0 < ao:
                movement_type = "SAP 만기경과 진입"
            elif bo > 0 >= ao:
                movement_type = "SAP 만기경과 해소"
            elif overdue_change >= 0.5:
                movement_type = "SAP 만기경과 증가"
            elif overdue_change <= -0.5:
                movement_type = "SAP 만기경과 감소"
            elif total_change >= 5:
                movement_type = "채권 증가"
            elif total_change <= -5:
                movement_type = "채권 감소"
            if movement_type:
                source = after or before
                movements.append({
                    "type": movement_type, "id": code, "name": source["name"],
                    "beforeTotal": round(bt, 4), "total": round(at, 4), "totalChange": round(total_change, 4),
                    "beforeOverdue": round(bo, 4), "overdue": round(ao, 4), "overdueChange": round(overdue_change, 4),
                    "beforeOver90": round(b90, 4), "over90": round(a90, 4),
                })
        movements.sort(key=lambda item: (0 if item["type"] in {"SAP 만기경과 진입", "SAP 만기경과 증가"} else 1, -abs(item["overdueChange"]), -abs(item["totalChange"])))

        ks, us = kp["summary"], up["summary"]
        matched = [row for row in customers if row["credit"]]
        unmatched = [row for row in customers if not row["credit"]]
        high_util = [row for row in customers if row["util"] is not None and row["util"] >= 80]
        over_limit = [row for row in customers if row["util"] is not None and row["util"] > 100]
        total_overdue = sum(row["overdue"] for row in customers)
        summary = {
            "currency": "KRW_EQ", "unit": "억원",
            "total": round(ks["total"] + usd_to_eok(us["total"]), 4),
            "gross": round(ks["gross"] + usd_to_eok(us["gross"]), 4),
            "credits": round(ks["credits"] + usd_to_eok(us["credits"]), 4),
            "overdue": round(ks["overdue"] + usd_to_eok(us["overdue"]), 4),
            "over90": round(ks["over90"] + usd_to_eok(us["over90"]), 4),
            "stress7": round(ks["stress7"] + usd_to_eok(us["stress7"]), 4),
            "stress30": round(ks["stress30"] + usd_to_eok(us["stress30"]), 4),
            "customerCount": len(customers),
            "matchedCount": len(matched), "matchedAmount": round(sum(row["total"] for row in matched), 4),
            "matchedOverdue": round(sum(row["overdue"] for row in matched), 4), "matchedOver90": round(sum(row["over90"] for row in matched), 4),
            "unmatchedCount": len(unmatched), "unmatchedAmount": round(sum(row["total"] for row in unmatched), 4),
            "unmatchedOverdue": round(sum(row["overdue"] for row in unmatched), 4), "unmatchedOver90": round(sum(row["over90"] for row in unmatched), 4),
            "unmatchedOverdueShare": round(sum(row["overdue"] for row in unmatched) / total_overdue * 100, 1) if total_overdue else 0,
            "highUtilCount": len(high_util), "highUtilAmount": round(sum(row["total"] for row in high_util), 4),
            "overLimitCount": len(over_limit), "overLimitAmount": round(sum(row["total"] for row in over_limit), 4),
            "rawRows": max(ks["rawRows"], us["rawRows"]), "includedRows": max(ks["includedRows"], us["includedRows"]),
            "currencyRows": ks["currencyRows"] + us["currencyRows"], "excludedRows": max(ks["excludedRows"], us["excludedRows"]),
            "negativeRows": ks["negativeRows"] + us["negativeRows"],
            "negativeAmount": round(ks["negativeAmount"] + usd_to_eok(us["negativeAmount"]), 4),
            "creditAvailable": is_latest and bool(credit_map),
        }
        high_risk = [row for row in customers if row["grade"] in {"R1", "R2"}]
        slowdown = [
            row for row in customers
            if float((row.get("behavior") or {}).get("speedDecline") or 0) >= 10
        ]
        deteriorating = [
            row for row in customers
            if row["grade"] in {"R1", "R2"}
            or (
                float((row.get("behavior") or {}).get("amountScore") or 0) >= 70
                and sum((
                    float((row.get("behavior") or {}).get("speedDecline") or 0) >= 10,
                    float((row.get("behavior") or {}).get("agedShare") or 0) >= 45,
                    int(row.get("maxOpenDays") or 0) > 270,
                    float((row.get("behavior") or {}).get("transitionRate") or 0) >= 50,
                )) >= 2
            )
        ]
        repayment_base = round(sum(max(0.0, row["repaymentBase"]) for row in customers), 4)
        estimated_repayment = round(sum(max(0.0, row["estimatedRepaymentAmount"]) for row in customers), 4)
        eligible_speed_customers = [row for row in customers if (row.get("behavior") or {}).get("speedAvailable")]
        portfolio_speed = weighted_mean(
            [(row["speed"], max(row["repaymentBase"], 0.0)) for row in eligible_speed_customers],
            0.0,
        )
        portfolio_prior_speed = weighted_mean(
            [(row["speedPrior"], max(row["repaymentBase"], 0.0)) for row in eligible_speed_customers],
            portfolio_speed,
        )
        transition_eligible = sum(max(0.0, row["agingTransitionEligibleAmount"]) for row in customers)
        transitioned = sum(max(0.0, row["agingTransitionedAmount"]) for row in customers)
        portfolio_transition_rate = transitioned / transition_eligible * 100 if transition_eligible else 0.0
        positive_open_amount = sum(max(0.0, row["positiveOpenAmount"]) for row in customers)
        aged90_amount = sum(max(0.0, row["aged90DocumentAmount"]) for row in customers)
        summary.update({
            "highRiskCount": len(high_risk),
            "highRiskAmount": round(sum(max(0.0, row["positiveOpenAmount"]) for row in high_risk), 4),
            "delayAmount": round(aged90_amount, 4),
            "positiveOpenAmount": round(positive_open_amount, 4),
            "slowdownCount": len(slowdown),
            "slowdownAmount": round(sum(max(0.0, row["positiveOpenAmount"]) for row in slowdown), 4),
            "deterioratingCount": len(deteriorating),
            "deterioratingAmount": round(sum(max(0.0, row["total"]) for row in deteriorating), 4),
            "openDocumentCount": sum(row["openDocumentCount"] for row in customers),
            "maxOpenDays": max((int(row.get("maxOpenDays") or 0) for row in customers), default=0),
            "aged90DocumentCount": sum(row["aged90DocumentCount"] for row in customers),
            "aged90DocumentAmount": round(aged90_amount, 4),
            "aged90AmountShare": round(aged90_amount / positive_open_amount * 100, 1) if positive_open_amount else 0.0,
            "aged361DocumentCount": sum(row["aged361DocumentCount"] for row in customers),
            "aged361DocumentAmount": round(sum(max(0.0, row["aged361DocumentAmount"]) for row in customers), 4),
            "aged361CustomerCount": sum(1 for row in customers if row["aged361DocumentAmount"] > 0),
            "reappearedDocumentCount": sum(row["reappearedDocumentCount"] for row in customers),
            "repeatDelayAmount": round(sum(max(0.0, row["aged90DocumentAmount"]) for row in customers), 4),
            "repaymentBase": repayment_base,
            "estimatedRepaymentAmount": estimated_repayment,
            "resolutionSpeed": round(portfolio_speed, 1),
            "priorResolutionSpeed": round(portfolio_prior_speed, 1),
            "recoverySpeedDecline": round(max(0.0, portfolio_prior_speed - portfolio_speed), 1),
            "agingTransitionRate": round(portfolio_transition_rate, 1),
            "agingTransitionEligibleAmount": round(transition_eligible, 4),
        })
        if repayment_base > 0 and elapsed_days > 0:
            portfolio_speed_history.append(portfolio_speed)
        kr, ur = kp["rollforward"], up["rollforward"]
        rollforward = {
            key: round(float(kr.get(key, 0)) + usd_to_eok(ur.get(key, 0)), 4)
            for key in (
                "opening", "disappeared", "new", "carriedChange", "closing",
                "repaymentBase", "estimatedRepayment", "disappearedRepayment",
                "partialRepayment", "over90Base", "over90Disappeared",
            )
        }
        rollforward.update({
            "openingDate": kr.get("openingDate") or ur.get("openingDate"),
            "over90DisappearanceRate": round(
                rollforward["over90Disappeared"] / rollforward["over90Base"] * 100,
                1,
            ) if rollforward["over90Base"] else 0.0,
            "anomaly": bool(kr["anomaly"] or ur["anomaly"]),
        })
        forecast = {
            "horizons": [
                {
                    "days": kh["days"],
                    **{field: round(kh[field] + usd_to_eok(uh[field]), 4) for field in ("conservative", "base", "optimistic")},
                }
                for kh, uh in zip(kp["forecast"]["horizons"], up["forecast"]["horizons"])
            ],
            "rates": [],
            "method": kp["forecast"]["method"],
        }
        usd_rates = {row["label"]: row for row in up["forecast"]["rates"]}
        for row in kp["forecast"]["rates"]:
            usd_row = usd_rates.get(row["label"], {"balance": 0, "rate": 0, "sampleCount": 0})
            balance = row["balance"] + usd_to_eok(usd_row["balance"])
            expected = row["balance"] * row["rate"] / 100 + usd_to_eok(usd_row["balance"]) * usd_row["rate"] / 100
            forecast["rates"].append({
                "label": row["label"],
                "balance": round(balance, 4),
                "rate": round(expected / balance * 100, 1) if balance else 0,
                "sampleCount": int(row.get("sampleCount") or 0) + int(usd_row.get("sampleCount") or 0),
            })
        quality = {
            "missingDueDate": kp["quality"]["missingDueDate"] + up["quality"]["missingDueDate"],
            "missingInvoiceDate": kp["quality"].get("missingInvoiceDate", 0) + up["quality"].get("missingInvoiceDate", 0),
            "missingCustomerName": kp["quality"]["missingCustomerName"] + up["quality"]["missingCustomerName"],
            "duplicateSuspects": kp["quality"]["duplicateSuspects"] + up["quality"]["duplicateSuspects"],
            "rowCount": kp["quality"]["rowCount"] + up["quality"]["rowCount"],
            "amountDifference": round(kp["quality"]["amountDifference"] + usd_to_eok(up["quality"]["amountDifference"]), 4),
        }
        kb = {row["key"]: row for row in kp["allowance"]["buckets"]}
        ub = {row["key"]: row for row in up["allowance"]["buckets"]}
        allowance_buckets = []
        for key in kb:
            krw_bucket, usd_bucket = kb[key], ub[key]
            allowance_buckets.append({
                **{field: krw_bucket[field] for field in ("key", "label", "appliedRate", "rawRate", "method", "risk")},
                "exposure": round(krw_bucket["exposure"] + allowance_usd_to_eok(usd_bucket["exposure"]), 4),
                "expectedLoss": round(krw_bucket["expectedLoss"] + allowance_usd_to_eok(usd_bucket["expectedLoss"]), 4),
                "customerCount": krw_bucket["customerCount"] + usd_bucket["customerCount"],
            })
        if is_model_date:
            baseline = {row["key"]: row for row in allowance_model.get("baselineBuckets", [])}
            for row in allowance_buckets:
                baseline_row = baseline.get(row["key"])
                if baseline_row:
                    row["exposure"] = round(float(baseline_row["exposureKRW"]) / 100_000_000, 4)
                    row["expectedLoss"] = round(float(baseline_row["expectedLossKRW"]) / 100_000_000, 4)
        allowance_total = sum(row["exposure"] for row in allowance_buckets)
        allowance_loss = sum(row["expectedLoss"] for row in allowance_buckets)
        scope_excluded_exposure = round(
            kp["allowance"]["scopeExcludedExposure"]
            + allowance_usd_to_eok(up["allowance"]["scopeExcludedExposure"]),
            4,
        )
        if is_model_date:
            allowance_loss = float(allowance_model["adjustedAllowance"]) / 100_000_000
            scope_excluded_exposure = round(
                float(allowance_model.get("baselineExcludedExposureKRW") or 0) / 100_000_000,
                4,
            )
        notices = [
            {"id": row["id"], "grade": row["grade"], "name": row["name"], "score": row["riskScore"], "amount": row["total"], "detail": row["signal"]}
            for row in customers if row["grade"] in {"R1", "R2"}
        ][:3]
        periods.append({
            "date": period_date, "label": kp["label"], "summary": summary,
            "aging": {"due": merge_amount_rows(kp["aging"]["due"], up["aging"]["due"]), "invoice": merge_amount_rows(kp["aging"]["invoice"], up["aging"]["invoice"]), "behavior": merge_amount_rows(kp["aging"]["behavior"], up["aging"]["behavior"])},
            "customers": customers, "movements": movements[:200], "rollforward": rollforward,
            "maturity": merge_amount_rows(kp["maturity"], up["maturity"]),
            "persistence": merge_amount_rows(kp["persistence"], up["persistence"]),
            "concentration": concentration(customers),
            "forecast": forecast,
            "allowance": {"buckets": allowance_buckets, "totalExposure": round(allowance_total, 4), "expectedLoss": round(allowance_loss, 4), "weightedRate": round(allowance_loss / allowance_total * 100, 2) if allowance_total else 0.0, "over360Exposure": next(row["exposure"] for row in allowance_buckets if row["key"] == "age_361_plus"), "unclassifiedExposure": round(kp["allowance"]["unclassifiedExposure"] + allowance_usd_to_eok(up["allowance"]["unclassifiedExposure"]), 4), "scopeExcludedExposure": scope_excluded_exposure, "otherReceivablesExposure": round(kp["allowance"]["otherReceivablesExposure"] + allowance_usd_to_eok(up["allowance"]["otherReceivablesExposure"]), 4), "scopeExcludedCustomerCodes": allowance_model.get("excludedCustomerCodes", []), "signedBasis": True, "calculationBasis": "확정모형 재현" if is_model_date else "동일모형 추정", "fxRate": allowance_fx_rate, "fxBasis": "대손시트 확정 외화평가" if is_model_date else "최신 보고기준일 관리환율"},
            "notices": notices,
            "coverage": {"supported": is_latest and bool(credit_map), "matchedOverdue": summary["matchedOverdue"], "unmatchedOverdue": summary["unmatchedOverdue"], "unmatchedOverdueShare": summary["unmatchedOverdueShare"]},
            "quality": quality,
        })
        prior_customers = current_by_code

    trend = [{"month": datetime.strptime(period["date"], "%Y-%m-%d").strftime("%m.%d"), "date": period["date"], "total": period["summary"]["total"], "overdue": period["summary"]["overdue"], "over90": period["summary"]["over90"]} for period in periods]
    latest = periods[-1]
    distribution = {"50% 미만": 0, "50–80%": 0, "80–100%": 0, "100% 초과": 0}
    for customer in latest["customers"]:
        util = customer["util"]
        if util is None:
            continue
        if util < 50: distribution["50% 미만"] += 1
        elif util < 80: distribution["50–80%"] += 1
        elif util <= 100: distribution["80–100%"] += 1
        else: distribution["100% 초과"] += 1
    return {
        "currency": "KRW_EQ", "unit": "억원", "fxRate": usd_krw_rate,
        "summary": latest["summary"], "customers": latest["customers"], "aging": latest["aging"],
        "trend": trend, "periods": periods, "movements": latest["movements"], "quality": latest["quality"],
        "creditDistribution": [{"label": label, "value": value, "risk": label in {"80–100%", "100% 초과"}} for label, value in distribution.items()],
    }


def select_display_periods(
    periods: list[dict[str, Any]],
    calendar_months: int = 12,
) -> list[dict[str, Any]]:
    """Return bounded month-end boards while retaining full-history analytics.

    High-frequency snapshots remain in SQLite and in the lightweight trend
    series.  The standalone HTML only needs one selectable customer board per
    calendar month, with the latest snapshot always taking precedence.
    """
    if not periods:
        return []
    latest_day = datetime.strptime(periods[-1]["date"], "%Y-%m-%d").date()
    latest_month_index = latest_day.year * 12 + latest_day.month - 1
    earliest_month_index = latest_month_index - max(1, calendar_months) + 1
    month_end_periods: dict[str, dict[str, Any]] = {}
    for period in periods:
        period_day = datetime.strptime(period["date"], "%Y-%m-%d").date()
        month_index = period_day.year * 12 + period_day.month - 1
        if month_index < earliest_month_index:
            continue
        month_end_periods[period["date"][:7]] = period
    selected = sorted(month_end_periods.values(), key=lambda period: period["date"])
    if not selected or selected[-1]["date"] != periods[-1]["date"]:
        selected.append(periods[-1])
    return selected


def advanced_dashboard_data(
    connection: sqlite3.Connection,
    usd_krw_rate: float,
    rate_basis: str,
    rate_date: str,
    rate_quote_date: str,
    allowance_model: dict[str, Any],
) -> dict[str, Any]:
    """Build a currency-safe, decision-oriented static dashboard payload."""
    summary_rows = rows_as_dicts(
        connection.execute(
            "SELECT * FROM currency_summary ORDER BY currency, snapshot_date"
        )
    )
    if not summary_rows:
        raise ValueError("적재된 통화별 Aging 스냅샷이 없습니다.")

    imports = rows_as_dicts(
        connection.execute(
            "SELECT snapshot_date,source_type,source_filename,raw_rows,included_rows "
            "FROM import_batch ORDER BY snapshot_date,source_type"
        )
    )
    aging_imports = {
        row["snapshot_date"]: row for row in imports if row["source_type"] == "aging"
    }
    latest_date = max(row["snapshot_date"] for row in summary_rows)
    credit_date_row = connection.execute(
        "SELECT MAX(snapshot_date) FROM credit_snapshot"
    ).fetchone()
    credit_date = credit_date_row[0] if credit_date_row else None
    credit_import = next(
        (
            row for row in reversed(imports)
            if row["source_type"] == "credit" and row["snapshot_date"] == credit_date
        ),
        None,
    )
    credit_rows = rows_as_dicts(
        connection.execute(
            """
            SELECT customer_code, MAX(customer_name) AS customer_name,
                   SUM(credit_limit) AS credit_limit,
                   SUM(credit_used) AS credit_used,
                   SUM(credit_available) AS credit_available
            FROM credit_snapshot
            WHERE snapshot_date=?
            GROUP BY customer_code
            """,
            (credit_date,),
        )
    ) if credit_date else []
    credit_map = {row["customer_code"]: row for row in credit_rows}

    currency_data: dict[str, dict[str, Any]] = {}
    currencies = sorted({row["currency"] for row in summary_rows}, key=lambda c: (c != "KRW", c))
    grade_order = {"R1": 0, "R2": 1, "R3": 2, "R4": 3}
    bucket_labels = {
        "age_0_30": "송장경과 30일 이하",
        "age_31_60": "송장경과 31–60일",
        "age_61_90": "송장경과 61–90일",
        "age_91_180": "송장경과 91–180일",
        "age_181_360": "송장경과 181–360일",
        "age_361_plus": "송장경과 361일 이상",
        "unclassified": "송장일 미분류",
    }

    for currency in currencies:
        currency_summaries = [row for row in summary_rows if row["currency"] == currency]
        dates = [row["snapshot_date"] for row in currency_summaries]
        summary_by_date = {row["snapshot_date"]: row for row in currency_summaries}
        customer_rows = rows_as_dicts(
            connection.execute(
                "SELECT * FROM customer_currency_summary WHERE currency=? "
                "ORDER BY snapshot_date, customer_code",
                (currency,),
            )
        )
        customers_by_date: dict[str, list[dict[str, Any]]] = {}
        for row in customer_rows:
            customers_by_date.setdefault(row["snapshot_date"], []).append(row)

        raw_rows = rows_as_dicts(
            connection.execute(
                """
                SELECT snapshot_date,source_row,customer_code,customer_name,document_no,
                       invoice_date,due_date,days_overdue,aging_days_source,due_bucket,
                       scope,amount,payment_terms,recon_account
                FROM ar_snapshot
                WHERE CASE WHEN TRIM(COALESCE(currency,''))='' THEN '미지정' ELSE TRIM(currency) END=?
                ORDER BY snapshot_date,customer_code,document_no,source_row
                """,
                (currency,),
            )
        )
        raw_by_date: dict[str, list[dict[str, Any]]] = {}
        for row in raw_rows:
            raw_by_date.setdefault(row["snapshot_date"], []).append(row)

        def document_map(period_date: str) -> dict[str, dict[str, Any]]:
            result: dict[str, dict[str, Any]] = {}
            snapshot_day = datetime.strptime(period_date, "%Y-%m-%d").date()
            for row in raw_by_date.get(period_date, []):
                document_id = row["document_no"] or f"ROW-{row['source_row']}"
                # SAP document numbers may repeat across years.  The source does not
                # expose company code/fiscal year/line item, so invoice date and the
                # reconciliation account are included in the best available stable
                # business key.  Amount is intentionally excluded because partial
                # clearing changes the balance of the same document.
                key = "|".join(
                    (
                        row["customer_code"],
                        document_id,
                        row["invoice_date"] or "",
                        row["recon_account"] or row["scope"],
                    )
                )
                invoice_date_value = row["invoice_date"]
                invoice_age = None
                if invoice_date_value:
                    invoice_age = (snapshot_day - datetime.strptime(invoice_date_value, "%Y-%m-%d").date()).days
                elif row["aging_days_source"] is not None:
                    invoice_age = int(float(row["aging_days_source"]))
                item = result.setdefault(
                    key,
                    {
                        "id": document_id,
                        "customer": row["customer_code"],
                        "name": row["customer_name"],
                        "scope": row["scope"],
                        "invoiceDate": invoice_date_value,
                        "invoiceAge": invoice_age,
                        "dueDate": row["due_date"],
                        "days": row["days_overdue"],
                        "bucket": row["due_bucket"],
                        "terms": row["payment_terms"],
                        "documentKey": key,
                        "amount": 0.0,
                    },
                )
                item["amount"] += float(row["amount"] or 0)
                if invoice_age is not None:
                    item["invoiceAge"] = max(item["invoiceAge"] if item["invoiceAge"] is not None else invoice_age, invoice_age)
                if row["days_overdue"] is not None:
                    item["days"] = max(item["days"] or row["days_overdue"], row["days_overdue"])
            return result

        doc_maps = {period_date: document_map(period_date) for period_date in dates}
        risk_doc_maps = {
            period_date: net_open_documents(doc_maps.get(period_date, {}))
            for period_date in dates
        }
        overdue_counts: dict[str, int] = {}
        disappearance_history: dict[str, list[float]] = {key: [] for key in bucket_labels}
        periods: list[dict[str, Any]] = []
        prior_customer_map: dict[str, dict[str, Any]] = {}
        prior_docs: dict[str, dict[str, Any]] = {}

        for date_index, period_date in enumerate(dates):
            period_summary = summary_by_date[period_date]
            period_rows = customers_by_date.get(period_date, [])
            current_map = {row["customer_code"]: row for row in period_rows}
            signed_current_docs = doc_maps[period_date]
            current_docs = risk_doc_maps[period_date]
            is_latest = period_date == latest_date
            credit_supported = currency == "KRW" and is_latest and bool(credit_rows)

            current_overdue_codes = {
                row["customer_code"] for row in period_rows if float(row["overdue_amount"] or 0) > 0
            }
            for code in current_overdue_codes:
                overdue_counts[code] = overdue_counts.get(code, 0) + 1

            if prior_docs:
                elapsed = max(1, (datetime.strptime(period_date, "%Y-%m-%d") - datetime.strptime(dates[date_index - 1], "%Y-%m-%d")).days)
                for bucket in bucket_labels:
                    eligible = {
                        key: item for key, item in prior_docs.items()
                        if repayment_age_bucket(item.get("invoiceAge")) == bucket and item["amount"] > 0
                    }
                    base = sum(item["amount"] for item in eligible.values())
                    vanished = sum(item["amount"] for key, item in eligible.items() if key not in current_docs)
                    partial = sum(
                        max(0.0, item["amount"] - float(current_docs[key].get("amount") or 0))
                        for key, item in eligible.items()
                        if key in current_docs and float(current_docs[key].get("amount") or 0) > 0
                    )
                    if base > 0:
                        normalized = normalized_repayment_rate(vanished + partial, base, elapsed)
                        disappearance_history[bucket].append(normalized)

            customer_documents: dict[str, list[dict[str, Any]]] = {}
            for item in current_docs.values():
                customer_documents.setdefault(item["customer"], []).append(item)

            matched_count = high_util_count = over_limit_count = 0
            matched_total = unmatched_total = high_util_total = over_limit_total = 0.0
            matched_overdue = unmatched_overdue = matched_over90 = unmatched_over90 = 0.0
            credit_distribution = {"50% 미만": 0, "50–80%": 0, "80–100%": 0, "100% 초과": 0}
            customer_payload: list[dict[str, Any]] = []
            for row in period_rows:
                code = row["customer_code"]
                total = float(row["total_amount"] or 0)
                overdue = float(row["overdue_amount"] or 0)
                over90 = float(row["over90_amount"] or 0)
                gross = float(row["gross_debit"] or 0)
                credits = float(row["credit_amount"] or 0)
                previous = prior_customer_map.get(code)
                previous_total = float(previous["total_amount"] or 0) if previous else 0.0
                change = total - previous_total
                persistence = overdue_counts.get(code, 0) if overdue > 0 else 0
                if overdue <= 0:
                    risk_type = "낮음"
                elif persistence >= 4:
                    risk_type = "만성"
                elif persistence == 1:
                    risk_type = "급성"
                else:
                    risk_type = "진행"

                credit_row = credit_map.get(code) if credit_supported else None
                limit_value = used_value = util = None
                has_credit = credit_row is not None
                if has_credit:
                    matched_count += 1
                    matched_total += total
                    matched_overdue += overdue
                    matched_over90 += over90
                    limit_value = float(credit_row["credit_limit"] or 0)
                    used_value = float(credit_row["credit_used"] or 0)
                    util = used_value / limit_value * 100 if limit_value > 0 else (999.0 if used_value > 0 else 0.0)
                    if util < 50:
                        credit_distribution["50% 미만"] += 1
                    elif util < 80:
                        credit_distribution["50–80%"] += 1
                    elif util <= 100:
                        credit_distribution["80–100%"] += 1
                    else:
                        credit_distribution["100% 초과"] += 1
                    if util >= 80:
                        high_util_count += 1
                        high_util_total += total
                    if util > 100:
                        over_limit_count += 1
                        over_limit_total += total
                elif credit_supported:
                    unmatched_total += total
                    unmatched_overdue += overdue
                    unmatched_over90 += over90

                over90_threshold = 500_000_000 if currency == "KRW" else 500_000
                overdue_threshold = 1_000_000_000 if currency == "KRW" else 1_000_000
                medium_threshold = 200_000_000 if currency == "KRW" else 200_000
                change_threshold = 500_000_000 if currency == "KRW" else 500_000
                if over90 >= over90_threshold or overdue >= overdue_threshold or (util is not None and util > 100):
                    grade = "R1"
                elif over90 > 0 or overdue >= medium_threshold:
                    grade = "R2"
                elif overdue > 0 or (util is not None and util >= 80):
                    grade = "R3"
                else:
                    grade = "R4"
                signals: list[str] = []
                if risk_type == "만성":
                    signals.append(f"만성 연체 {persistence}회")
                elif risk_type == "급성":
                    signals.append("신규 급성 연체")
                elif overdue > 0:
                    signals.append("연체 진행")
                if over90 > 0:
                    signals.append("90일 초과")
                if util is not None and util > 100:
                    signals.append("여신한도 초과")
                elif util is not None and util >= 80:
                    signals.append("여신 고사용률")
                if credit_supported and not has_credit:
                    signals.append("여신정보 없음")
                if abs(change) >= change_threshold:
                    signals.append("직전 대비 급변")
                if not signals:
                    signals.append("정상 모니터링")

                document_limit = 15 if is_latest else 0
                documents = sorted(
                    customer_documents.get(code, []),
                    key=lambda item: (-(item.get("invoiceAge") if item.get("invoiceAge") is not None else -99999), -abs(item["amount"])),
                )[:document_limit]
                customer_payload.append(
                    {
                        "id": code,
                        "name": row["customer_name"],
                        "grade": grade,
                        "riskType": risk_type,
                        "persistence": persistence,
                        "total": display_amount(total, currency),
                        "gross": display_amount(gross, currency),
                        "credits": display_amount(credits, currency),
                        "overdue": display_amount(overdue, currency),
                        "over90": display_amount(over90, currency),
                        "limit": None if limit_value is None else display_amount(limit_value, currency),
                        "used": None if used_value is None else display_amount(used_value, currency),
                        "util": None if util is None else round(util, 1),
                        "credit": has_credit,
                        "scope": row["scope"],
                        "signal": " · ".join(signals[:4]),
                        "action": "재무 중점 점검" if grade in {"R1", "R2"} else "정기 모니터링",
                        "change": display_amount(change, currency),
                        "documents": [
                            {
                                "id": item["id"],
                                "invoiceDate": item["invoiceDate"],
                                "invoiceAge": item["invoiceAge"],
                                "dueDate": item["dueDate"],
                                "days": item["days"],
                                "scope": item["scope"],
                                "amount": display_amount(item["amount"], currency),
                                "terms": item["terms"],
                            }
                            for item in documents
                        ],
                    }
                )

            customer_payload.sort(
                key=lambda item: (grade_order[item["grade"]], -item["over90"], -item["overdue"], -item["total"])
            )

            movements: list[dict[str, Any]] = []
            for code in set(prior_customer_map) | set(current_map):
                before = prior_customer_map.get(code)
                after = current_map.get(code)
                bt = float(before["total_amount"] or 0) if before else 0.0
                at = float(after["total_amount"] or 0) if after else 0.0
                bo = float(before["overdue_amount"] or 0) if before else 0.0
                ao = float(after["overdue_amount"] or 0) if after else 0.0
                b90 = float(before["over90_amount"] or 0) if before else 0.0
                a90 = float(after["over90_amount"] or 0) if after else 0.0
                overdue_change = ao - bo
                total_change = at - bt
                movement_type = None
                threshold = 50_000_000 if currency == "KRW" else 50_000
                total_threshold = 500_000_000 if currency == "KRW" else 500_000
                if bo <= 0 < ao:
                    movement_type = "SAP 만기경과 진입"
                elif bo > 0 >= ao:
                    movement_type = "SAP 만기경과 해소"
                elif overdue_change >= threshold:
                    movement_type = "SAP 만기경과 증가"
                elif overdue_change <= -threshold:
                    movement_type = "SAP 만기경과 감소"
                elif total_change >= total_threshold:
                    movement_type = "채권 증가"
                elif total_change <= -total_threshold:
                    movement_type = "채권 감소"
                if movement_type:
                    source = after or before
                    movements.append(
                        {
                            "type": movement_type,
                            "id": code,
                            "name": source["customer_name"],
                            "beforeTotal": display_amount(bt, currency),
                            "total": display_amount(at, currency),
                            "totalChange": display_amount(total_change, currency),
                            "beforeOverdue": display_amount(bo, currency),
                            "overdue": display_amount(ao, currency),
                            "overdueChange": display_amount(overdue_change, currency),
                            "beforeOver90": display_amount(b90, currency),
                            "over90": display_amount(a90, currency),
                        }
                    )
            movements.sort(key=lambda item: (0 if item["type"] in {"SAP 만기경과 진입", "SAP 만기경과 증가"} else 1, -abs(item["overdueChange"]), -abs(item["totalChange"])))

            period_day = datetime.strptime(period_date, "%Y-%m-%d").date()
            prior_month_end = period_day.replace(day=1) - timedelta(days=1)
            prior_month_start = prior_month_end.replace(day=1)
            rollforward_base_dates = [
                candidate
                for candidate in dates[:date_index]
                if prior_month_start
                <= datetime.strptime(candidate, "%Y-%m-%d").date()
                <= prior_month_end
            ]
            rollforward_base_date = max(rollforward_base_dates, default=None)
            rollforward_prior_docs = risk_doc_maps.get(rollforward_base_date, {}) if rollforward_base_date else {}
            prior_positive = {
                key: item
                for key, item in rollforward_prior_docs.items()
                if float(item.get("amount") or 0) > 0
            }
            current_positive = {
                key: item for key, item in current_docs.items() if float(item.get("amount") or 0) > 0
            }
            opening = sum(item["amount"] for item in prior_positive.values())
            closing = sum(item["amount"] for item in current_positive.values())
            common_keys = set(prior_positive) & set(current_positive)
            disappeared = sum(prior_positive[key]["amount"] for key in set(prior_positive) - set(current_positive))
            new_amount = sum(current_positive[key]["amount"] for key in set(current_positive) - set(prior_positive))
            carried_change = sum(current_positive[key]["amount"] - prior_positive[key]["amount"] for key in common_keys)
            repayment_base, estimated_repayment, disappeared_repayment, partial_repayment = (
                estimated_repayment_components(prior_positive, current_positive)
            )
            over90_open = sum(item["amount"] for item in prior_positive.values() if (item.get("invoiceAge") or 0) > 90)
            over90_disappeared = sum(item["amount"] for key, item in prior_positive.items() if key not in current_positive and (item.get("invoiceAge") or 0) > 90)
            over90_disappearance_rate = over90_disappeared / over90_open * 100 if over90_open else 0.0
            rollforward = {
                "openingDate": rollforward_base_date,
                "opening": display_amount(opening, currency),
                "disappeared": display_amount(disappeared, currency),
                "new": display_amount(new_amount, currency),
                "carriedChange": display_amount(carried_change, currency),
                "closing": display_amount(closing, currency),
                "repaymentBase": display_amount(repayment_base, currency),
                "estimatedRepayment": display_amount(estimated_repayment, currency),
                "disappearedRepayment": display_amount(disappeared_repayment, currency),
                "partialRepayment": display_amount(partial_repayment, currency),
                "over90Base": display_amount(over90_open, currency),
                "over90Disappeared": display_amount(over90_disappeared, currency),
                "over90DisappearanceRate": round(over90_disappearance_rate, 1),
                "anomaly": over90_disappearance_rate >= 50 and over90_open > 0,
            }

            maturity_raw = {
                "만기경과": sum(item["amount"] for item in current_docs.values() if (item["days"] or 0) > 0),
                "당일": sum(item["amount"] for item in current_docs.values() if item["days"] == 0),
                "1–7일": sum(item["amount"] for item in current_docs.values() if item["days"] is not None and -7 <= item["days"] <= -1),
                "8–14일": sum(item["amount"] for item in current_docs.values() if item["days"] is not None and -14 <= item["days"] <= -8),
                "15–30일": sum(item["amount"] for item in current_docs.values() if item["days"] is not None and -30 <= item["days"] <= -15),
                "31일+": sum(item["amount"] for item in current_docs.values() if item["days"] is not None and item["days"] < -30),
            }
            maturity = [{"label": label, "value": display_amount(value, currency), "risk": label in {"만기경과", "당일", "1–7일"}} for label, value in maturity_raw.items()]
            stress7 = maturity_raw["만기경과"] + maturity_raw["당일"] + maturity_raw["1–7일"]
            stress30 = stress7 + maturity_raw["8–14일"] + maturity_raw["15–30일"]

            persistence_groups = {"급성(1회)": 0.0, "진행(2–3회)": 0.0, "만성(4–6회)": 0.0, "고착(7회+)": 0.0}
            for customer in customer_payload:
                if customer["overdue"] <= 0:
                    continue
                raw_value = customer["overdue"]
                if customer["persistence"] == 1:
                    persistence_groups["급성(1회)"] += raw_value
                elif customer["persistence"] <= 3:
                    persistence_groups["진행(2–3회)"] += raw_value
                elif customer["persistence"] <= 6:
                    persistence_groups["만성(4–6회)"] += raw_value
                else:
                    persistence_groups["고착(7회+)"] += raw_value

            top_total = sorted((max(0.0, float(row["total_amount"] or 0)) for row in period_rows), reverse=True)
            top_overdue = sorted((max(0.0, float(row["overdue_amount"] or 0)) for row in period_rows), reverse=True)
            top_over90 = sorted((max(0.0, float(row["over90_amount"] or 0)) for row in period_rows), reverse=True)
            total_positive = sum(top_total)
            overdue_positive = sum(top_overdue)
            over90_positive = sum(top_over90)
            concentration = {
                "top1Total": round((top_total[0] / total_positive * 100) if total_positive and top_total else 0, 1),
                "top5Total": round(sum(top_total[:5]) / total_positive * 100 if total_positive else 0, 1),
                "top5Overdue": round(sum(top_overdue[:5]) / overdue_positive * 100 if overdue_positive else 0, 1),
                "top5Over90": round(sum(top_over90[:5]) / over90_positive * 100 if over90_positive else 0, 1),
            }

            rates = {
                bucket: (statistics.median(values) if values else 0.0)
                for bucket, values in disappearance_history.items()
            }
            positive_by_bucket = {
                bucket: sum(
                    item["amount"] for item in current_docs.values()
                    if repayment_age_bucket(item.get("invoiceAge")) == bucket and item["amount"] > 0
                )
                for bucket in bucket_labels
            }
            def scenario_amount(days_forward: int, factor: float) -> float:
                expected = 0.0
                for bucket, balance in positive_by_bucket.items():
                    p30 = min(0.95, rates[bucket] * factor)
                    probability = 1 - (1 - p30) ** (days_forward / 30)
                    expected += balance * probability
                return display_amount(expected, currency)
            forecast = {
                "rates": [
                    {
                        "label": bucket_labels[bucket],
                        "rate": round(rate * 100, 1),
                        "balance": display_amount(positive_by_bucket[bucket], currency),
                        "sampleCount": len(disappearance_history[bucket]),
                    }
                    for bucket, rate in rates.items()
                ],
                "horizons": [
                    {
                        "days": horizon,
                        "conservative": scenario_amount(horizon, 0.75),
                        "base": scenario_amount(horizon, 1.0),
                        "optimistic": scenario_amount(horizon, 1.25),
                    }
                    for horizon in (7, 14, 30)
                ],
                "method": "동일 송장경과 구간의 과거 채권회수율 중앙값",
            }

            period_import = aging_imports[period_date]
            currency_row_count = len(raw_by_date.get(period_date, []))
            negative_rows = sum(1 for row in raw_by_date.get(period_date, []) if float(row["amount"] or 0) < 0)
            negative_amount = sum(float(row["amount"] or 0) for row in raw_by_date.get(period_date, []) if float(row["amount"] or 0) < 0)
            missing_due = sum(1 for row in raw_by_date.get(period_date, []) if not row["due_date"])
            duplicate_count = connection.execute(
                """
                SELECT COALESCE(SUM(cnt-1),0) FROM (
                    SELECT COUNT(*) AS cnt FROM ar_snapshot
                    WHERE snapshot_date=? AND CASE WHEN TRIM(COALESCE(currency,''))='' THEN '미지정' ELSE TRIM(currency) END=?
                    GROUP BY customer_code,document_no,due_date,amount HAVING COUNT(*)>1
                )
                """,
                (period_date, currency),
            ).fetchone()[0]
            period_summary_payload = {
                "currency": currency,
                "unit": "억원" if currency == "KRW" else "백만 달러",
                "total": display_amount(period_summary["total_amount"], currency),
                "gross": display_amount(sum(max(0.0, item["amount"]) for item in signed_current_docs.values()), currency),
                "credits": display_amount(sum(min(0.0, item["amount"]) for item in signed_current_docs.values()), currency),
                "overdue": display_amount(period_summary["overdue_amount"], currency),
                "over90": display_amount(period_summary["over90_amount"], currency),
                "customerCount": int(period_summary["customer_count"]),
                "highUtilCount": high_util_count,
                "highUtilAmount": display_amount(high_util_total, currency),
                "overLimitCount": over_limit_count,
                "overLimitAmount": display_amount(over_limit_total, currency),
                "matchedCount": matched_count,
                "matchedAmount": display_amount(matched_total, currency),
                "matchedOverdue": display_amount(matched_overdue, currency),
                "matchedOver90": display_amount(matched_over90, currency),
                "unmatchedCount": len(period_rows) - matched_count if credit_supported else 0,
                "unmatchedAmount": display_amount(unmatched_total, currency),
                "unmatchedOverdue": display_amount(unmatched_overdue, currency),
                "unmatchedOver90": display_amount(unmatched_over90, currency),
                "unmatchedOverdueShare": round(unmatched_overdue / float(period_summary["overdue_amount"] or 1) * 100, 1) if credit_supported else None,
                "rawRows": int(period_import["raw_rows"]),
                "includedRows": int(period_import["included_rows"]),
                "currencyRows": currency_row_count,
                "excludedRows": int(period_import["raw_rows"] - period_import["included_rows"]),
                "negativeRows": negative_rows,
                "negativeAmount": display_amount(negative_amount, currency),
                "creditAvailable": credit_supported,
                "stress7": display_amount(stress7, currency),
                "stress30": display_amount(stress30, currency),
            }
            period_aging = {
                "due": [
                    {"label": "만기 미도래", "value": display_amount(period_summary["not_due_amount"], currency), "risk": False},
                    {"label": "당일 만기", "value": display_amount(period_summary["due_today_amount"], currency), "risk": False},
                    {"label": "1–30일 만기경과", "value": display_amount(period_summary["overdue_1_30"], currency), "risk": True},
                    {"label": "31–60일 만기경과", "value": display_amount(period_summary["overdue_31_60"], currency), "risk": True},
                    {"label": "61–90일 만기경과", "value": display_amount(period_summary["overdue_61_90"], currency), "risk": True},
                    {"label": "90일 초과 만기경과", "value": display_amount(period_summary["over90_amount"], currency), "risk": True},
                ],
                "invoice": [
                    {"label": "송장 1–30일", "value": display_amount(period_summary["invoice_1_30"], currency), "risk": False},
                    {"label": "송장 31–60일", "value": display_amount(period_summary["invoice_31_60"], currency), "risk": False},
                    {"label": "송장 61–90일", "value": display_amount(period_summary["invoice_61_90"], currency), "risk": False},
                    {"label": "송장 91일 이상", "value": display_amount(period_summary["invoice_91_plus"], currency), "risk": True},
                ],
            }
            periods.append(
                {
                    "date": period_date,
                    "label": datetime.strptime(period_date, "%Y-%m-%d").strftime("%Y.%m.%d"),
                    "summary": period_summary_payload,
                    "aging": period_aging,
                    "customers": customer_payload,
                    "movements": movements[:200],
                    "rollforward": rollforward,
                    "maturity": maturity,
                    "persistence": [{"label": label, "value": round(value, 4), "risk": "만성" in label or "고착" in label} for label, value in persistence_groups.items()],
                    "concentration": concentration,
                    "forecast": forecast,
                    "coverage": {
                        "supported": credit_supported,
                        "matchedOverdue": display_amount(matched_overdue, currency),
                        "unmatchedOverdue": display_amount(unmatched_overdue, currency),
                        "unmatchedOverdueShare": period_summary_payload["unmatchedOverdueShare"],
                    },
                    "quality": {
                        "missingDueDate": missing_due,
                        "missingInvoiceDate": sum(1 for row in raw_by_date.get(period_date, []) if not row["invoice_date"]),
                        "missingCustomerName": sum(1 for row in raw_by_date.get(period_date, []) if not str(row["customer_name"] or "").strip()),
                        "duplicateSuspects": int(duplicate_count or 0),
                        "rowCount": currency_row_count,
                        "amountDifference": display_amount(sum(item["amount"] for item in signed_current_docs.values()) - float(period_summary["total_amount"] or 0), currency),
                    },
                }
            )
            prior_customer_map = current_map
            prior_docs = current_docs

        enrich_native_periods(periods, doc_maps, dates, currency, allowance_model, usd_krw_rate)

        trend = [
            {
                "month": datetime.strptime(period["date"], "%Y-%m-%d").strftime("%m.%d"),
                "date": period["date"],
                "total": period["summary"]["total"],
                "overdue": period["summary"]["overdue"],
                "over90": period["summary"]["over90"],
            }
            for period in periods
        ]
        latest_period = periods[-1]
        currency_data[currency] = {
            "currency": currency,
            "unit": latest_period["summary"]["unit"],
            "summary": latest_period["summary"],
            "customers": latest_period["customers"],
            "aging": latest_period["aging"],
            "trend": trend,
            "periods": periods,
            "movements": latest_period["movements"],
            "quality": latest_period["quality"],
            "creditDistribution": [
                {"label": label, "value": value, "risk": label in {"80–100%", "100% 초과"}}
                for label, value in credit_distribution.items()
            ],
        }

    if "KRW" in currency_data and "USD" in currency_data:
        currency_data["KRW_EQ"] = combine_currency_views(
            currency_data["KRW"], currency_data["USD"], usd_krw_rate, credit_map,
            allowance_model,
        )
    # Keep only fields consumed by the standalone report.  This avoids repeating
    # internal calculation intermediates across every historical snapshot.
    unused_customer_fields = {
        "persistence", "gross", "credits", "overdue", "over90", "action",
        "delayAmountChange", "delayAmountChangePct", "durationDelta",
        "durationMixDelta", "speedBase", "resolvedAmount", "delayFrequency",
        "delayEligible", "delayAmount", "delayDays", "weightedObservedDays",
        "aged90DocumentShare", "reappearedDocumentCount",
        "allowanceScopeExcluded",
        "otherReceivablesExposure", "subsidiaryCandidate",
        "_openAgeBasis",
    }
    compact_behavior_fields = {
        "amountScore", "shareScore", "durationScore", "speedScore",
        "transitionScore", "speedDecline", "transitionRate",
        "speedAvailable", "transitionAvailable", "transitionStatus", "coverage",
        "agedAmount", "collectionRiskScore", "creditPressureScore",
        "creditAdjustment", "creditAvailable", "creditMateriality",
        "averageOpenDays",
    }
    combined_detail_available = "KRW_EQ" in currency_data
    for currency_code, view in currency_data.items():
        for period_index, period in enumerate(view["periods"]):
            is_latest_period = period_index == len(view["periods"]) - 1
            for customer in period["customers"]:
                for field in unused_customer_fields:
                    customer.pop(field, None)
                behavior = customer.get("behavior") or {}
                customer["behavior"] = {
                    key: value for key, value in behavior.items() if key in compact_behavior_fields
                }
                if not customer["behavior"].get("creditAvailable"):
                    for field in (
                        "collectionRiskScore", "creditPressureScore", "creditAdjustment",
                        "creditAvailable", "creditMateriality",
                    ):
                        customer["behavior"].pop(field, None)
                if not is_latest_period:
                    for field in (
                        "repaymentBase", "estimatedRepaymentAmount",
                        "disappearedRepaymentAmount", "partialRepaymentAmount",
                        "openAgeBuckets", "averageOpenDayContributors",
                        "agingTransitionEligibleAmount", "agingTransitionedAmount",
                        "agingTransitionDocuments", "agingTransitionEligibleEok",
                        "allowanceOver180", "allowanceOver360",
                    ):
                        customer.pop(field, None)
                if not is_latest_period or (combined_detail_available and currency_code != "KRW_EQ"):
                    customer.pop("documents", None)
                    customer.pop("closedDocuments", None)

    # The database and all calculations retain every snapshot.  The standalone
    # report carries only the boards that a user can select: the latest snapshot
    # in each of the most recent 12 calendar months.  Trend arrays still contain
    # every snapshot, so daily/weekly loading improves analytics without causing
    # customer-level JSON to grow without bound.
    for view in currency_data.values():
        view["periods"] = select_display_periods(view["periods"], calendar_months=12)
        # These values are already present in periods[-1].  Removing the aliases
        # avoids serializing the complete latest customer list twice per currency.
        for duplicate_key in ("summary", "customers", "aging", "movements", "quality"):
            view.pop(duplicate_key, None)
    display_currencies = [code for code in ("KRW_EQ", "KRW", "USD") if code in currency_data]
    default_currency = "KRW_EQ" if "KRW_EQ" in currency_data else ("KRW" if "KRW" in currency_data else currencies[0])
    snapshots = [
        {"date": row["snapshot_date"], "label": "현재 Aging" if row["snapshot_date"] == latest_date else "과거 Aging"}
        for row in sorted(
            {row["snapshot_date"]: row for row in summary_rows}.values(),
            key=lambda row: row["snapshot_date"],
        )
    ]
    history_start = snapshots[0]["date"]
    history_days = (datetime.strptime(latest_date, "%Y-%m-%d") - datetime.strptime(history_start, "%Y-%m-%d")).days
    history_months = round(history_days / 30.4375, 1)
    if history_days < 180:
        history_stage = "기초 이력"
    elif history_days < 365:
        history_stage = "안정화 중"
    elif history_days < 548:
        history_stage = "12개월 기준 확보"
    else:
        history_stage = "계절성 비교 가능"
    return {
        "meta": {
            "asOf": latest_date,
            "reportDateBasis": "최종 Aging 기준일",
            "creditAsOf": credit_date,
            "creditSourceFile": credit_import["source_filename"] if credit_import else None,
            "creditBasis": "최신 적재 여신현황",
            "generatedAt": datetime.now().astimezone().strftime("%Y-%m-%d %H:%M"),
            "readonly": True,
            "currencyNotice": "원화환산 전체는 보고기준일에 적용되는 서울외국환중개 USD/KRW 매매기준율 사용(휴일은 직전 고시일)",
            "usdKrwRate": usd_krw_rate,
            "rateBasis": rate_basis,
            "rateDate": rate_quote_date,
            "rateReportDate": rate_date,
            "rateSourceUrl": "http://www.smbs.biz/ExRate/StdExRate.jsp",
            "history": {
                "startDate": history_start,
                "endDate": latest_date,
                "days": history_days,
                "months": history_months,
                "snapshotCount": len(snapshots),
                "stage": history_stage,
                "minimumRecommendedMonths": 12,
                "seasonalityRecommendedMonths": 24,
            },
            "payload": {
                "analysisSnapshotCount": len(snapshots),
                "displaySnapshotCount": len(currency_data[default_currency]["periods"]),
                "displayPeriodPolicy": "최근 12개월 월별 마지막 기준일",
                "trendScope": "전체 적재 기준일 집계",
                "customerDetailScope": "최신 기준일",
            },
            "allowanceModel": {
                "file": allowance_model["file"],
                "modelDate": allowance_model["modelDate"],
                "basis": allowance_model["basis"],
                "rawBasis": allowance_model["rawBasis"],
                "previousBasis": allowance_model["previousBasis"],
                "methodNote": allowance_model["methodNote"],
                "transitionSteps": allowance_model["transitionSteps"],
                "buckets": allowance_model["buckets"],
                "modelFxRate": allowance_model["modelFxRate"],
                "excludedCustomerCodes": allowance_model["excludedCustomerCodes"],
                "baselineExcludedExposureKRW": allowance_model["baselineExcludedExposureKRW"],
                "collectiveAllowanceKRW": allowance_model["collectiveAllowance"],
                "individualLongTermAllowanceKRW": allowance_model["individualLongTermAllowance"],
                "adjustedAllowanceKRW": allowance_model["adjustedAllowance"],
                "controlChecks": allowance_model["controlChecks"],
            },
        },
        "currencies": display_currencies,
        "currencyLabels": {"KRW_EQ": "원화환산 전체", "KRW": "원화(KRW)", "USD": "달러(USD)"},
        "defaultCurrency": default_currency,
        "currencyData": currency_data,
        "snapshots": snapshots,
        "imports": [
            {
                "date": row["snapshot_date"],
                "type": row["source_type"],
                "file": row["source_filename"],
                "rawRows": int(row["raw_rows"]),
                "includedRows": int(row["included_rows"]),
            }
            for row in imports
        ],
    }


def write_dashboard(data: dict[str, Any]) -> tuple[Path, Path]:
    template = TEMPLATE_PATH.read_text(encoding="utf-8")
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    replacement = f"/*__DASHBOARD_DATA_START__*/\n    const dashboardData = {payload};\n    /*__DASHBOARD_DATA_END__*/"
    rendered, count = re.subn(
        r"/\*__DASHBOARD_DATA_START__\*/.*?/\*__DASHBOARD_DATA_END__\*/",
        replacement,
        template,
        count=1,
        flags=re.DOTALL,
    )
    if count != 1:
        raise ValueError("대시보드 템플릿의 데이터 삽입 위치를 찾지 못했습니다.")
    # Keep the source template lightweight while refreshing the locally opened copy.
    WORKING_HTML_PATH.write_text(rendered, encoding="utf-8")
    PUBLISH_DIR.mkdir(parents=True, exist_ok=True)
    latest_path = PUBLISH_DIR / "채권리스크대시보드_최신.html"
    dated_path = PUBLISH_DIR / f"채권리스크대시보드_{data['meta']['asOf']}.html"
    latest_path.write_text(rendered, encoding="utf-8")
    dated_path.write_text(rendered, encoding="utf-8")
    return latest_path, dated_path


def backup_database() -> None:
    if not DB_PATH.exists():
        return
    backup_dir = PRIVATE_DIR / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(DB_PATH, backup_dir / f"receivables_{stamp}.db")


def main() -> int:
    args = parse_args()
    source_dir = args.source_dir.expanduser().resolve()
    if not source_dir.exists():
        raise FileNotFoundError(f"SAP 파일 폴더를 찾을 수 없습니다: {source_dir}")
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"대시보드 템플릿을 찾을 수 없습니다: {TEMPLATE_PATH}")

    source_files, discovery_warnings = discover_source_workbooks(source_dir)
    aging_files = source_files["aging"]
    credit_files = source_files["credit"]
    if not aging_files:
        raise FileNotFoundError(f"날짜와 '에이징'이 포함된 XLSX 파일이 없습니다: {source_dir}")
    if args.latest_date_only:
        print(max(source_workbook_date(path, "aging") for path in aging_files).isoformat())
        return 0

    for warning in discovery_warnings:
        print(f"파일 인식 안내: {warning}")

    PRIVATE_DIR.mkdir(parents=True, exist_ok=True)
    temp_db = PRIVATE_DIR / "receivables.building.db"
    if temp_db.exists():
        temp_db.unlink()

    connection = sqlite3.connect(temp_db)
    try:
        allowance_path = find_allowance_workbook(source_dir, args.allowance_file, source_files["allowance"])
        allowance_model = load_allowance_model(allowance_path)
        create_schema(connection)
        aging_results = [import_aging(connection, path) for path in aging_files]
        latest_aging_date = max(result.snapshot_date for result in aging_results)
        if credit_files:
            latest_credit = latest_source_workbook(credit_files, "credit")
            credit_date = source_workbook_date(latest_credit, "credit")
            import_credit(connection, latest_credit, credit_date)
        connection.execute("PRAGMA optimize")
        connection.commit()
        config = load_config()
        configured_rate_date = str(config["rateDate"])
        if configured_rate_date != latest_aging_date.isoformat():
            raise ValueError(
                "최종 보고일의 서울외국환중개 환율이 필요합니다. "
                f"dashboard-settings.json의 rateDate={configured_rate_date}, "
                f"최종 Aging 기준일={latest_aging_date.isoformat()}"
            )
        usd_krw_rate = float(args.usd_krw_rate or config["usdKrwRate"])
        if usd_krw_rate <= 0:
            raise ValueError("USD/KRW 관리환율은 0보다 커야 합니다.")
        rate_quote_date = str(config.get("rateQuoteDate", configured_rate_date))
        quote_date = datetime.strptime(rate_quote_date, "%Y-%m-%d").date()
        quote_lookback_days = (latest_aging_date - quote_date).days
        if quote_lookback_days < 0 or quote_lookback_days > 7:
            raise ValueError(
                "서울외국환중개 환율 고시일은 최종 Aging 기준일과 같거나 "
                f"직전 7일 이내여야 합니다: {rate_quote_date}"
            )
        dashboard_data = advanced_dashboard_data(
            connection,
            usd_krw_rate,
            str(config["rateBasis"]),
            configured_rate_date,
            rate_quote_date,
            allowance_model,
        )
    finally:
        connection.close()

    backup_database()
    os.replace(temp_db, DB_PATH)
    latest_html, dated_html = write_dashboard(dashboard_data)

    default_view = dashboard_data["currencyData"][dashboard_data["defaultCurrency"]]
    default_period = default_view["periods"][-1]
    print(f"Aging snapshots: {len(aging_files)}")
    print(f"As of: {dashboard_data['meta']['asOf']}")
    print(f"Latest credit as of: {dashboard_data['meta'].get('creditAsOf') or 'not loaded'}")
    print(f"Display boards: {len(default_view['periods'])} monthly snapshots")
    print(f"Trade receivables: {default_period['summary']['total']:.4f} EOK KRW")
    print(f"Overdue: {default_period['summary']['overdue']:.4f} EOK KRW")
    print(f"Over 90 days: {default_period['summary']['over90']:.4f} EOK KRW")
    print(f"High-risk customer exposure: {default_period['summary']['highRiskAmount']:.4f} EOK KRW")
    print(f"Estimated allowance: {dashboard_data['currencyData'][dashboard_data['defaultCurrency']]['periods'][-1]['allowance']['expectedLoss']:.4f} {dashboard_data['currencyData'][dashboard_data['defaultCurrency']]['unit']}")
    print("SQLite database updated.")
    print("Read-only dashboards generated.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"오류: {exc}", file=sys.stderr)
        raise
