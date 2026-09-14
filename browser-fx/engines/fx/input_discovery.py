from __future__ import annotations

import re
from pathlib import Path

import openpyxl


LEDGER_NAMES = {"외화계정원장", "환차손익원장"}
PERIOD_RE = re.compile(r"^\d{2}년말$|^\d{2}년\s*\d{1,2}월$")
MONTH_FILE_RE = re.compile(r"(?:^|_)평가_(\d{2})년_(\d{1,2})월(?:\s+RAW)?$", re.IGNORECASE)
YEAR_END_FILE_RE = re.compile(r"(?:^|_)평가_(\d{2})년말(?:\s+RAW)?$", re.IGNORECASE)


def logical_sheet_name(path: Path, actual_sheet_name: str) -> str:
    actual = actual_sheet_name.strip()
    if actual in LEDGER_NAMES or PERIOD_RE.match(actual):
        return actual
    stem = path.stem.strip()
    if "외화계정원장" in stem:
        return "외화계정원장"
    if "환차손익원장" in stem:
        return "환차손익원장"
    if match := MONTH_FILE_RE.search(stem):
        return f"{int(match.group(1)):02d}년 {int(match.group(2))}월"
    if match := YEAR_END_FILE_RE.search(stem):
        return f"{int(match.group(1)):02d}년말"
    return actual


def input_priority(path: Path, actual_sheet_name: str) -> int:
    # 사용자가 SAP에서 새로 내려받아 넣은 RAW가 같은 기간의 분해본보다 우선한다.
    if re.search(r"(?:^|\s)RAW$", path.stem, re.IGNORECASE):
        return 30
    if actual_sheet_name.strip() not in {"Sheet1", "시트1"}:
        return 20
    return 10


def discover_input_files(input_dir: Path):
    selected: dict[str, tuple[Path, str, int]] = {}
    notices: list[str] = []
    for path in sorted(input_dir.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if len(wb.sheetnames) != 1:
            wb.close()
            raise RuntimeError(f"적재파일은 시트가 1개여야 합니다: {path.name}")
        actual = wb.sheetnames[0].strip()
        wb.close()
        logical = logical_sheet_name(path, actual)
        priority = input_priority(path, actual)
        if logical not in selected:
            selected[logical] = (path, actual, priority)
            continue
        previous_path, previous_actual, previous_priority = selected[logical]
        if priority == previous_priority:
            raise RuntimeError(f"같은 논리 시트의 적재파일이 중복되었습니다: {logical} ({previous_path.name}, {path.name})")
        if priority > previous_priority:
            selected[logical] = (path, actual, priority)
            notices.append(f"{logical}: {path.name} 우선, {previous_path.name} 제외")
        else:
            notices.append(f"{logical}: {previous_path.name} 우선, {path.name} 제외")
    return selected, notices
