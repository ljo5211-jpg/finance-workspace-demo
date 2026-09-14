from __future__ import annotations

import argparse
import json
import math
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import openpyxl

from input_discovery import discover_input_files


YEAR_END_RE = re.compile(r"^(\d{2})년말$")
MONTH_RE = re.compile(r"^(\d{2})년\s*(\d{1,2})월$")
FX_PL_CODES = {"70210001", "72010002", "74030001", "74140001", "74450008"}
BANK_ACCOUNT = "외화보통예금"
REFERENCE_ACCOUNT_ORDER = {
    "매출채권": 1,
    "장기대여금": 2,
    "미수금(기타)": 3,
    "미수금(거래처)": 4,
    "단기차입금": 5,
    "미지급금(거래처)": 6,
    "장기차입금": 7,
    "매입채무": 8,
    BANK_ACCOUNT: 9,
}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def parse_rate(value: Any) -> float:
    s = text(value)
    if not s:
        return 0.0
    first_dot = s.find(".")
    last_dot = s.rfind(".")
    if first_dot >= 0 and first_dot != last_dot:
        s = s[:last_dot].replace(".", "") + "." + s[last_dot + 1 :]
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        result = float(s)
    except ValueError:
        return 0.0
    return result * 1000 if 0 < result < 10 else result


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        numeric = float(value)
        compact = str(int(numeric))
        if len(compact) == 8 and compact[:4].isdigit() and int(compact[:4]) >= 1900:
            try:
                return datetime.strptime(compact, "%Y%m%d").date()
            except ValueError:
                pass
        if numeric > 0:
            try:
                return date(1899, 12, 30) + timedelta(days=numeric)
            except OverflowError:
                return None
    s = text(value)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def ledger_month(value: Any) -> int | None:
    if isinstance(value, (int, float)) and float(value).is_integer() and 1 <= int(value) <= 12:
        return int(value)
    d = parse_date(value)
    if d:
        return d.month
    s = text(value).replace("년", "-").replace("월", "").replace(".", "-").replace("/", "-")
    parts = s.split("-")
    if len(parts) >= 2 and parts[1].strip().isdigit():
        m = int(parts[1])
        return m if 1 <= m <= 12 else None
    return None


def normalize_evidence(value: Any) -> str:
    value = text(value)
    return "" if value in {"24", "33"} else value


def normalize_account(name: str) -> str:
    s = text(name)
    mappings = (
        ("외화보통", BANK_ACCOUNT),
        ("보통예금", "보통예금"),
        ("매출채권", "매출채권"),
        ("매입채무", "매입채무"),
        ("장기대여금", "장기대여금"),
        ("단기차입금", "단기차입금"),
        ("장기차입금", "장기차입금"),
        ("단기금융상품", "단기금융상품"),
        ("미수금(거래처)", "미수금(거래처)"),
        ("미수금(기타)", "미수금(기타)"),
        ("미지급금(거래처)", "미지급금(거래처)"),
        ("미지급금(기타)", "미지급금(기타)"),
    )
    for needle, replacement in mappings:
        if needle in s:
            return replacement
    return s


def account_by_gl(code: str) -> str:
    code = text(code)
    if code.startswith(("1032", "1031")):
        return BANK_ACCOUNT
    if code.startswith(("1053", "1051")):
        return "매출채권"
    if code.startswith(("2011", "2012")):
        return "매입채무"
    exact = {
        "10600002": "미수금(거래처)",
        "10600099": "미수금(기타)",
        "20150007": "미지급금(거래처)",
        "20150099": "미지급금(거래처)",
        "20160099": "단기차입금",
        "20160002": "단기차입금",
        "25101099": "장기차입금",
        "25101002": "장기차입금",
        "25101001": "장기차입금",
    }
    if code in exact:
        return exact[code]
    if code.startswith("1553"):
        return "장기대여금"
    if code.startswith("1552"):
        return "단기금융상품"
    return ""


def economic_origin_by_gl(code: str) -> str:
    code = text(code)
    if not code or code.endswith("99"):
        return ""
    owner = account_by_gl(code)
    return "" if owner == BANK_ACCOUNT else owner


def semantic_origin(counter: str) -> str:
    return {
        "74010001": "단기차입금",
        "90210002": BANK_ACCOUNT,
        "104513": "매출채권",
        "3110": "매출채권",
    }.get(text(counter), "")


def account_type(name: str) -> str:
    if name in {BANK_ACCOUNT, "매출채권", "미수금(거래처)", "미수금(기타)", "장기대여금", "단기금융상품"}:
        return "asset"
    if any(word in name for word in ("채무", "미지급", "차입", "부채")):
        return "liability"
    return "asset"


def is_bank_code(code: str) -> bool:
    return text(code).startswith("1032")


def is_fx_pl(code: str) -> bool:
    return text(code) in FX_PL_CODES


@dataclass
class Account:
    name: str
    kind: str
    code: str
    month_count: int
    balances: list[float] = field(init=False)
    pls: list[float] = field(init=False)
    rates: list[float] = field(init=False)
    prevs: list[float] = field(init=False)
    realized: list[float] = field(init=False)
    increases: list[float] = field(init=False)
    decreases: list[float] = field(init=False)
    attributed_prevs: list[float] = field(init=False)
    reallocation: list[float] = field(init=False)
    attributed_realized: list[float] = field(init=False)
    realized_reallocation: list[float] = field(init=False)
    cost_rates: list[float] = field(init=False)
    items: list[list[dict[str, Any]]] = field(init=False)
    pre_pl: float | None = None

    def __post_init__(self):
        self.balances = [0.0] * self.month_count
        self.pls = [0.0] * self.month_count
        self.rates = [0.0] * self.month_count
        self.prevs = [0.0] * self.month_count
        self.realized = [0.0] * self.month_count
        self.increases = [0.0] * self.month_count
        self.decreases = [0.0] * self.month_count
        self.attributed_prevs = [0.0] * self.month_count
        self.reallocation = [0.0] * self.month_count
        self.attributed_realized = [0.0] * self.month_count
        self.realized_reallocation = [0.0] * self.month_count
        self.cost_rates = [0.0] * self.month_count
        self.items = [[] for _ in range(self.month_count)]


@dataclass(frozen=True)
class LedgerRow:
    row_no: int
    slip: str
    clearing: str
    gl: str
    gl_name: str
    posting_key: str
    business_area: str
    description: str
    customer: str
    vendor: str
    counter: str
    reversal: str
    tx_date: date | None
    local: float
    currency: str
    foreign: float


@dataclass(frozen=True)
class PLRow:
    row_no: int
    slip: str
    clearing: str
    gl: str
    business_area: str
    customer: str
    vendor: str
    counter: str
    tx_date_raw: Any
    tx_date: date | None
    local: float


def load_single_sheet(path: Path, values_only: bool = True):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=values_only)
    if len(wb.sheetnames) != 1:
        wb.close()
        raise RuntimeError(f"적재파일은 시트가 1개여야 합니다: {path.name}")
    return wb, wb[wb.sheetnames[0]]


def discover_inputs(input_dir: Path):
    selected, notices = discover_input_files(input_dir)
    for notice in notices:
        print(f"[우선순위] {notice}")
    return {logical: values[0] for logical, values in selected.items()}


def determine_periods(sheets: dict[str, Path]):
    month_candidates: list[tuple[int, int, str]] = []
    year_ends: list[tuple[int, str]] = []
    for name in sheets:
        if match := MONTH_RE.match(name):
            month_candidates.append((2000 + int(match.group(1)), int(match.group(2)), name))
        elif match := YEAR_END_RE.match(name):
            year_ends.append((2000 + int(match.group(1)), name))
    if not month_candidates:
        raise RuntimeError("월별 외화평가 시트가 없습니다.")
    report_year = max(year for year, _, _ in month_candidates)
    monthly = sorted((m, name) for year, m, name in month_candidates if year == report_year)
    if not monthly:
        raise RuntimeError("보고연도 월별 시트를 결정할 수 없습니다.")
    baseline_candidates = sorted((year, name) for year, name in year_ends if year < report_year)
    if not baseline_candidates:
        raise RuntimeError("보고연도 직전 연말 시트가 없습니다.")
    baseline_year, baseline = baseline_candidates[-1]
    pre_baseline = next((name for year, name in reversed(baseline_candidates[:-1]) if year < baseline_year), "")
    expected_months = list(range(1, monthly[-1][0] + 1))
    actual_months = [month for month, _ in monthly]
    if actual_months != expected_months:
        raise RuntimeError(f"월별 평가 시트가 연속적이지 않습니다: {actual_months}")
    return report_year, baseline, pre_baseline, [name for _, name in monthly]


def ensure_account(accounts: dict[str, Account], name: str, month_count: int, code: str = "") -> Account:
    name = normalize_account(name)
    if name not in accounts:
        accounts[name] = Account(name, account_type(name), code, month_count)
    elif not accounts[name].code and code:
        accounts[name].code = code
    return accounts[name]


def parse_evaluation_sheet(path: Path, month_index: int, month_count: int, accounts: dict[str, Account]) -> float:
    wb, ws = load_single_sheet(path)
    rate = 0.0
    for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_no == 1:
            continue
        values = list(row) + [None] * max(0, 19 - len(row))
        currency = text(values[11]).upper()
        if currency != "USD":
            continue
        raw_name = text(values[3])
        if not raw_name or "합계" in raw_name or "소계" in raw_name:
            continue
        current_rate = parse_rate(values[10])
        if not rate and current_rate:
            rate = current_rate
        account = ensure_account(accounts, raw_name, month_count, text(values[2]))
        foreign = number(values[12])
        book = number(values[14])
        evaluated = number(values[16])
        posting = number(values[18])
        account.balances[month_index] += foreign
        account.pls[month_index] += evaluated - book
        account.prevs[month_index] += posting
        account.rates[month_index] = current_rate or rate
        tx_date = parse_date(values[9])
        # VBA 기준 상세 모달은 H열 '전표번호'를 사용한다. I열 Invoice NO를
        # 사용하면 금액은 같아도 전표 추적 번호가 기준 대시보드와 달라진다.
        invoice = text(values[7])
        if account.name == BANK_ACCOUNT and invoice == "0":
            invoice = ""
        account.items[month_index].append({
            "q": clean_number(foreign),
            "date": tx_date.isoformat() if tx_date else "",
            "inv": invoice,
        })
    wb.close()
    for account in accounts.values():
        if not account.rates[month_index]:
            account.rates[month_index] = rate
    return rate


def parse_pre_baseline(path: Path, accounts: dict[str, Account]):
    wb, ws = load_single_sheet(path)
    totals: defaultdict[str, float] = defaultdict(float)
    for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_no == 1:
            continue
        values = list(row) + [None] * max(0, 19 - len(row))
        if text(values[11]).upper() != "USD":
            continue
        name = normalize_account(text(values[3]))
        if name in accounts:
            totals[name] += number(values[16]) - number(values[14])
    wb.close()
    for name, value in totals.items():
        accounts[name].pre_pl = value


def load_fc_ledger(path: Path) -> list[LedgerRow]:
    wb, ws = load_single_sheet(path)
    rows: list[LedgerRow] = []
    for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_no == 1:
            continue
        values = list(row) + [None] * max(0, 31 - len(row))
        rows.append(LedgerRow(
            row_no=row_no,
            slip=normalize_evidence(values[1]),
            clearing=normalize_evidence(values[2]),
            gl=normalize_evidence(values[3]),
            gl_name=text(values[4]),
            posting_key=text(values[5]),
            business_area=normalize_evidence(values[8]),
            description=text(values[12]),
            customer=normalize_evidence(values[13]),
            vendor=normalize_evidence(values[14]),
            counter=normalize_evidence(values[16]),
            reversal=normalize_evidence(values[17]),
            tx_date=parse_date(values[23]),
            local=number(values[27]),
            currency=text(values[28]).upper(),
            foreign=number(values[29]),
        ))
    wb.close()
    return rows


def load_pl_ledger(path: Path) -> list[PLRow]:
    wb, ws = load_single_sheet(path)
    rows: list[PLRow] = []
    for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_no == 1:
            continue
        values = list(row) + [None] * max(0, 31 - len(row))
        rows.append(PLRow(
            row_no=row_no,
            slip=normalize_evidence(values[1]),
            clearing=normalize_evidence(values[2]),
            gl=normalize_evidence(values[3]),
            business_area=normalize_evidence(values[8]),
            customer=normalize_evidence(values[13]),
            vendor=normalize_evidence(values[14]),
            counter=normalize_evidence(values[16]),
            tx_date_raw=values[23],
            tx_date=parse_date(values[23]),
            local=number(values[27]),
        ))
    wb.close()
    return rows


def add_unique(mapping: dict[str, str | None], key: str, owner: str):
    if not key or not owner:
        return
    if key not in mapping:
        mapping[key] = owner
    elif mapping[key] != owner:
        mapping[key] = None


def build_evidence(fc_rows: list[LedgerRow]):
    maps = {name: {} for name in ("partner", "slip_ba", "clear_ba", "slip", "clear")}
    tx_rows: defaultdict[tuple[str, date], list[LedgerRow]] = defaultdict(list)
    bank_foreign: defaultdict[tuple[str, date], float] = defaultdict(float)
    suspense: list[LedgerRow] = []
    for row in fc_rows:
        if row.tx_date and row.slip:
            tx_rows[(row.slip, row.tx_date)].append(row)
        if row.tx_date and row.slip and is_bank_code(row.gl) and row.currency == "USD":
            bank_foreign[(row.slip, row.tx_date)] += row.foreign
        owner = economic_origin_by_gl(row.gl)
        if owner:
            add_unique(maps["partner"], row.customer, owner)
            add_unique(maps["partner"], row.vendor, owner)
            add_unique(maps["slip"], row.slip, owner)
            add_unique(maps["clear"], row.clearing, owner)
            if row.business_area:
                add_unique(maps["slip_ba"], f"{row.slip}|{row.business_area}", owner)
                add_unique(maps["clear_ba"], f"{row.clearing}|{row.business_area}", owner)
        if row.counter == "90210001" and owner in {"매출채권", "미수금(거래처)"} and row.tx_date:
            suspense.append(row)
    return maps, tx_rows, bank_foreign, suspense


def weighted_origins(rows: Iterable[LedgerRow], business_area: str, partner_code: str) -> dict[str, float]:
    buckets: dict[str, defaultdict[str, float]] = {
        name: defaultdict(float) for name in ("all", "ba", "partner", "partner_ba", "counter", "counter_ba")
    }
    for row in rows:
        owner = economic_origin_by_gl(row.gl)
        if not owner:
            continue
        weight = abs(row.foreign) or abs(row.local)
        if weight <= 0:
            continue
        buckets["all"][owner] += weight
        if business_area and row.business_area == business_area:
            buckets["ba"][owner] += weight
        if partner_code:
            if row.counter == partner_code:
                buckets["counter"][owner] += weight
                if business_area and row.business_area == business_area:
                    buckets["counter_ba"][owner] += weight
            if row.customer == partner_code or row.vendor == partner_code:
                buckets["partner"][owner] += weight
                if business_area and row.business_area == business_area:
                    buckets["partner_ba"][owner] += weight
    for name in ("counter_ba", "counter", "partner_ba", "partner", "ba", "all"):
        if buckets[name]:
            return dict(buckets[name])
    return {}


def resolve_suspense(suspense_rows: list[LedgerRow], tx_date: date | None, business_area: str) -> str:
    if not tx_date:
        return ""
    buckets = {name: defaultdict(float) for name in ("date_ba", "date", "month_ba", "month")}
    for row in suspense_rows:
        if not row.tx_date or (row.tx_date.year, row.tx_date.month) != (tx_date.year, tx_date.month):
            continue
        owner = economic_origin_by_gl(row.gl)
        weight = abs(row.foreign) or abs(row.local)
        buckets["month"][owner] += weight
        if business_area and row.business_area == business_area:
            buckets["month_ba"][owner] += weight
        if row.tx_date == tx_date:
            buckets["date"][owner] += weight
            if business_area and row.business_area == business_area:
                buckets["date_ba"][owner] += weight
    for name in ("date_ba", "date", "month_ba", "month"):
        if buckets[name]:
            return max(buckets[name], key=buckets[name].get)
    return ""


def resolve_from_maps(counter: str, slip: str, clearing: str, ba: str, maps) -> str:
    direct = account_by_gl(counter)
    if direct and direct != BANK_ACCOUNT:
        return direct
    partner = maps["partner"].get(counter)
    if partner:
        return partner
    candidates: list[str] = []
    if ba:
        for mapping, key in ((maps["slip_ba"], f"{slip}|{ba}"), (maps["clear_ba"], f"{clearing}|{ba}")):
            value = mapping.get(key)
            if value:
                candidates.append(value)
        if candidates:
            return candidates[0] if len(set(candidates)) == 1 else ""
    candidates = [value for value in (maps["slip"].get(slip), maps["clear"].get(clearing)) if value]
    return candidates[0] if candidates and len(set(candidates)) == 1 else ""


def allocate(target: defaultdict[str, list[float]], owner: str, month: int, amount: float, month_count: int):
    if owner not in target:
        target[owner] = [0.0] * month_count
    target[owner][month] += amount


def allocate_weighted(target, weights: dict[str, float], month: int, amount: float, month_count: int) -> bool:
    total = sum(weights.values())
    if total <= 0:
        return False
    for owner, weight in weights.items():
        allocate(target, owner, month, amount * weight / total, month_count)
    return True


def build_realized(
    accounts: dict[str, Account],
    fc_rows: list[LedgerRow],
    pl_rows: list[PLRow],
    month_count: int,
):
    maps, tx_rows, bank_foreign, suspense = build_evidence(fc_rows)
    classified: defaultdict[str, list[float]] = defaultdict(lambda: [0.0] * month_count)
    ledger_totals = [0.0] * month_count
    pl_companions: defaultdict[tuple[str, date], list[PLRow]] = defaultdict(list)
    for row in pl_rows:
        if row.tx_date and row.slip:
            pl_companions[(row.slip, row.tx_date)].append(row)

    for row in pl_rows:
        month = ledger_month(row.tx_date_raw)
        if not month or not (1 <= month < month_count) or not row.local:
            continue
        amount = -row.local
        ledger_totals[month] += amount
        owner = account_by_gl(row.counter)
        allocated = False
        if not owner and row.tx_date:
            weights = weighted_origins(tx_rows.get((row.slip, row.tx_date), []), row.business_area, row.counter)
            allocated = allocate_weighted(classified, weights, month, amount, month_count)
        if allocated:
            continue
        if not owner and row.counter == "90210001":
            owner = resolve_suspense(suspense, row.tx_date, row.business_area)
        if not owner:
            owner = semantic_origin(row.counter)
        if not owner:
            owner = resolve_from_maps(row.counter, row.slip, row.clearing, row.business_area, maps)
        if not owner and row.tx_date:
            companion_owners = set()
            for companion in pl_companions.get((row.slip, row.tx_date), []):
                if companion.row_no == row.row_no:
                    continue
                candidate = account_by_gl(companion.counter) or semantic_origin(companion.counter)
                if companion.counter == "90210001" and not candidate:
                    candidate = resolve_suspense(suspense, companion.tx_date, companion.business_area)
                if candidate:
                    companion_owners.add(candidate)
            if len(companion_owners) == 1:
                owner = next(iter(companion_owners))
        if not owner and row.tx_date:
            foreign = bank_foreign.get((row.slip, row.tx_date), 0.0)
            if row.counter.isdigit() and len(row.counter) <= 7:
                owner = "매입채무" if foreign < -0.005 else "매출채권" if foreign > 0.005 else ""
            elif is_fx_pl(row.counter) and abs(foreign) > 0.005:
                owner = BANK_ACCOUNT
        # 손익계정 상호대체도 별도 표시하지 않고 증거가 남는 현금에 보수적으로 둔다.
        if not owner:
            owner = BANK_ACCOUNT
        allocate(classified, owner, month, amount, month_count)

    for name, values in classified.items():
        account = ensure_account(accounts, name, month_count)
        account.realized = values[:]
        account.attributed_realized = values[:]
    bank = ensure_account(accounts, BANK_ACCOUNT, month_count)
    for month in range(1, month_count):
        current = sum(account.realized[month] for account in accounts.values())
        difference = ledger_totals[month] - current
        bank.realized[month] += difference
        bank.attributed_realized[month] += difference
    return ledger_totals, maps, tx_rows, suspense


def mark_reversal_pairs(rows: list[LedgerRow]) -> set[int]:
    by_slip: defaultdict[str, list[LedgerRow]] = defaultdict(list)
    for row in rows:
        if row.slip:
            by_slip[row.slip].append(row)
    handled: set[int] = set()
    for row in rows:
        if row.row_no in handled or not is_bank_code(row.gl) or not row.reversal or not row.tx_date:
            continue
        for candidate in by_slip.get(row.reversal, []):
            if candidate.row_no in handled:
                continue
            if candidate.reversal != row.slip or candidate.gl != row.gl or candidate.currency != row.currency or candidate.tx_date != row.tx_date:
                continue
            if abs(candidate.foreign + row.foreign) < 0.005 and abs(candidate.local + row.local) < 0.5:
                handled.update({row.row_no, candidate.row_no})
                break
    return handled


def baseline_cash_states(path: Path):
    wb, ws = load_single_sheet(path)
    states: defaultdict[str, dict[str, float]] = defaultdict(lambda: {"fc": 0.0, "lc": 0.0})
    for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_no == 1:
            continue
        values = list(row) + [None] * max(0, 17 - len(row))
        gl = text(values[2])
        name = text(values[3])
        if text(values[11]).upper() == "USD" and (is_bank_code(gl) or "외화보통" in name):
            key = gl or "1032_POOL"
            states[key]["fc"] += number(values[12])
            states[key]["lc"] += number(values[16])
    wb.close()
    return states


def weighted_rate(states) -> float:
    total_fc = sum(state["fc"] for state in states.values())
    total_lc = sum(state["lc"] for state in states.values())
    return total_lc / total_fc if abs(total_fc) > 1e-6 else 0.0


def cash_state(states, code: str):
    key = code or "1032_POOL"
    if key not in states:
        states[key] = {"fc": 0.0, "lc": 0.0}
    return states[key]


def owner_weights_for_row(row: LedgerRow, tx_rows, maps, suspense) -> dict[str, float]:
    weights = weighted_origins(tx_rows.get((row.slip, row.tx_date), []), row.business_area, row.counter) if row.tx_date else {}
    if weights:
        return weights
    owner = ""
    if row.counter == "90210001":
        owner = resolve_suspense(suspense, row.tx_date, row.business_area)
    if not owner:
        owner = semantic_origin(row.counter)
    if not owner:
        owner = resolve_from_maps(row.counter, row.slip, row.clearing, row.business_area, maps)
    if not owner and row.counter.isdigit() and len(row.counter) <= 7:
        owner = "매입채무"
    if not owner and (is_fx_pl(row.counter) or row.counter in {"70110001", "62210001"}):
        owner = BANK_ACCOUNT
    return {owner or BANK_ACCOUNT: 1.0}


def build_valuation_attribution(
    accounts: dict[str, Account],
    baseline_path: Path,
    fc_rows: list[LedgerRow],
    maps,
    tx_rows,
    suspense,
    month_count: int,
):
    for account in accounts.values():
        account.attributed_prevs = account.prevs[:]
        account.reallocation = [0.0] * month_count
    bank = ensure_account(accounts, BANK_ACCOUNT, month_count)
    states = baseline_cash_states(baseline_path)
    if not states:
        states["1032_POOL"] = {"fc": bank.balances[0], "lc": bank.balances[0] * bank.rates[0]}
    bank.cost_rates[0] = weighted_rate(states)

    handled = mark_reversal_pairs(fc_rows)
    bank_rows = [
        row for row in fc_rows
        if row.tx_date and is_bank_code(row.gl) and row.currency == "USD" and abs(row.foreign) > 0.005
    ]
    grouped: defaultdict[tuple[date, str], list[LedgerRow]] = defaultdict(list)
    for row in bank_rows:
        grouped[(row.tx_date, row.slip)].append(row)

    # 은행 간 대체를 먼저 짝지어 같은 원가층을 이동시킨다.
    transfer_pairs: defaultdict[date, list[tuple[LedgerRow, LedgerRow]]] = defaultdict(list)
    for (tx_date, _slip), group in grouped.items():
        outs = [row for row in group if row.row_no not in handled and is_bank_code(row.counter) and row.foreign < -0.005]
        ins = [row for row in group if row.row_no not in handled and is_bank_code(row.counter) and row.foreign > 0.005]
        for out in outs:
            match = next((row for row in ins if row.row_no not in handled and abs(row.foreign + out.foreign) < 0.01), None)
            if match:
                handled.update({out.row_no, match.row_no})
                transfer_pairs[tx_date].append((out, match))

    settlement: defaultdict[str, list[float]] = defaultdict(lambda: [0.0] * month_count)
    rows_by_date: defaultdict[date, list[LedgerRow]] = defaultdict(list)
    for row in bank_rows:
        if row.row_no not in handled:
            rows_by_date[row.tx_date].append(row)

    for month in range(1, month_count):
        month_dates = sorted({d for d in set(rows_by_date) | set(transfer_pairs) if d.month == month})
        for tx_date in month_dates:
            for outgoing, incoming in transfer_pairs.get(tx_date, []):
                source = cash_state(states, outgoing.gl)
                destination = cash_state(states, incoming.gl)
                rate = source["lc"] / source["fc"] if abs(source["fc"]) > 1e-6 else weighted_rate(states)
                fc = -outgoing.foreign
                cost = fc * rate
                source["fc"] -= fc
                source["lc"] -= cost
                destination["fc"] += fc
                destination["lc"] += cost
            rows_today = rows_by_date.get(tx_date, [])
            for row in (item for item in rows_today if item.foreign > 0.005):
                state = cash_state(states, row.gl)
                state["fc"] += row.foreign
                state["lc"] += row.local
            for row in (item for item in rows_today if item.foreign < -0.005):
                state = cash_state(states, row.gl)
                rate = state["lc"] / state["fc"] if abs(state["fc"]) > 1e-6 else weighted_rate(states)
                fc = -row.foreign
                cost = fc * rate
                economic_difference = (-row.local) - cost
                weights = owner_weights_for_row(row, tx_rows, maps, suspense)
                total_weight = sum(weights.values()) or 1.0
                for owner, weight in weights.items():
                    settlement[owner][month] += economic_difference * weight / total_weight
                state["fc"] -= fc
                state["lc"] -= cost

        rate = weighted_rate(states) or bank.rates[month]
        economic_bank = bank.balances[month] * (bank.rates[month] - rate)
        bank.attributed_prevs[month] = economic_bank
        bank.reallocation[month] = economic_bank - bank.prevs[month]
        bank.cost_rates[month] = rate
        current_difference = bank.prevs[month] - economic_bank
        prior_difference = 0.0 if month == 1 else bank.prevs[month - 1] - bank.attributed_prevs[month - 1]
        required = current_difference - prior_difference
        explained = sum(values[month] for values in settlement.values())
        settlement[BANK_ACCOUNT][month] += required - explained

    for owner, values in settlement.items():
        account = ensure_account(accounts, owner, month_count)
        cumulative = 0.0
        for month in range(1, month_count):
            cumulative += values[month]
            account.attributed_prevs[month] += cumulative
            account.reallocation[month] += cumulative

    for month in range(1, month_count):
        difference = sum(account.reallocation[month] for account in accounts.values())
        bank.attributed_prevs[month] -= difference
        bank.reallocation[month] -= difference


def add_activity(accounts: dict[str, Account], fc_rows: list[LedgerRow], month_count: int):
    for row in fc_rows:
        if row.currency != "USD" or abs(row.foreign) <= 0.01 or not row.tx_date:
            continue
        month = row.tx_date.month
        if not (1 <= month < month_count):
            continue
        name = normalize_account(row.gl_name)
        if name not in accounts:
            continue
        account = accounts[name]
        if account.kind == "liability":
            if row.foreign < 0:
                account.increases[month] += row.foreign
            else:
                account.decreases[month] += row.foreign
        else:
            if row.foreign > 0:
                account.increases[month] += row.foreign
            else:
                account.decreases[month] += row.foreign


def relevant(account: Account) -> bool:
    if any(abs(value) >= 1e-6 for value in account.balances[1:]):
        return True
    krw_values = account.prevs[1:] + account.attributed_prevs[1:] + account.realized[1:] + account.attributed_realized[1:]
    return any(abs(value) >= 0.5 for value in krw_values)


def clean_number(value: float) -> int | float:
    """VBA Double의 JSON 표기와 같은 15자리 유효숫자로 직렬화한다."""
    if abs(value) < 1e-9:
        return 0
    cleaned = float(format(float(value), ".15g"))
    return int(cleaned) if cleaned.is_integer() else cleaned


def build_payload(project_root: Path):
    settings = json.loads((project_root / "dashboard-settings.json").read_text(encoding="utf-8"))
    sheets = discover_inputs(project_root / "적재")
    for required in ("외화계정원장", "환차손익원장"):
        if required not in sheets:
            raise RuntimeError(f"필수 적재파일 누락: {required}")
    report_year, baseline, pre_baseline, monthly = determine_periods(sheets)
    month_names = [baseline] + monthly
    month_count = len(month_names)
    accounts: dict[str, Account] = {}
    for index, name in enumerate(month_names):
        parse_evaluation_sheet(sheets[name], index, month_count, accounts)
    if pre_baseline:
        parse_pre_baseline(sheets[pre_baseline], accounts)

    fc_rows = load_fc_ledger(sheets["외화계정원장"])
    pl_rows = load_pl_ledger(sheets["환차손익원장"])
    add_activity(accounts, fc_rows, month_count)
    realized_totals, maps, tx_rows, suspense = build_realized(accounts, fc_rows, pl_rows, month_count)
    build_valuation_attribution(accounts, sheets[baseline], fc_rows, maps, tx_rows, suspense, month_count)

    accounts.pop("가수금", None)
    accounts.pop("손익분류조정", None)
    display_accounts = [account for account in accounts.values() if relevant(account)]
    display_accounts.sort(key=lambda a: (REFERENCE_ACCOUNT_ORDER.get(a.name, 90), a.name))

    checks = []
    for month in range(1, month_count):
        valuation_ledger = sum(a.prevs[month] for a in accounts.values())
        valuation_attributed = sum(a.attributed_prevs[month] for a in accounts.values())
        realized_classified = sum(a.attributed_realized[month] for a in accounts.values())
        checks.append({
            "month": month_names[month],
            "valuationLedger": clean_number(valuation_ledger),
            "valuationAttributed": clean_number(valuation_attributed),
            "valuationDifference": clean_number(valuation_attributed - valuation_ledger),
            "realizedLedger": clean_number(realized_totals[month]),
            "realizedAttributed": clean_number(realized_classified),
            "realizedDifference": clean_number(realized_classified - realized_totals[month]),
        })

    manifest_path = project_root / "private" / "source-manifest.json"
    source_manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    payload_accounts = []
    for account in display_accounts:
        payload_account = {
            "acc": account.name,
            "type": account.kind,
            "code": account.code,
            "balances": [clean_number(v) for v in account.balances],
            "pls": [clean_number(v) for v in account.pls],
            "prevs": [clean_number(v) for v in account.prevs],
            "attributedPrevs": [clean_number(v) for v in account.attributed_prevs],
            "reallocation": [clean_number(v) for v in account.reallocation],
            "rates": [clean_number(v) for v in account.rates],
            "items": account.items,
            "realized": [clean_number(v) for v in account.realized],
            "attributedRealized": [clean_number(v) for v in account.attributed_realized],
            "realizedReallocation": [clean_number(v) for v in account.realized_reallocation],
            "increases": [clean_number(v) for v in account.increases],
            "decreases": [clean_number(v) for v in account.decreases],
            "costRates": [clean_number(v) for v in account.cost_rates],
        }
        if account.pre_pl is not None:
            payload_account["preBaselinePL"] = clean_number(account.pre_pl)
        payload_accounts.append(payload_account)
    return {
        "months": month_names,
        "reportMonth": month_names[-1],
        "reportYear": report_year,
        "accounts": payload_accounts,
        "checks": checks,
        "settings": settings,
        "source": {
            "sourceWorkbook": source_manifest.get("sourceWorkbook", "개별 적재파일"),
            "sourceSha256": source_manifest.get("sourceSha256", ""),
            "builtAt": datetime.now().isoformat(timespec="seconds"),
            "inputFiles": [path.name for path in sorted(set(sheets.values()))],
            "fcRows": len(fc_rows),
            "plRows": len(pl_rows),
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", required=True)
    args = parser.parse_args()
    project_root = Path(args.project_root).resolve()
    payload = build_payload(project_root)
    template = (project_root / "dashboard-template.html").read_text(encoding="utf-8")
    if "__FX_DATA_JSON__" not in template:
        raise RuntimeError("dashboard-template.html 데이터 자리표시자가 없습니다.")
    dashboard_database = {
        "months": payload["months"],
        "reportMonth": payload["reportMonth"],
        "accounts": payload["accounts"],
    }
    data_json = json.dumps(dashboard_database, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = template.replace("__FX_DATA_JSON__", data_json)
    timestamp = datetime.now().strftime("%H%M%S")
    (project_root / "publish").mkdir(parents=True, exist_ok=True)
    (project_root / "private").mkdir(parents=True, exist_ok=True)
    output_path = project_root / "publish" / f"FX_Analysis_Dashboard_{timestamp}.html"
    output_path.write_text(html, encoding="utf-8")
    (project_root / "private" / "latest-data.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    failed = [check for check in payload["checks"] if abs(check["valuationDifference"]) > 1 or abs(check["realizedDifference"]) > 1]
    if failed:
        raise RuntimeError(f"원장 대사 실패: {failed}")
    print(f"[생성] {output_path}")
    print(f"[대사] {len(payload['checks'])}개월 평가·환차손익 Net 일치")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
