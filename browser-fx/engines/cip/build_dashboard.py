from __future__ import annotations

import argparse
import json
import re
from calendar import monthrange
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from input_discovery import Period, discover_inputs, previous_final


WBS_RE = re.compile(r"\b\d{5}-CT[A-Z]\d{2}\b", re.IGNORECASE)
SHEET_ORDER = ["HQ", "화학_QL", "화학_기타", "DP_8.6G", "DP_기타", "전장_기타"]
USEFUL_LIFE_YEARS = {
    "토지": None,
    "건물": 30,
    "구축물": 15,
    "기계장치": 5,
    "집기비품": 5,
    "개발비": 5,
    "기타무형자산": 5,
    "차량운반구": 5,
    "소프트웨어": 5,
    "비용처리": None,
}


def text(value: Any) -> str:
    return "" if value is None else " ".join(str(value).replace("\xa0", " ").split())


def number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0.0


def clean_number(value: float) -> int | float:
    if abs(value) < 1e-8:
        return 0
    rounded = round(float(value), 6)
    return int(rounded) if rounded.is_integer() else rounded


def iso_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def month_end(period: Period) -> date:
    year = 2000 + period.year
    return date(year, period.month, monthrange(year, period.month)[1])


def months_between(start: date | None, end: date) -> int | None:
    if start is None:
        return None
    return max(0, (end.year - start.year) * 12 + end.month - start.month)


def normalize_asset_id(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return text(value)


def extract_wbs(*values: Any) -> str:
    for value in values:
        match = WBS_RE.search(text(value))
        if match:
            return match.group(0).upper()
    return ""


def infer_business(asset: dict[str, Any], sheet_name: str = "") -> str:
    if sheet_name.startswith("HQ"):
        return "HQ"
    if sheet_name.startswith("화학"):
        return "화학"
    if sheet_name.startswith("DP") or sheet_name.startswith("6G"):
        return "디스플레이"
    if sheet_name.startswith("전장"):
        return "전자/전장"
    combined = f"{asset.get('name', '')} {asset.get('description', '')}".upper()
    if "CT_IN_CH" in combined or "CHEMTRONICS_CHEM" in combined:
        return "화학"
    if "CT_IN_DP" in combined or "CHEMTRONICS_DP" in combined or "8.6G" in combined:
        return "디스플레이"
    if "CHEMTRONICS_AM" in combined:
        return "자율주행"
    if "CHEMTRONICS_ELEC" in combined:
        return "전자"
    if asset.get("businessArea") == "1100":
        return "HQ"
    if asset.get("businessArea") == "1120":
        return "화학"
    if asset.get("businessArea") == "1130":
        return "디스플레이"
    return "기타"


def infer_sheet(asset: dict[str, Any]) -> str:
    business = infer_business(asset)
    if business == "HQ":
        return "HQ"
    if business == "화학":
        return "화학_QL" if "QL" in asset.get("name", "").upper() else "화학_기타"
    if business == "디스플레이":
        combined = f"{asset.get('name', '')} {asset.get('wbs', '')}".upper()
        return "DP_8.6G" if "8.6G" in combined or asset.get("assetId") in {"990020001", "990020002", "990020003"} else "DP_기타"
    return "전장_기타"


def infer_asset_type(asset: dict[str, Any], type_totals: dict[str, float] | None = None) -> str:
    if type_totals:
        usable = [(name, abs(value)) for name, value in type_totals.items() if name and name != "미분류" and abs(value) > 0.5]
        if usable:
            return max(usable, key=lambda item: item[1])[0]
    class_name = asset.get("className", "")
    name = asset.get("name", "")
    if "건물" in class_name or "사옥" in name or "성남금토" in name:
        return "건물"
    if "기계장치" in class_name:
        return "기계장치"
    if "차입원가" in name or "자본화" in name or "MES" in name.upper() or "LIMS" in name.upper():
        return "기타무형자산"
    return "기계장치" if asset.get("wbs") else "기타무형자산"


def load_register(path: Path, period: Period) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = [text(cell.value) for cell in worksheet[1]]
    index = {name: position for position, name in enumerate(headers) if name}
    assets = []
    totals = defaultdict(float)
    class_totals: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        class_name = text(row[index["자산 클래스 이름"]])
        if "건설중인자산" not in class_name:
            continue
        acquisition = row[index["취득일자"]]
        acquisition_date = acquisition.date() if isinstance(acquisition, datetime) else acquisition if isinstance(acquisition, date) else None
        item = {
            "assetId": normalize_asset_id(row[index["자산번호"]]),
            "subId": normalize_asset_id(row[index.get("하위 번호", 1)]),
            "name": text(row[index["자산명"]]),
            "classCode": text(row[index["클래스"]]),
            "className": class_name,
            "wbs": text(row[index["WBS 요소"]]).upper(),
            "description": text(row[index["내역"]]),
            "beginning": number(row[index["기초자산가액"]]),
            "ytdAdditions": number(row[index["당기자산증가"]]),
            "ytdTransfers": number(row[index["당기자산감소"]]),
            "ending": number(row[index["기말자산가액"]]),
            "acquisitionDate": acquisition_date,
            "costCenter": text(row[index.get("코스트센터", 34)]),
            "costCenterName": text(row[index.get("이름", 35)]),
            "businessArea": text(row[index.get("BusA", 36)]),
            "functionArea": text(row[index.get("기능 영역", 38)]),
            "costType": text(row[index.get("비용구분", 40)]),
            "vendorCode": text(row[index.get("공급업체", 44)]),
            "vendorName": text(row[index.get("이름 1", 45)]),
        }
        assets.append(item)
        for field in ("beginning", "ytdAdditions", "ytdTransfers", "ending"):
            totals[field] += item[field]
            class_totals[class_name][field] += item[field]
        class_totals[class_name]["count"] += 1
    workbook.close()
    return {
        "period": period,
        "file": path.name,
        "assets": assets,
        "totals": dict(totals),
        "classTotals": {name: dict(values) for name, values in class_totals.items()},
    }


def header_positions(worksheet: Any) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, worksheet.max_column + 1):
        label = text(worksheet.cell(5, column).value)
        if label and label not in result:
            result[label] = column - 1
    return result


def detail_value(row: tuple[Any, ...], positions: dict[str, int], name: str) -> Any:
    position = positions.get(name)
    return row[position] if position is not None and position < len(row) else None


def load_reference_final(path: Path, period: Period) -> dict[str, Any]:
    formula_book = openpyxl.load_workbook(path, read_only=False, data_only=False)
    value_book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    groups: list[dict[str, Any]] = []
    for sheet_name in formula_book.sheetnames:
        if sheet_name == "요약" or "감가상각" in sheet_name:
            continue
        formula_sheet = formula_book[sheet_name]
        value_sheet = value_book[sheet_name]
        positions = header_positions(formula_sheet)
        paid_position = positions.get("지급액 금액")
        contract_position = positions.get("계약금액")
        if paid_position is None:
            continue
        group_rows: list[int] = []
        for row_number in range(6, formula_sheet.max_row + 1):
            asset_id = normalize_asset_id(formula_sheet.cell(row_number, 1).value)
            paid_formula = formula_sheet.cell(row_number, paid_position + 1).value
            if asset_id and isinstance(paid_formula, str) and paid_formula.startswith("=") and "SUM" in paid_formula.upper():
                group_rows.append(row_number)
        for group_index, group_row in enumerate(group_rows):
            end_row = group_rows[group_index + 1] - 1 if group_index + 1 < len(group_rows) else formula_sheet.max_row
            asset_id = normalize_asset_id(formula_sheet.cell(group_row, 1).value)
            details = []
            type_totals: defaultdict[str, float] = defaultdict(float)
            expected_depreciation = 0.0
            for row_number in range(group_row + 1, end_row + 1):
                row_values = tuple(value_sheet.cell(row_number, column).value for column in range(1, formula_sheet.max_column + 1))
                if not any(value not in (None, "") for value in row_values[: paid_position + 1]):
                    continue
                if normalize_asset_id(row_values[0]) not in {"", asset_id}:
                    continue
                paid = number(row_values[paid_position])
                asset_type = text(detail_value(row_values, positions, "자산구분")) or "미분류"
                type_totals[asset_type] += paid
                depreciation_position = positions.get("예상 월 감가상각비")
                if depreciation_position is not None:
                    expected_depreciation += number(row_values[depreciation_position])
                details.append({
                    "sourceRow": row_number,
                    "assetId": asset_id,
                    "wbs": text(detail_value(row_values, positions, "WBS 코드") or detail_value(row_values, positions, "코드")),
                    "businessArea": text(row_values[2] if len(row_values) > 2 else None),
                    "postingDate": iso_date(row_values[3] if len(row_values) > 3 else None),
                    "description": text(row_values[4] if len(row_values) > 4 else None),
                    "assetType": asset_type,
                    "costCenter": text(detail_value(row_values, positions, "CCTR")),
                    "division": text(detail_value(row_values, positions, "사업부")),
                    "vendorCode": text(detail_value(row_values, positions, "코드")),
                    "vendorName": text(detail_value(row_values, positions, "거래처")),
                    "contractAmount": clean_number(number(detail_value(row_values, positions, "계약금액"))),
                    "paymentTerm": text(detail_value(row_values, positions, "지급조건")),
                    "paidAmount": clean_number(paid),
                    "approval": text(detail_value(row_values, positions, "현업 품의서")),
                    "drafter": text(detail_value(row_values, positions, "기안자")),
                    "waApproval": text(detail_value(row_values, positions, "WA품의")),
                    "document": text(detail_value(row_values, positions, "현업 전표") or detail_value(row_values, positions, "선급금 전표")),
                    "temporaryAsset": text(detail_value(row_values, positions, "고정자산 임시")),
                    "transferDocument": text(detail_value(row_values, positions, "건자 대체전표")),
                    "usefulLife": number(detail_value(row_values, positions, "내용연수")) or None,
                    "sourceStatus": "전월 최종본 승계",
                })
            if not expected_depreciation:
                for asset_type, amount in type_totals.items():
                    life = USEFUL_LIFE_YEARS.get(asset_type)
                    if life:
                        expected_depreciation += amount / life / 12
            paid_total = number(value_sheet.cell(group_row, paid_position + 1).value)
            contract_total = number(value_sheet.cell(group_row, (contract_position if contract_position is not None else paid_position) + 1).value)
            groups.append({
                "sourceRow": group_row,
                "sheet": sheet_name,
                "assetId": asset_id,
                "wbs": text(formula_sheet.cell(group_row, 2).value),
                "businessArea": text(formula_sheet.cell(group_row, 3).value),
                "name": text(formula_sheet.cell(group_row, 5).value),
                "contractTotal": clean_number(contract_total),
                "paidTotal": clean_number(paid_total),
                "typeTotals": {name: clean_number(value) for name, value in type_totals.items()},
                "expectedMonthlyDepreciation": clean_number(expected_depreciation),
                "details": details,
            })
    summary_value = number(value_book["요약"]["D3"].value) if "요약" in value_book.sheetnames else sum(number(group["paidTotal"]) for group in groups)
    formula_book.close()
    value_book.close()
    return {"file": path.name, "period": period, "summaryValue": summary_value, "groups": groups}


def load_ledger(path: Path, report_period: Period, assets_by_wbs: dict[str, dict[str, Any]], fallback_building_asset: str) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = [text(cell.value) for cell in worksheet[1]]
    index = {name: position for position, name in enumerate(headers) if name}
    cutoff = month_end(report_period)
    rows = []
    relevant_sum = 0.0
    current_month_rows = []
    future_rows = []
    for source in worksheet.iter_rows(min_row=2, values_only=True):
        account_name = text(source[index["G/L 계정명"]])
        if "건설중인자산" not in account_name:
            continue
        posting_value = source[index["전기일"]]
        posting_date = posting_value.date() if isinstance(posting_value, datetime) else posting_value if isinstance(posting_value, date) else None
        description = text(source[index["텍스트"]])
        wbs = extract_wbs(source[index["WBS 요소"]], description)
        matched_asset = assets_by_wbs.get(wbs, {}).get("assetId", "") if wbs else ""
        if not matched_asset and "건물" in account_name:
            matched_asset = fallback_building_asset
        amount = number(source[index["금액(현지 통화)"]])
        item = {
            "document": text(source[index["전표 번호"]]),
            "postingDate": posting_date.isoformat() if posting_date else None,
            "accountCode": text(source[index["계정"]]),
            "accountName": account_name,
            "postingKey": text(source[index["전기 키"]]),
            "amount": clean_number(amount),
            "description": description,
            "counterAccount": text(source[index["상계 계정"]]),
            "wbs": wbs,
            "matchedAssetId": matched_asset,
        }
        rows.append(item)
        if posting_date and posting_date <= cutoff:
            relevant_sum += amount
            if posting_date.year == cutoff.year and posting_date.month == cutoff.month:
                current_month_rows.append(item)
        elif posting_date and posting_date > cutoff:
            future_rows.append(item)
    workbook.close()
    return {
        "file": path.name,
        "rows": rows,
        "sumThroughCutoff": clean_number(relevant_sum),
        "currentMonthRows": current_month_rows,
        "futureRows": future_rows,
    }


def sheet_for_asset(asset: dict[str, Any], group: dict[str, Any] | None) -> str:
    if group and group["sheet"] in SHEET_ORDER:
        return group["sheet"]
    return infer_sheet(asset)


def last_addition(asset, group, ledger, history, cutoff):
    candidates = []
    for detail in (group or {}).get("details", []):
        value = detail.get("postingDate")
        if value and number(detail.get("paidAmount")) > .5 and value <= cutoff.isoformat():
            candidates.append((value, "최종본 증가 전기일"))
    for row in ledger["rows"]:
        # A building-account fallback is not an asset-level match.
        if row.get("wbs") and row["matchedAssetId"] == asset["assetId"] and row["amount"] > .5 and row.get("postingDate") and row["postingDate"] <= cutoff.isoformat():
            candidates.append((row["postingDate"], "원장 WBS 증가 전기일"))
    best = max(candidates, default=(None, "증가일 확인 필요"))
    previous = None
    for register in history:
        current = next((a for a in register["assets"] if a["assetId"] == asset["assetId"]), None)
        if current:
            base = previous[1] if previous and previous[0] == register["period"].year else 0
            if current["ytdAdditions"] - base > .5:
                end = month_end(register["period"]).isoformat()
                if best[0] is None or end[:7] > best[0][:7]:
                    best = (end, "자산대장 최종 증가월 · 월말 기준")
            previous = (register["period"].year, current["ytdAdditions"])
    return best


def depreciation_components(asset, group, use_defaults=True):
    buckets = defaultdict(float)
    for detail in (group or {}).get("details", []):
        kind = detail["assetType"]
        life = number(detail.get("usefulLife"))
        if not use_defaults:
            life=life if kind in USEFUL_LIFE_YEARS and .1<=life<=100 else None
        else:
            life=life or USEFUL_LIFE_YEARS.get(kind)
        if kind in {"토지", "비용처리"}:
            life = None
        buckets[(kind, life)] += number(detail["paidAmount"])
    total = sum(buckets.values())
    remaining = max(0, asset["ending"])
    if use_defaults and abs(total - remaining) > 1:
        # Never assign an unclassified balance to the dominant asset class.
        if total > remaining and total > 0:
            buckets = {key: amount * remaining / total for key, amount in buckets.items()}
        else:
            buckets[("미분류", None)] += remaining - total
    result = []
    for (kind, life), amount in buckets.items():
        if abs(amount) <= .5:
            continue
        exempt = kind in {"토지", "비용처리"}
        result.append({"assetType": kind, "amount": clean_number(amount), "usefulLife": life,
                       "monthlyDepreciation": clean_number(amount / life / 12) if life else 0,
                       "basis": "비상각" if exempt else ("명세 내용연수 / 기본 가정" if use_defaults else "최종본 내용연수") if life else "분류·내용연수 확인 필요"})
    return result


def current_final_groups(reference):
    from management_rules import SHEETS
    # Prior zero-balance placeholders can repeat an asset moved to another sheet.
    return [g for g in reference['groups'] if g['sheet'] in SHEETS and
            (abs(g['paidTotal'])>1 or any(abs(r['paidAmount'])>1 for r in g['details']))]


def disappeared_final_projects(previous, current, latest_map, prior_map):
    """History only: absent asset IDs, not changed names/sheets or partial reductions."""
    current_ids = {g['assetId'] for g in current_final_groups(current)}
    prior_groups = current_final_groups(previous)
    ids = [g['assetId'] for g in prior_groups]
    if len(ids) != len(set(ids)):
        raise ValueError('전월 최종본에 중복 자산번호가 있어 대체 비교를 할 수 없습니다.')
    projects = []
    for group in prior_groups:
        aid = group['assetId']
        amount = number(group['paidTotal'])
        if amount <= 1 or aid in current_ids:
            continue
        source = latest_map.get(aid, prior_map.get(aid, {}))
        dates = sorted(d['postingDate'] for d in group['details'] if d.get('postingDate') and number(d.get('paidAmount')) > 0)
        basis = f"전월 최종본에서 사라짐 · {previous['file']} / {group['sheet']} {group['sourceRow']}행 → {current['file']}"
        projects.append({
            'assetId': aid, 'wbs': group['wbs'], 'name': group['name'],
            'business': infer_business(group, group['sheet']), 'sheet': group['sheet'],
            'className': source.get('className', '건설중인자산'),
            'plannedAssetType': infer_asset_type(group, group['typeTotals']),
            'acquisitionDate': iso_date(source.get('acquisitionDate')), 'ageMonths': None,
            'lastAdditionDate': dates[-1] if dates else None, 'lastAdditionBasis': '전월 최종본 구성행',
            'depreciationComponents': [], 'unclassifiedAmount': 0,
            'beginning': 0, 'monthAdditions': 0, 'monthTransfers': clean_number(amount),
            'ending': 0, 'previousEnding': clean_number(amount), 'netChange': clean_number(-amount),
            'status': '당월 대체', 'reviewReason': basis,
            'transferBasis': 'prior-final-disappearance', 'transferBasisLabel': '최종본 비교 · 전월 잔액 기준',
            'previousFinalFile': previous['file'], 'previousFinalRow': group['sourceRow'],
            'currentFinalFile': current['file'], 'priorTypeTotals': group['typeTotals'],
            'schedulePaid': 0, 'scheduleGap': 0, 'contractTotal': 0,
            'expectedMonthlyDepreciation': 0, 'typeTotals': {},
            'sourceRegisterAmount': clean_number(number(latest_map.get(aid, {}).get('ending'))),
            'sourceDifference': clean_number(-number(latest_map.get(aid, {}).get('ending'))),
            'sourceRegisterMissing': aid not in latest_map,
            'costCenter': source.get('costCenter', ''), 'businessArea': group['businessArea'],
            'vendorCode': source.get('vendorCode', ''), 'vendorName': source.get('vendorName', ''),
        })
    return projects


def build_payload(root: Path, edited_final=False, final_file=None) -> dict[str, Any]:
    inputs = discover_inputs(root)
    if edited_final and final_file:
        from edited_final_update import select_final
        _,inputs['reference_final']=select_final(root,final_file)
        inputs['reference_period']=inputs['report_period']
    register_history = [load_register(path, period) for period, path in inputs["registers"].items()]
    latest = register_history[-1]
    prior = register_history[-2] if len(register_history) > 1 else None
    latest_map = {asset["assetId"]: asset for asset in latest["assets"]}
    prior_map = {asset["assetId"]: asset for asset in prior["assets"]} if prior else {}
    reference = load_reference_final(inputs["reference_final"], inputs["reference_period"])
    if edited_final:
        from management_rules import SHEETS
        if inputs['reference_period']!=inputs['report_period']:
            raise ValueError('당월 최종본이 필요합니다. 먼저 1단계를 실행해 주세요.')
        reference['groups']=current_final_groups(reference)
        for group in reference['groups']:
            for detail in group['details']:detail['sourceStatus']='편집 최종본 반영'
    group_map = {group["assetId"]: group for group in reference["groups"]}
    dashboard_assets=latest['assets']
    source_differences=[]
    metadata_issues=[]
    if edited_final:
        dashboard_assets=[]
        for aid in sorted(set(latest_map)|set(group_map)):
            source_asset=latest_map.get(aid)
            group=group_map.get(aid)
            source_ending=number(source_asset['ending']) if source_asset else 0
            final_ending=number(group['paidTotal']) if group else 0
            if abs(final_ending-source_ending)>1:
                source_differences.append({'assetId':aid,'registerAmount':source_ending,'finalAmount':final_ending,
                    'difference':clean_number(final_ending-source_ending),
                    'note':'최종본에만 존재' if not source_asset else '최종본 명세 없음 · 대시보드 미포함' if not group else '추가 전기·원천 추출시점 등 확인'})
            if not group:continue  # Never restore source-only balances as fake final detail.
            asset=dict(source_asset) if source_asset else {'assetId':aid,'className':'건설중인자산',
                'beginning':0,'ytdAdditions':0,'ytdTransfers':0,'acquisitionDate':None,
                'costCenter':'','vendorCode':'','vendorName':''}
            asset['ending']=final_ending
            asset['name']=group['name'];asset['wbs']=group['wbs'];asset['businessArea']=group['businessArea']
            asset['sourceRegisterMissing']=source_asset is None
            dashboard_assets.append(asset)
            if final_ending>1:
                for detail in group['details']:
                    if abs(detail['paidAmount'])<=1:continue
                    kind=detail['assetType'];life=number(detail.get('usefulLife'))
                    if kind not in USEFUL_LIFE_YEARS or (kind not in ('토지','비용처리') and not .1<=life<=100):
                        metadata_issues.append({'assetId':aid,'sheet':group['sheet'],'row':detail['sourceRow'],
                            'amount':detail['paidAmount'],'message':'분류·내용연수 확인 필요'})

    for asset in dashboard_assets:
        group = group_map.get(asset["assetId"])
        if edited_final and group:
            asset['name']=group['name'];asset['wbs']=group['wbs'];asset['businessArea']=group['businessArea']
        asset["business"] = infer_business(asset, group["sheet"] if group else "")
        asset["sheet"] = sheet_for_asset(asset, group)

    assets_by_wbs = {asset["wbs"]: asset for asset in dashboard_assets if asset["wbs"]}
    building_candidates = [asset for asset in dashboard_assets if "건물" in asset["className"] and asset["ending"] > 0]
    fallback_building_asset = max(building_candidates, key=lambda asset: asset["ending"])["assetId"] if building_candidates else ""
    ledger = load_ledger(inputs["ledger"], inputs["report_period"], assets_by_wbs, fallback_building_asset)

    periods = []
    for index, register in enumerate(register_history):
        previous = register_history[index - 1] if index else None
        current_totals = register["totals"]
        previous_totals = previous["totals"] if previous and previous["period"].year == register["period"].year else defaultdict(float)
        additions = current_totals.get("ytdAdditions", 0) - previous_totals.get("ytdAdditions", 0)
        transfers = current_totals.get("ytdTransfers", 0) - previous_totals.get("ytdTransfers", 0)
        periods.append({
            "key": register["period"].label,
            "label": register["period"].korean_label,
            "asOf": month_end(register["period"]).isoformat(),
            "ending": clean_number(current_totals.get("ending", 0)),
            "additions": clean_number(additions),
            "transfers": clean_number(transfers),
            "netChange": clean_number(additions - transfers),
            "assetCount": sum(1 for asset in register["assets"] if asset["ending"] > 0.5),
            "sourceFile": register["file"],
        })

    if edited_final:
        periods[-1].update(ending=clean_number(sum(g['paidTotal'] for g in reference['groups'])),
            assetCount=sum(a['ending']>.5 for a in dashboard_assets),sourceFile=reference['file'],
            balanceBasis='edited-final',movementSourceFile=latest['file'])

    report_end = month_end(inputs["report_period"])
    projects = []
    validation_groups: dict[str, list[dict[str, Any]]] = {name: [] for name in SHEET_ORDER}
    current_ledger_by_asset: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in ledger["currentMonthRows"]:
        if row["matchedAssetId"]:
            current_ledger_by_asset[row["matchedAssetId"]].append(row)

    candidate_assets = [
        asset for asset in dashboard_assets
        if abs(asset["ending"]) > 0.5
        or abs(asset["ytdTransfers"] - number(prior_map.get(asset["assetId"], {}).get("ytdTransfers"))) > 0.5
    ]
    for asset in candidate_assets:
        previous = prior_map.get(asset["assetId"])
        previous_additions = number(previous.get("ytdAdditions")) if previous else 0.0
        previous_transfers = number(previous.get("ytdTransfers")) if previous else 0.0
        month_additions = asset["ytdAdditions"] - previous_additions
        month_transfers = asset["ytdTransfers"] - previous_transfers
        previous_ending = number(previous.get("ending")) if previous else 0.0
        group = group_map.get(asset["assetId"])
        type_totals = group["typeTotals"] if group else {}
        dominant_type = infer_asset_type(asset, type_totals)
        last_date, last_basis = last_addition(asset, group, ledger, register_history, report_end)
        age_months = months_between(date.fromisoformat(last_date), report_end) if last_date else None
        schedule_paid = number(group.get("paidTotal")) if group else 0.0
        schedule_gap = asset["ending"] - schedule_paid
        if abs(asset["ending"]) <= 0.5 and month_transfers > 0.5:
            status = "당월 대체"
        elif abs(asset["ending"]) <= 0.5:
            status = "대체 완료"
        elif age_months is not None and age_months >= 24:
            status = "24개월 이상 미발생"
        elif age_months is not None and age_months >= 12:
            status = "12개월 이상 미발생"
        elif age_months is None:
            status = "최종 발생일 확인"
        elif previous is None and asset["ending"] > 0:
            status = "신규"
        elif month_additions > 0.5:
            status = "당월 증가"
        else:
            status = "진행중"
        reasons = []
        if age_months is None:
            reasons.append("최종 증가일 자료 확인 필요")
        if age_months is not None and age_months >= 12 and asset["ending"] > 0:
            reasons.append(f"마지막 증가 후 {age_months}개월 미발생")
        if month_transfers > 0.5:
            reasons.append("당월 대체 발생")
        if abs(schedule_gap) > 1:
            reasons.append("전월 명세 보완 필요")
        if not reasons:
            reasons.append("정상 진행")
        components = depreciation_components(asset, group,use_defaults=not edited_final)
        monthly_depreciation = sum(item["monthlyDepreciation"] for item in components)
        project = {
            "assetId": asset["assetId"],
            "wbs": asset["wbs"],
            "name": asset["name"],
            "business": asset["business"],
            "sheet": asset["sheet"],
            "className": asset["className"],
            "plannedAssetType": dominant_type,
            "acquisitionDate": asset["acquisitionDate"].isoformat() if asset["acquisitionDate"] else None,
            "ageMonths": age_months,
            "lastAdditionDate": last_date,
            "lastAdditionBasis": last_basis,
            "depreciationComponents": components,
            "unclassifiedAmount": clean_number(sum(abs(item["amount"]) for item in components if item["basis"] == "분류·내용연수 확인 필요")),
            "beginning": clean_number(asset["beginning"]),
            "monthAdditions": clean_number(month_additions),
            "monthTransfers": clean_number(month_transfers),
            "ending": clean_number(asset["ending"]),
            "previousEnding": clean_number(previous_ending),
            "netChange": clean_number(asset["ending"] - previous_ending),
            "status": status,
            "reviewReason": " · ".join(reasons),
            "schedulePaid": clean_number(schedule_paid),
            "scheduleGap": clean_number(schedule_gap),
            "sourceRegisterAmount": clean_number(number(latest_map.get(asset['assetId'],{}).get('ending'))),
            "sourceDifference": clean_number(asset['ending']-number(latest_map.get(asset['assetId'],{}).get('ending'))) if edited_final else 0,
            "sourceRegisterMissing": asset.get('sourceRegisterMissing',False),
            "contractTotal": clean_number(number(group.get("contractTotal"))) if group else 0,
            "expectedMonthlyDepreciation": clean_number(monthly_depreciation),
            "typeTotals": type_totals,
            "costCenter": asset["costCenter"],
            "businessArea": asset["businessArea"],
            "vendorCode": asset["vendorCode"],
            "vendorName": asset["vendorName"],
        }
        projects.append(project)

        details = list(group["details"]) if group else []
        if abs(schedule_gap) > 1 and not edited_final:
            documents = [row["document"] for row in current_ledger_by_asset.get(asset["assetId"], [])]
            details.append({
                "assetId": asset["assetId"],
                "wbs": asset["wbs"],
                "businessArea": asset["businessArea"],
                "postingDate": report_end.isoformat(),
                "description": "SAP 자산대장 증감 자동반영 (품의·거래처 정보 보완 필요)",
                "assetType": dominant_type,
                "costCenter": asset["costCenter"],
                "division": asset["business"],
                "vendorCode": asset["vendorCode"],
                "vendorName": asset["vendorName"],
                "contractAmount": 0,
                "paymentTerm": "자동검증",
                "paidAmount": clean_number(schedule_gap),
                "approval": "",
                "drafter": "",
                "waApproval": "",
                "document": "",
                "temporaryAsset": "",
                "transferDocument": ", ".join(documents[:8]),
                "usefulLife": USEFUL_LIFE_YEARS.get(dominant_type),
                "sourceStatus": "자동반영 · 상세 보완",
            })
        validation_groups.setdefault(asset["sheet"], []).append({
            "assetId": asset["assetId"],
            "wbs": asset["wbs"],
            "businessArea": asset["businessArea"],
            "name": asset["name"],
            "plannedAssetType": dominant_type,
            "contractTotal": project["contractTotal"],
            "paidTotal": project["ending"],
            "status": status,
            "details": details,
        })

    final_comparison = None
    if edited_final:
        previous_period, previous_path, comparison_note = previous_final(root, inputs['report_period'])
        transfers = []
        if previous_path:
            previous_reference = load_reference_final(previous_path, previous_period)
            transfers = disappeared_final_projects(previous_reference, reference, latest_map, prior_map)
            projects.extend(transfers)
        final_comparison = {'period': previous_period.label,
            'file': previous_path.name if previous_path else None,
            'status': '완료' if previous_path else '확인 필요', 'note': comparison_note,
            'count': len(transfers), 'amount': clean_number(sum(p['monthTransfers'] for p in transfers)) if previous_path else None}
        transferred_ids = {p['assetId'] for p in transfers}
        for difference in source_differences:
            if difference['assetId'] in transferred_ids:
                difference['note'] = '당월 최종본 명세 없음 · 최종본 비교 대체로 표시 (잔액 0)'

    projects.sort(key=lambda item: (
        0 if item["ageMonths"] is not None and item["ageMonths"] >= 12 else 1,
        -number(item["ending"]),
        item["assetId"],
    ))
    for groups in validation_groups.values():
        groups.sort(key=lambda item: item["assetId"])

    latest_total = number(latest["totals"].get("ending"))
    beginning_total = number(latest["totals"].get("beginning"))
    ledger_ending = beginning_total + number(ledger["sumThroughCutoff"])
    prior_ending_total = number(prior["totals"].get("ending")) if prior else beginning_total
    latest_additions = number(latest["totals"].get("ytdAdditions")) - (number(prior["totals"].get("ytdAdditions")) if prior else 0)
    latest_transfers = number(latest["totals"].get("ytdTransfers")) - (number(prior["totals"].get("ytdTransfers")) if prior else 0)
    bridge_ending = prior_ending_total + latest_additions - latest_transfers

    def check(name: str, expected: float, actual: float, tolerance: float = 1.0, note: str = "") -> dict[str, Any]:
        difference = actual - expected
        return {
            "name": name,
            "expected": clean_number(expected),
            "actual": clean_number(actual),
            "difference": clean_number(difference),
            "status": "일치" if abs(difference) <= tolerance else "확인 필요",
            "note": note,
        }

    checks = [
        check("자산대장 롤포워드", beginning_total + number(latest["totals"].get("ytdAdditions")) - number(latest["totals"].get("ytdTransfers")), latest_total),
        check("월간 증감 브리지", bridge_ending, latest_total),
        check("원장 잔액 대사", latest_total, ledger_ending, note="기초자산가액 + 기준월까지 건설중인자산원장"),
        check("편집 최종본 명세 대사" if edited_final else "전월 최종본 명세 대사", latest_total, reference["summaryValue"], note=f"반영 파일: {reference['file']}" if edited_final else f"승계 기준: {reference['file']}"),
    ]
    if edited_final:
        for item in checks:item['blocking']=False
        checks.append({'name': '전월 최종본 대체 비교', 'expected': 0, 'actual': 0 if final_comparison['file'] else None,
            'difference': 0 if final_comparison['file'] else None, 'status': '일치' if final_comparison['file'] else '확인 필요',
            'note': final_comparison['note'] or f"{final_comparison['file']} · 사라진 프로젝트 {final_comparison['count']}건 · 전월 잔액 기준",
            'blocking': False})
        checks[3]['name']='최종본 · 자산대장 차이 (참고)'
        checks[3]['note']=f'최종본 우선 반영 · 차이 {reference["summaryValue"]-latest_total:,.0f}원 · 원천자료 재추출 시 확인'
        checks.extend([
            {**check('최종본 · 원장 차이 (참고)',ledger_ending,reference['summaryValue'],note='추가 전기·추출시점 차이는 갱신을 막지 않습니다.'),'blocking':False},
            {**check('자산별 원천 차이',0,len(source_differences),tolerance=0,note='자산별 차이는 아래 표에서 확인할 수 있습니다.'),'unit':'건','blocking':False},
            {**check('상각정보 보완',0,len(metadata_issues),tolerance=0,note='정보가 없는 구성행은 잔액에 포함하고 상각 예측에서 제외합니다.'),'unit':'행','blocking':False},
            {**check('대시보드 · 최종본 잔액',reference['summaryValue'],sum(p['ending'] for p in projects),note='대시보드 잔액은 선택한 최종본을 기준으로 합니다.'),'blocking':True},
        ])
    validation_file = f"{inputs['report_period'].label} 건설중인자산 검증_{datetime.now().strftime('%y%m%d')}.xlsx"
    status_counts = Counter(project["status"] for project in projects)
    aging_counts = Counter()
    aging_amounts = defaultdict(float)
    for project in projects:
        if number(project["ending"]) <= 0.5:
            continue
        age = project["ageMonths"]
        bucket = "발생일 미상" if age is None else "0~3개월" if age <= 3 else "4~6개월" if age <= 6 else "7~12개월" if age <= 12 else "13~24개월" if age <= 24 else "24개월 초과"
        aging_counts[bucket] += 1
        aging_amounts[bucket] += number(project["ending"])

    business_summary: dict[str, dict[str, Any]] = defaultdict(lambda: {"count": 0, "ending": 0.0, "monthAdditions": 0.0, "monthTransfers": 0.0, "expectedMonthlyDepreciation": 0.0})
    for project in projects:
        if project.get('transferBasis') == 'prior-final-disappearance':
            continue  # Detail history must not inflate the active forecast business counts.
        if number(project["ending"]) <= 0.5 and number(project["monthTransfers"]) <= 0.5:
            continue
        bucket = business_summary[project["business"]]
        bucket["count"] += 1
        for field in ("ending", "monthAdditions", "monthTransfers", "expectedMonthlyDepreciation"):
            bucket[field] += number(project[field])

    return {
        "schemaVersion": 1,
        "assetTypeDefaults": USEFUL_LIFE_YEARS,
        "reportPeriod": inputs["report_period"].label,
        "reportMonth": inputs["report_period"].korean_label,
        "reportDate": report_end.isoformat(),
        "builtAt": datetime.now().isoformat(timespec="seconds"),
        "validationFile": validation_file,
        "periods": periods,
        "projects": projects,
        "businessSummary": {name: {field: clean_number(value) for field, value in values.items()} for name, values in business_summary.items()},
        "statusCounts": dict(status_counts),
        "aging": [
            {"bucket": bucket, "count": aging_counts[bucket], "amount": clean_number(aging_amounts[bucket])}
            for bucket in ["0~3개월", "4~6개월", "7~12개월", "13~24개월", "24개월 초과", "발생일 미상"]
            if aging_counts[bucket]
        ],
        "checks": checks,
        "sourceDifferences":source_differences,
        "finalComparison":final_comparison,
        "metadataIssues":metadata_issues,
        "ledger": ledger,
        "validationSheets": validation_groups,
        "source": {
            "balanceBasis":'edited-final' if edited_final else 'asset-register',
            "finalTotal":reference['summaryValue'] if edited_final else None,
            "assetRegister": latest["file"],
            "priorAssetRegister": prior["file"] if prior else None,
            "ledger": ledger["file"],
            "referenceFinal": reference["file"],
            "referencePeriod": inputs["reference_period"].label,
            "futureLedgerRowsExcluded": len(ledger["futureRows"]),
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", required=True)
    parser.add_argument("--validation-file")
    parser.add_argument("--edited-final",action='store_true')
    parser.add_argument('--final-file')
    args = parser.parse_args()
    root = Path(args.project_root).resolve()
    payload = build_payload(root,edited_final=args.edited_final,final_file=args.final_file)
    if args.validation_file:
        if Path(args.validation_file).name!=args.validation_file or not args.validation_file.lower().endswith('.xlsx'):
            raise ValueError('검증 다운로드 파일명 오류')
        payload['validationFile']=args.validation_file
    private_directory = root / "private"
    publish_directory = root / "publish"
    private_directory.mkdir(parents=True, exist_ok=True)
    publish_directory.mkdir(parents=True, exist_ok=True)
    (private_directory / "latest-data.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (private_directory / "validation-data.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    template_path = root / "dashboard-template.html"
    template = template_path.read_text(encoding="utf-8")
    marker = "__CIP_DATA_JSON__"
    if marker not in template:
        raise RuntimeError("대시보드 템플릿에 데이터 자리표시자가 없습니다.")
    embedded = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = template.replace(marker, embedded)
    output_path = publish_directory / f"건설중인자산_대시보드_{datetime.now().strftime('%H%M%S')}.html"
    output_path.write_text(html, encoding="utf-8")
    required_checks=[c for c in payload['checks'] if c.get('blocking')] if args.edited_final else payload['checks'][:3]
    failed = [check for check in required_checks if check["status"] != "일치"]
    if failed:
        raise RuntimeError(f"핵심 원천 대사 실패: {failed}")
    print(f"[생성] {output_path}")
    print('[기준] 편집 최종본 잔액 반영 · 원천 차이는 참고 표시' if args.edited_final else f"[대사] 자산대장·원장·월간 브리지 {len(payload['checks'][:3])}건 일치")
    if payload["source"]["futureLedgerRowsExcluded"]:
        print(f"[제외] 기준월 후 원장 {payload['source']['futureLedgerRowsExcluded']:,}건")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
